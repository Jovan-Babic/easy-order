from fastapi import FastAPI, APIRouter, Header, HTTPException, Depends, File, Query, Request, UploadFile
from fastapi.staticfiles import StaticFiles
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from contextlib import asynccontextmanager
from motor.motor_asyncio import AsyncIOMotorClient
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError
import asyncio
import hashlib
import json
import os
import re
import hmac
import time
import secrets
import smtplib
import logging
import mimetypes
from pathlib import Path
from pydantic import BaseModel, Field, EmailStr, PrivateAttr
from typing import List, Literal, Optional, Dict, Any, Tuple
from enum import Enum
import uuid
import bcrypt
import jwt
from datetime import date, datetime, timezone, timedelta
from email.message import EmailMessage
from email.utils import formataddr
import html as html_lib
from urllib.parse import quote
import cloudinary
import cloudinary.uploader

from calc import compute_order_totals, effective_qty, line_net, line_vat


ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Production = explicit APP_ENV=production, or any public Vercel deployment
# (production and preview URLs are both reachable from the internet).
IS_PRODUCTION = (
    os.environ.get("APP_ENV", "").lower() == "production"
    or os.environ.get("VERCEL_ENV", "") in ("production", "preview")
)
DEV_JWT_SECRET = "dev-insecure-secret-change-me"
DEFAULT_SUPERADMIN_PASSWORD = "ChangeMe123!"

JWT_SECRET = os.environ.get("JWT_SECRET", "")
if not JWT_SECRET:
    if IS_PRODUCTION:
        # Fail fast: with the public dev fallback anyone could sign a valid
        # superadmin token, so refusing to start is the only safe option.
        raise RuntimeError("JWT_SECRET must be set in production")
    logger.warning("JWT_SECRET not set - using insecure dev secret (never do this in production)")
    JWT_SECRET = DEV_JWT_SECRET
JWT_EXPIRE_MINUTES = int(os.environ.get("JWT_EXPIRE_MINUTES", "720"))
PASSWORD_RESET_EXPIRE_MINUTES = int(os.environ.get("PASSWORD_RESET_EXPIRE_MINUTES", "30"))
SUPERADMIN_EMAIL = os.environ.get("SUPERADMIN_EMAIL", "admin@easyorder.dev")
SUPERADMIN_PASSWORD = os.environ.get("SUPERADMIN_PASSWORD", DEFAULT_SUPERADMIN_PASSWORD)
DEMO_ADMIN_EMAIL = os.environ.get("DEMO_ADMIN_EMAIL", "demo-admin@easyorder.dev")
DEMO_ADMIN_PASSWORD = os.environ.get("DEMO_ADMIN_PASSWORD", "ChangeMe123!")
# Demo client/admin/customers/products for local tests and CI. See seed_data().
SEED_DEMO_DATA = os.environ.get("SEED_DEMO_DATA", "false").lower() == "true"
CORS_ORIGINS = os.environ.get("CORS_ORIGINS", "*").split(",")
SMTP_HOST = os.environ.get("SMTP_HOST", "")
SMTP_PORT = int(os.environ.get("SMTP_PORT", "587"))
SMTP_USER = os.environ.get("SMTP_USER", "")
SMTP_PASSWORD = os.environ.get("SMTP_PASSWORD", "")
SMTP_FROM = os.environ.get("SMTP_FROM", "noreply@easyorder.dev")
SMTP_USE_TLS = os.environ.get("SMTP_USE_TLS", "true").lower() == "true"
RESET_WEB_URL = os.environ.get("RESET_WEB_URL", "http://localhost:3000/reset-password")
RESET_MOBILE_SCHEME = os.environ.get("RESET_MOBILE_SCHEME", "easy-order://reset-password")
PASSWORD_RESET_DEBUG_TOKEN_IN_RESPONSE = os.environ.get("PASSWORD_RESET_DEBUG_TOKEN_IN_RESPONSE", "false").lower() == "true"
APP_UPDATE_FILE = ROOT_DIR / "public" / "app" / "app-update.json"

# Public URLs put into invite emails. On Vercel the backend URL falls back to
# the project's production domain (a Vercel system env var).
PUBLIC_BACKEND_URL = (
    os.environ.get("PUBLIC_BACKEND_URL")
    or (f"https://{os.environ['VERCEL_PROJECT_PRODUCTION_URL']}" if os.environ.get("VERCEL_PROJECT_PRODUCTION_URL") else "")
    or "http://localhost:8000"
).rstrip("/")
ADMIN_WEB_URL = (os.environ.get("ADMIN_WEB_URL") or RESET_WEB_URL.rsplit("/reset-password", 1)[0]).rstrip("/")

# Brute-force protection (counted per email, stored in Mongo so it survives
# serverless cold starts and is shared by all instances).
LOGIN_MAX_FAILURES = int(os.environ.get("LOGIN_MAX_FAILURES", "5"))
LOGIN_LOCKOUT_MINUTES = int(os.environ.get("LOGIN_LOCKOUT_MINUTES", "15"))
FORGOT_PASSWORD_MAX_PER_HOUR = int(os.environ.get("FORGOT_PASSWORD_MAX_PER_HOUR", "3"))

# Product images. Vercel rejects request bodies over 4.5 MB, so the limit
# stays under that; Cloudinary additionally shrinks what it stores.
IMAGE_MAX_BYTES = int(os.environ.get("IMAGE_MAX_BYTES", str(4 * 1024 * 1024)))
IMAGE_ALLOWED_TYPES = {"image/jpeg", "image/png", "image/webp", "image/heic", "image/heif"}
PRODUCT_IMAGE_FOLDER = "easy-order/products"
CLIENT_LOGO_FOLDER = "easy-order/clients"

mongo_client: AsyncIOMotorClient = None
db = None  # AsyncIOMotorDatabase

def configure_cloudinary() -> None:
    """Configure Cloudinary from either CLOUDINARY_URL or explicit CLOUDINARY_* vars."""
    cloudinary_url = os.environ.get("CLOUDINARY_URL")
    cloud_name = os.environ.get("CLOUDINARY_CLOUD_NAME")
    api_key = os.environ.get("CLOUDINARY_API_KEY")
    api_secret = os.environ.get("CLOUDINARY_API_SECRET")

    if cloudinary_url:
        cloudinary.config(cloudinary_url=cloudinary_url, secure=True)
        return

    cloudinary.config(
        cloud_name=cloud_name,
        api_key=api_key,
        api_secret=api_secret,
        secure=True,
    )


configure_cloudinary()


@asynccontextmanager
async def lifespan(app: FastAPI):
    global mongo_client, db
    mongo_client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = mongo_client[os.environ["DB_NAME"]]
    await db.customers.create_index("client_id")
    await db.products.create_index("client_id")
    await db.orders.create_index("client_id")
    await db.orders.create_index([("client_id", 1), ("status", 1), ("created_at", -1)])
    await db.orders.create_index([("client_id", 1), ("created_by_user_id", 1), ("created_at", -1)])
    try:
        await db.orders.create_index(
            [("client_id", 1), ("invoice_number", 1)],
            unique=True,
            partialFilterExpression={"invoice_number": {"$type": "string"}},
        )
    except Exception as exc:
        logger.error("Could not create unique index on orders.invoice_number: %s", exc)
    try:
        await db.products.create_index(
            [("client_id", 1), ("barcode", 1)],
            unique=True,
            partialFilterExpression={"barcode": {"$type": "string"}},
        )
    except Exception as exc:
        logger.error("Could not create unique index on products.barcode: %s", exc)
    try:
        await db.products.create_index(
            [("client_id", 1), ("package_barcode", 1)],
            unique=True,
            partialFilterExpression={"package_barcode": {"$type": "string"}},
        )
    except Exception as exc:
        logger.error("Could not create unique index on products.package_barcode: %s", exc)
    try:
        await db.payments.create_index([("client_id", 1), ("created_at", -1)])
    except Exception as exc:
        logger.error("Could not create index on payments: %s", exc)
    try:
        await db.subscription_events.create_index([("client_id", 1), ("created_at", -1)])
    except Exception as exc:
        logger.error("Could not create index on subscription_events: %s", exc)
    try:
        # One batch per product and expiry date (expiry_date null = undated stock).
        await db.stock_batches.create_index([("product_id", 1), ("expiry_date", 1)], unique=True)
        await db.stock_batches.create_index([("client_id", 1), ("expiry_date", 1)])
    except Exception as exc:
        logger.error("Could not create indexes on stock_batches: %s", exc)
    await _migrate_user_emails_to_lowercase()
    try:
        await db.users.create_index("email", unique=True)
    except Exception as exc:  # e.g. two legacy accounts that only differed by case
        logger.error("Could not create unique index on users.email: %s", exc)
    await _seed_units()
    await refresh_settings(force=True)
    await _migrate_calendar_billing()
    await db.password_reset_tokens.create_index("token_hash", unique=True)
    await db.password_reset_tokens.create_index("expires_at")
    await db.password_reset_tokens.create_index("user_id")
    await db.auth_attempts.create_index("key")
    # TTL cleanup only - the limits themselves are enforced by querying created_at.
    await db.auth_attempts.create_index("created_at", expireAfterSeconds=24 * 3600)
    await seed_data()
    logger.info("Backend started with MongoDB (%s)", os.environ["DB_NAME"])
    yield
    mongo_client.close()


DEFAULT_EXPIRY_ALERT_DAYS = [5, 15, 30]  # ascending (default 30/15/5 days)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---------------- Auth / role models ----------------
class Role(str, Enum):
    SUPERADMIN = "superadmin"
    ADMIN = "admin"
    OPERATOR = "operator"  # UI: "Komercijalista" - creates orders
    WAREHOUSE = "warehouse"  # UI: "Magacin" - processes orders


# Roles an admin may create/edit/delete (always within their own client).
ADMIN_MANAGEABLE_ROLES = {Role.OPERATOR, Role.WAREHOUSE}


# ---------------- Modules (PLAN_MODULI.md) ----------------
# What a client may use is a list of modules the superadmin turns on per
# client. Everything not listed here is the always-on base: catalog,
# customers, orders (create/edit/cancel), invoice from the phone, Excel import
# of products/customers.
class Module(str, Enum):
    WAREHOUSE = "warehouse"  # warehouse role, order flow, packing, barcode + scanner, delivery note
    STOCK = "stock"  # stock levels, receipts, counts, movements, deduction on shipping
    EXPIRY = "expiry"  # expiry dates, batches, FEFO, alerts
    REPORTS = "reports"  # order reports


ALL_MODULES = [m.value for m in Module]
# A module can only be on when the ones it builds on are on too.
MODULE_REQUIRES: Dict[Module, List[Module]] = {Module.STOCK: [Module.WAREHOUSE], Module.EXPIRY: [Module.STOCK]}


def effective_modules(client_doc: Optional[dict]) -> List[str]:
    """Clients created before modules existed (no field) keep everything."""
    stored = (client_doc or {}).get("modules")
    return list(ALL_MODULES) if stored is None else [m for m in ALL_MODULES if m in stored]


def validated_modules(values: List[Module]) -> List[str]:
    chosen = {Module(v) for v in values}
    for module in Module:
        for needed in MODULE_REQUIRES.get(module, []):
            if module in chosen and needed not in chosen:
                raise HTTPException(
                    status_code=400, detail=f"Module '{module.value}' requires module '{needed.value}'"
                )
    return [m for m in ALL_MODULES if Module(m) in chosen]


# ---------------- Subscriptions (PLAN_PRETPLATE.md) ----------------
# A client may have a subscription: a package (plan) valid until a date. No
# subscription = no expiry (older clients, internal/free accounts).
#   active -> (ends_at passes) grace -> locked -> (90 days locked) data purge
SUBSCRIPTION_GRACE_DAYS = int(os.environ.get("SUBSCRIPTION_GRACE_DAYS", "14"))
SUBSCRIPTION_PURGE_AFTER_DAYS = int(os.environ.get("SUBSCRIPTION_PURGE_AFTER_DAYS", "90"))
SUBSCRIPTION_REMINDER_DAYS = int(os.environ.get("SUBSCRIPTION_REMINDER_DAYS", "14"))
CRON_SECRET = os.environ.get("CRON_SECRET", "")
# The purge really deletes only when this is "true"; otherwise the daily job
# just reports what it would delete (a dry run).
AUTO_PURGE_ENABLED = os.environ.get("AUTO_PURGE_ENABLED", "false").lower() == "true"


# ---------------- System settings (PLAN_NAPLATA_KALENDARSKA.md, N1) ----------------
# Business values the superadmin edits in the portal ("Podešavanja"). One document
# in `app_settings`; a value that isn't stored falls back to env, then to code.
# Secrets, URLs and safety switches (AUTO_PURGE_ENABLED) stay in env.
SETTINGS_DEFAULTS: Dict[str, Any] = {
    # billing cycle
    "draft_day": 25,  # charge drafts for next month are made on this day
    "due_day": 1,  # payment is due on this day of the paid month
    "block_after_day": 5,  # unpaid after this day -> blocked
    "first_month_free": True,
    "reminder_days": SUBSCRIPTION_REMINDER_DAYS,  # before an annual subscription ends / before a purge
    "purge_after_days": SUBSCRIPTION_PURGE_AFTER_DAYS,
    "expiry_alert_days": list(DEFAULT_EXPIRY_ALERT_DAYS),  # default for new clients
    # payment details shown to clients (+ IPS QR)
    "payment_instructions_sr": "",
    "payment_instructions_en": "",
    "ips_account": "",
    "ips_recipient_name": "",
    "ips_payment_code": "289",
    "ips_purpose": "Pretplata Easy Order",
    "ips_reference_from_number": True,
    # issuer on the pro-forma invoice (all optional)
    "issuer_name": "",
    "issuer_address": "",
    "issuer_pib": "",
    "issuer_account": "",
    "issuer_note": "",
    # billing emails
    "billing_from_email": "",
    "billing_from_name": "",
    "billing_reply_to": "",
    "bilingual_emails": True,
}
SETTINGS_TTL_SECONDS = 60
_settings_cache: Dict[str, Any] = {}
_settings_loaded_at = 0.0


async def refresh_settings(force: bool = False) -> None:
    """Reload the stored settings at most once a minute (or now, after a change),
    so every serverless instance picks up an edit within a minute."""
    global _settings_cache, _settings_loaded_at
    if not force and time.monotonic() - _settings_loaded_at < SETTINGS_TTL_SECONDS:
        return
    try:
        doc = await db.app_settings.find_one({"_id": "global"}) or {}
    except Exception as exc:  # keep the last known values
        logger.warning("Could not load settings: %s", exc)
        return
    _settings_cache = {k: v for k, v in doc.items() if k in SETTINGS_DEFAULTS and v is not None}
    _settings_loaded_at = time.monotonic()


def setting(key: str) -> Any:
    value = _settings_cache.get(key)
    return SETTINGS_DEFAULTS[key] if value is None else value


def effective_settings() -> Dict[str, Any]:
    return {key: setting(key) for key in SETTINGS_DEFAULTS}


def subscription_state(sub: Optional[dict], today: Optional[date] = None) -> dict:
    """Where a subscription stands today. status: none | active | grace | locked."""
    if not sub:
        return {"status": "none"}
    today = today or datetime.now(timezone.utc).date()
    ends = date.fromisoformat(sub["ends_at"])
    calendar = bool(sub.get("billing_cycle"))
    block_day = int(setting("block_after_day"))
    if calendar:
        # Calendar billing (PLAN_NAPLATA_KALENDARSKA.md): paid up to the end of a
        # month, then usable until `block_after_day` of the next month.
        following = _add_months(date(ends.year, ends.month, 1), 1)
        grace_ends = date(following.year, following.month, block_day)
        extension = sub.get("due_extension_until")
        if extension and date.fromisoformat(extension) > grace_ends:
            grace_ends = date.fromisoformat(extension)
    else:
        grace_ends = ends + timedelta(days=SUBSCRIPTION_GRACE_DAYS)
    state = {
        "status": "active",
        "ends_at": ends.isoformat(),
        "days_left": (ends - today).days,
        "grace_ends_at": grace_ends.isoformat(),
        "locked_since": None,
        "purge_at": None,
    }
    locked_since = None
    if sub.get("canceled_at"):
        locked_since = date.fromisoformat(str(sub["canceled_at"])[:10])
    elif today > grace_ends:
        locked_since = grace_ends + timedelta(days=1)
    elif today > ends:
        state["status"] = "grace"
    if locked_since is not None and locked_since <= today:
        state["status"] = "locked"
        state["locked_since"] = locked_since.isoformat()
        state["purge_at"] = (locked_since + timedelta(days=setting("purge_after_days"))).isoformat()
        state["lock_reason"] = "canceled" if sub.get("canceled_at") else "expired"
    elif calendar:
        # Blocks that don't come from the end date: they lift as soon as the
        # cause is gone and never lead to a purge.
        extension = sub.get("due_extension_until")
        extended = bool(extension) and date.fromisoformat(extension) >= today
        monthly = float(sub.get("monthly_amount") or 0)
        # Only debt whose due date has passed counts (a charge for next month is not late yet).
        overdue = sum(c["open"] for c in (sub.get("open_charges") or []) if c.get("due_date") and c["due_date"] < today.isoformat())
        if sub.get("manual_block"):
            state.update(status="locked", lock_reason="manual", locked_since=sub["manual_block"].get("since"))
        elif not extended and today.day > block_day and monthly > 0 and overdue >= monthly - 0.005:
            state.update(status="locked", lock_reason="debt", locked_since=date(today.year, today.month, block_day + 1).isoformat())
    if calendar:
        state["billing_cycle"] = sub["billing_cycle"]
        state["balance_due"] = round(float(sub.get("balance_due") or 0), 2)
        state["credit"] = round(float(sub.get("credit") or 0), 2)
        state["due_extension_until"] = sub.get("due_extension_until")
    return state


def _add_months(day: date, months: int) -> date:
    month_index = day.month - 1 + months
    year, month = day.year + month_index // 12, month_index % 12 + 1
    last = [31, 29 if year % 4 == 0 and (year % 100 != 0 or year % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][month - 1]
    return date(year, month, min(day.day, last))


async def client_has_module(client_id: Optional[str], module: Module) -> bool:
    if not client_id:  # superadmin: not tied to a client
        return True
    doc = await db.clients.find_one({"id": client_id}, {"_id": 0, "modules": 1})
    return module.value in effective_modules(doc)


class Client(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    address: Optional[str] = ""
    email: Optional[str] = ""
    phone: Optional[str] = ""
    pib: Optional[str] = ""
    # Invoice settings (RBAC_PLAN.md 5a). Edited by the superadmin only.
    registration_number: Optional[str] = ""
    bank_account: Optional[str] = ""
    logo: Optional[str] = ""
    invoice_prefix: Optional[str] = ""
    invoice_numbering: str = "auto"  # "auto" | "manual"
    # Days before expiry at which batches show up as warnings (warehouse/admin).
    expiry_alert_days: List[int] = Field(default_factory=lambda: list(setting("expiry_alert_days")))
    # Enabled modules (Module values); a stored client without the field has all.
    modules: List[str] = Field(default_factory=lambda: list(ALL_MODULES))
    # Subscription (superadmin manages it in /subscriptions): plan_id, plan_name,
    # starts_at, ends_at, note, source, canceled_at, purge_paused, purged_at.
    subscription: Optional[Dict[str, Any]] = None
    accounting_email: Optional[str] = None  # also gets the billing emails (PUT /billing/clients/{id}/emails)
    subscription_state: Optional[Dict[str, Any]] = None  # computed
    user_count: Optional[int] = None  # computed in GET /clients: active users of the client
    last_activity_at: Optional[str] = None  # computed in GET /clients: last time a user was active / ordered
    last_order_at: Optional[str] = None  # computed in GET /clients
    orders_30d: Optional[int] = None  # computed in GET /clients
    # Read-only: the number the next shipped order will get (auto numbering).
    invoice_next_seq: Optional[int] = None
    active: bool = True
    created_at: str = Field(default_factory=now_iso)


class ClientInput(BaseModel):
    name: str
    address: Optional[str] = ""
    email: Optional[str] = ""
    phone: Optional[str] = ""
    pib: Optional[str] = ""
    registration_number: Optional[str] = ""
    bank_account: Optional[str] = ""
    logo: Optional[str] = ""
    invoice_prefix: Optional[str] = ""
    invoice_numbering: str = "auto"
    # None = keep the stored thresholds (default 30/15/5 days).
    expiry_alert_days: Optional[List[int]] = None
    # Superadmin only. None = keep (new client: all modules).
    modules: Optional[List[Module]] = None
    # Only ever raises the counter (to continue a series from an old system).
    invoice_next_seq: Optional[int] = Field(default=None, ge=1)


class ClientCreateInput(ClientInput):
    # No password field: the admin gets a generated temporary password by
    # email (see _invite_user). An old client still sending admin_password
    # is ignored.
    admin_name: str
    admin_email: EmailStr
    # Optional: start the client on a package. Then `modules` comes from the plan.
    plan_id: Optional[str] = None
    subscription_ends_at: Optional[str] = None  # YYYY-MM-DD (moved to its month end); None = this month
    billing_cycle: Literal["monthly", "annual"] = "monthly"


class User(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    email: EmailStr
    name: str
    phone: Optional[str] = ""
    role: Role
    client_id: Optional[str] = None  # None only for SUPERADMIN
    active: bool = True
    # True until the user replaces their temporary (invite) password. While
    # set, only /auth/me, /auth/change-password and /auth/logout work.
    must_change_password: bool = False
    created_at: str = Field(default_factory=now_iso)
    # Set per request from the user's client (see get_authenticated_user); not stored.
    _modules: Optional[List[str]] = PrivateAttr(default=None)
    _subscription: Optional[dict] = PrivateAttr(default=None)


class UserRow(User):
    """Row of GET /users: the user plus the name of their client (superadmin list)."""
    client_name: Optional[str] = None


class UserOut(User):
    """/auth/me and login: the user plus the modules their client has
    (superadmin: all)."""
    modules: List[str] = Field(default_factory=lambda: list(ALL_MODULES))
    # Only when the client has a subscription: status active|grace, ends_at, days_left, grace_ends_at.
    subscription: Optional[Dict[str, Any]] = None
    # Live announcements from the system owner: [{id, level, message_sr, message_en}].
    announcements: List[Dict[str, Any]] = Field(default_factory=list)
    # Calendar billing banner (PLAN_NAPLATA_KALENDARSKA.md, N3); amounts only for the admin.
    billing: Optional[Dict[str, Any]] = None


class InviteResult(BaseModel):
    invite_sent: bool
    # Only returned when the invite email could NOT be sent (SMTP missing or
    # failing), so the creator can hand the password over another way.
    temporary_password: Optional[str] = None


class UserInviteResponse(User, InviteResult):
    pass


class ClientCreateResponse(InviteResult):
    client: Client
    admin_user: User


class UserInput(BaseModel):
    email: EmailStr
    name: str
    phone: Optional[str] = ""
    role: Role
    client_id: Optional[str] = None


class ChangePasswordInput(BaseModel):
    current_password: str
    new_password: str


class UserUpdateInput(BaseModel):
    name: Optional[str] = None
    phone: Optional[str] = None
    password: Optional[str] = None
    active: Optional[bool] = None
    role: Optional[Role] = None


class LoginInput(BaseModel):
    email: EmailStr
    password: str


class ForgotPasswordChannel(str, Enum):
    WEB = "web"
    MOBILE = "mobile"


class ForgotPasswordInput(BaseModel):
    email: EmailStr
    channel: ForgotPasswordChannel = ForgotPasswordChannel.WEB


class ResetPasswordInput(BaseModel):
    token: str
    new_password: str


class TestEmailInput(BaseModel):
    to_email: EmailStr
    subject: str = "Easy Order SMTP test"


class TestEmailResponse(BaseModel):
    ok: bool = True
    sent_to: EmailStr
    subject: str


class ForgotPasswordResponse(BaseModel):
    ok: bool = True
    message: str = "If this account exists, a password reset link has been sent"
    reset_token: Optional[str] = None


class GenericOkResponse(BaseModel):
    ok: bool = True


class TokenResponse(BaseModel):
    access_token: str
    token_type: str = "bearer"
    user: UserOut


# ---------------- Domain models ----------------
class Customer(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_id: str
    name: str
    address: Optional[str] = ""
    email: Optional[str] = ""
    phone: Optional[str] = ""
    pib: Optional[str] = ""
    created_at: str = Field(default_factory=now_iso)


class CustomerInput(BaseModel):
    name: str
    address: Optional[str] = ""
    email: Optional[str] = ""
    phone: Optional[str] = ""
    pib: Optional[str] = ""
    client_id: Optional[str] = None  # only honored for SUPERADMIN writes


class Product(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_id: str
    name: str
    image: Optional[str] = ""  # remote url or base64 data uri
    manufacturer: Optional[str] = ""
    price_no_vat: Optional[float] = 0
    vat_rate: Optional[float] = 20
    discount: Optional[float] = 0
    discounts: List[float] = Field(default_factory=lambda: [0])
    additional_discounts: List[int] = Field(default_factory=lambda: [0])
    pieces_per_package: Optional[int] = 0
    boxes_per_transport: Optional[int] = 0
    # Barcode of one piece (EAN etc.), unique per client. None = not set.
    barcode: Optional[str] = None
    # Barcode of one box (pieces_per_package pieces): a scan in receipt/count
    # adds the whole box. Unique per client together with `barcode` (no code
    # can be both). Needs pieces_per_package > 0. None = not set.
    package_barcode: Optional[str] = None
    # False = delisted/draft: hidden from sales reps and can't be ordered, but
    # stock, history and past orders stay. Documents without the field are active.
    active: bool = True
    # Expiry tracking (PLAN_ROK_TRAJANJA.md): stock is then kept in batches
    # (stock_batches: expiry date + pieces) and stock_qty is their sum.
    track_expiry: bool = False
    # Unit of measure (units.code). Every quantity of this product - stock,
    # ordered/picked, batches - is a whole number of the unit's BASE unit
    # (g, ml, mm, kom); the price is per `unit` (PLAN_JEDINICE_MERE.md).
    unit: str = "kom"
    # Pieces in the warehouse. None = stock is not tracked for this product.
    # Only changed through /stock/* and order shipments, never via ProductInput.
    stock_qty: Optional[int] = None
    created_at: str = Field(default_factory=now_iso)


class ProductOut(Product):
    """List response: stock plus what open orders (new/in_progress) reserve."""
    reserved_qty: int = 0
    available_qty: Optional[int] = None  # stock_qty - reserved_qty - expired_qty; None when untracked
    # Expiry info is for warehouse/admin only; sales reps always get the defaults.
    expired_qty: int = 0
    next_expiry: Optional[str] = None  # earliest expiry date among batches still in stock
    # Only set by GET /products/by-barcode: what the scanned code was.
    scan_unit: Optional[str] = None  # "piece" | "package"
    scan_qty: int = 1  # pieces one scan stands for (pieces_per_package for a box)


class ProductInput(BaseModel):
    name: str
    image: Optional[str] = ""
    manufacturer: Optional[str] = ""
    price_no_vat: Optional[float] = 0
    vat_rate: Optional[float] = 20
    # None = "not sent". On update this keeps the stored value; a default of 0
    # made "not sent" and "explicitly 0" indistinguishable.
    discount: Optional[float] = None
    discounts: Optional[List[float]] = None
    additional_discounts: Optional[List[int]] = None
    pieces_per_package: Optional[int] = 0
    boxes_per_transport: Optional[int] = 0
    # None = keep the stored barcode, "" = clear it.
    barcode: Optional[str] = None
    # Same rules as barcode (None = keep, "" = clear).
    package_barcode: Optional[str] = None
    # None = keep. Activating a product that has no price yet is refused.
    active: Optional[bool] = None
    # None = keep. Switching it on turns the current stock into an undated batch.
    track_expiry: Optional[bool] = None
    # Unit code. None = keep (create: "kom"). Fixed once the product has stock or orders.
    unit: Optional[str] = None
    client_id: Optional[str] = None  # only honored for SUPERADMIN writes


class PickedBatch(BaseModel):
    """Part of a packed line taken from one batch (snapshot kept on the order)."""
    batch_id: Optional[str] = None
    expiry_date: Optional[str] = None  # None = undated stock
    qty: int


class OrderItem(BaseModel):
    product_id: str
    name: str
    image: Optional[str] = ""
    manufacturer: Optional[str] = ""
    price_no_vat: Optional[float] = 0
    vat_rate: Optional[float] = 20
    pieces_per_package: Optional[int] = 0
    boxes_per_transport: Optional[int] = 0
    discount: Optional[float] = 0
    additional_discount: Optional[float] = 0
    # Snapshot of the product's unit: quantities are whole BASE units, the price
    # is per `unit` (= unit_factor base units). Older orders: kom / 1.
    unit: str = "kom"
    unit_symbol: str = "kom"
    unit_factor: int = 1
    ordered_qty: int = 0
    # Set by the warehouse while packing; None = not checked yet.
    picked_qty: Optional[int] = None
    # Only for products with expiry tracking, filled when packed/shipped.
    # Never shown to sales reps.
    picked_batches: Optional[List[PickedBatch]] = None


class OrderStatus(str, Enum):
    NEW = "new"
    IN_PROGRESS = "in_progress"
    SHIPPED = "shipped"
    REJECTED = "rejected"
    CANCELED = "canceled"


class StatusChange(BaseModel):
    from_status: Optional[str] = None
    to_status: str
    changed_by_user_id: Optional[str] = None
    changed_by_name: Optional[str] = None
    changed_by_role: Optional[str] = None
    changed_at: str = Field(default_factory=now_iso)
    note: Optional[str] = None


class Order(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_id: str
    customer_id: str
    customer_name: str
    items: List[OrderItem] = []
    item_count: Optional[int] = None
    client_name: Optional[str] = None
    status: OrderStatus = OrderStatus.NEW
    status_history: List[StatusChange] = []
    # Warehouse user who took the order (set on new -> in_progress).
    assigned_to_user_id: Optional[str] = None
    assigned_to_name: Optional[str] = None
    shipped_at: Optional[str] = None
    updated_at: Optional[str] = None  # last edit of a new order
    # Assigned once, when the order is first shipped; never changes or is freed.
    invoice_seq: Optional[int] = None
    invoice_number: Optional[str] = None
    # Who placed it. None on orders created before this field existed.
    created_by_user_id: Optional[str] = None
    created_by_name: Optional[str] = None
    created_at: str = Field(default_factory=now_iso)


class OrderItemOut(OrderItem):
    line_net: float = 0


class OrderTotals(BaseModel):
    subtotal: float
    vat: float
    grand: float


class OrderOut(Order):
    """Response shape: stored order + totals computed by calc.py, so clients
    can display amounts without re-implementing the discount/VAT formula.
    `totals` follow picked_qty once shipped; `ordered_totals` always follow
    the ordered quantities."""
    items: List[OrderItemOut] = []
    totals: OrderTotals
    ordered_totals: OrderTotals


class OrderItemInput(BaseModel):
    """What the client may decide per line. Everything else (name, price, VAT,
    packaging...) is snapshotted server-side from the stored Product. Extra
    fields that older clients still send (price_no_vat, name, ...) are ignored."""
    product_id: str
    ordered_qty: int
    discount: Optional[float] = None  # must be one of product.discounts; None = product default
    additional_discount: Optional[float] = None  # must be one of product.additional_discounts


class OrderInput(BaseModel):
    customer_id: str
    customer_name: Optional[str] = None  # ignored - taken from the stored Customer
    items: List[OrderItemInput] = []
    client_id: Optional[str] = None  # only honored for SUPERADMIN writes


class ClientStats(BaseModel):
    client_id: str
    client_name: str
    order_count: int
    customer_count: int
    product_count: int
    total_net: float
    total_vat: float
    total_grand: float


class StatsResponse(BaseModel):
    scope: str  # "global" | "client"
    totals: ClientStats
    by_client: List[ClientStats]


class DashboardSummary(BaseModel):
    order_count: int
    customer_count: int
    product_count: int
    active_customer_count: int
    total_net: float
    total_vat: float
    total_grand: float
    average_order_value: float
    average_items_per_order: float


class RevenuePoint(BaseModel):
    period: str
    order_count: int
    total_net: float
    total_vat: float
    total_grand: float


class ProductSalesStats(BaseModel):
    product_id: str
    name: str
    manufacturer: Optional[str] = ""
    ordered_qty: int  # sold quantity in whole base units (g, ml, mm, kom); shipped orders count what was packed
    unit_symbol: str = "kom"  # unit of the latest order line; ordered_qty / unit_factor is the amount in it
    unit_factor: int = 1
    order_count: int
    total_net: float
    total_vat: float
    total_grand: float


class CustomerSalesStats(BaseModel):
    customer_id: str
    customer_name: str
    order_count: int
    total_grand: float


class RecentOrderStats(BaseModel):
    id: str
    customer_name: str
    item_count: int
    total_grand: float
    created_at: str


class DashboardStatsResponse(BaseModel):
    scope: str
    from_date: Optional[str] = None
    to_date: Optional[str] = None
    summary: DashboardSummary
    revenue_by_period: List[RevenuePoint]
    top_products: List[ProductSalesStats]
    top_customers: List[CustomerSalesStats]
    recent_orders: List[RecentOrderStats]


# ---------------- Auth helpers ----------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(password: str, hashed: str) -> bool:
    return bcrypt.checkpw(password.encode("utf-8"), hashed.encode("utf-8"))


def create_access_token(user: dict) -> str:
    now = datetime.now(timezone.utc)
    payload = {
        "sub": user["id"],
        "email": user["email"],
        "role": user["role"],
        "client_id": user.get("client_id"),
        # Token version: bumped on password change/reset and role change, which
        # invalidates every token issued before. Missing = 0 (older tokens).
        "tv": user.get("token_version", 0),
        "iat": now,
        "exp": now + timedelta(minutes=JWT_EXPIRE_MINUTES),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm="HS256")


PASSWORD_RULES_MESSAGE = "Password must be at least 8 characters long and contain at least one letter and one digit"


def _password_is_strong_enough(password: str) -> bool:
    return (
        len(password) >= 8
        and any(c.isalpha() for c in password)
        and any(c.isdigit() for c in password)
    )


def ensure_strong_password(password: str) -> None:
    if not _password_is_strong_enough(password):
        raise HTTPException(status_code=400, detail=PASSWORD_RULES_MESSAGE)


def generate_temporary_password() -> str:
    # No look-alike characters (0/O, 1/l/I) since people may retype it from
    # the email by hand. Always contains a letter and a digit.
    letters = "abcdefghjkmnpqrstuvwxyzABCDEFGHJKLMNPQRSTUVWXYZ"
    digits = "23456789"
    chars = [secrets.choice(letters) for _ in range(8)] + [secrets.choice(digits) for _ in range(4)]
    secrets.SystemRandom().shuffle(chars)
    return "".join(chars)


def normalize_email(email: str) -> str:
    return email.strip().lower()


def _hash_reset_token(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _build_reset_link(token: str, channel: ForgotPasswordChannel) -> str:
    # Always the https web page, also for requests from the mobile app: email
    # clients (Gmail etc.) don't make custom-scheme links like easy-order://
    # clickable and spam filters distrust them, so those emails got lost. The
    # web reset page needs no login and works for every role; the new password
    # applies to the app right away. `channel` is still stored on the token.
    encoded = quote(token, safe="")
    sep = "&" if "?" in RESET_WEB_URL else "?"
    return f"{RESET_WEB_URL}{sep}token={encoded}"


def _to_aware_utc(value: datetime) -> datetime:
    if value.tzinfo is None:
        return value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc)


def _smtp_is_configured() -> bool:
    return bool(SMTP_HOST and SMTP_USER and SMTP_PASSWORD and SMTP_FROM)


def _send_smtp_email_sync(to_email: str, subject: str, plain_text_body: str) -> None:
    msg = EmailMessage()
    msg["From"] = SMTP_FROM
    msg["To"] = to_email
    msg["Subject"] = subject
    msg.set_content(plain_text_body)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as smtp:
        if SMTP_USE_TLS:
            smtp.starttls()
        smtp.login(SMTP_USER, SMTP_PASSWORD)
        smtp.send_message(msg)


def _build_reset_email_body(user_name: str, reset_link: str) -> str:
    return (
        f"Zdravo {user_name or ''},\n\n"
        "Primili smo zahtev za promenu lozinke za Easy Order.\n"
        f"Novu lozinku postavite preko ovog linka:\n{reset_link}\n\n"
        f"Link važi {PASSWORD_RESET_EXPIRE_MINUTES} minuta. Ako niste vi poslali zahtev, zanemarite ovaj email.\n\n"
        "---\n"
        f"Hi {user_name or 'there'},\n\n"
        "We received a request to reset your Easy Order password.\n"
        f"Use this link to set a new password:\n{reset_link}\n\n"
        f"This link expires in {PASSWORD_RESET_EXPIRE_MINUTES} minutes.\n"
        "If you did not request this, you can ignore this email."
    )


def _mobile_app_download_url() -> Optional[str]:
    """Absolute APK link from app-update.json, or None if no release is published."""
    try:
        data = json.loads(APP_UPDATE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if not data.get("enabled") or not data.get("download_url"):
        return None
    url = data["download_url"]
    return url if url.startswith("http") else f"{PUBLIC_BACKEND_URL}{url}"


def _build_invite_email_body(name: str, email: str, role: Role, temporary_password: str) -> str:
    lines = [
        f"Zdravo {name},",
        "",
        "Otvoren vam je nalog u aplikaciji Easy Order.",
        "",
        f"Email za prijavu: {email}",
        f"Privremena lozinka: {temporary_password}",
        "",
        "Pri prvoj prijavi bićete zamoljeni da postavite novu lozinku.",
        "",
    ]
    apk_url = _mobile_app_download_url()
    if apk_url:
        lines += [f"Android aplikacija (APK): {apk_url}", ""]
    if role in (Role.ADMIN, Role.SUPERADMIN, Role.WAREHOUSE):
        lines += [f"Administratorski portal: {ADMIN_WEB_URL}", ""]
    lines += [
        "---",
        f"Hi {name}, an Easy Order account was created for you.",
        f"Login: {email} / temporary password: {temporary_password}",
        "You will be asked to set a new password on first login.",
    ]
    return "\n".join(lines)


async def _send_invite_email(name: str, email: str, role: Role, temporary_password: str) -> bool:
    """Returns True if the email went out. Never raises - a failed email must
    not fail the account creation (the caller falls back to showing the
    password to the creator)."""
    if not _smtp_is_configured():
        logger.warning("SMTP not configured - invite email for %s not sent", email)
        return False
    body = _build_invite_email_body(name, email, role, temporary_password)
    try:
        await asyncio.to_thread(_send_smtp_email_sync, email, "Easy Order - pristup aplikaciji / account access", body)
        return True
    except Exception as exc:
        logger.exception("Failed to send invite email to %s: %s", email, exc)
        return False


async def _invite_user(user: User) -> InviteResult:
    """Stores a new user with a generated temporary password, emails it, and
    forces a password change on first login."""
    temporary_password = generate_temporary_password()
    raw = user.dict()
    raw["email"] = normalize_email(raw["email"])
    raw["must_change_password"] = True
    raw["token_version"] = 0
    raw["password_hash"] = hash_password(temporary_password)
    await db.users.insert_one({**raw, "_id": user.id})
    user.must_change_password = True
    sent = await _send_invite_email(user.name, raw["email"], user.role, temporary_password)
    return InviteResult(invite_sent=sent, temporary_password=None if sent else temporary_password)


# ---------------- Brute-force limits ----------------
async def _recent_attempts(key: str, window: timedelta) -> int:
    since = datetime.now(timezone.utc) - window
    return await db.auth_attempts.count_documents({"key": key, "created_at": {"$gte": since}})


async def _record_attempt(key: str) -> None:
    await db.auth_attempts.insert_one({"_id": str(uuid.uuid4()), "key": key, "created_at": datetime.now(timezone.utc)})


def _login_key(email: str) -> str:
    return f"login:{normalize_email(email)}"


def _build_test_email_body() -> str:
    return (
        "This is a test email from Easy Order.\n\n"
        "If you received this, Brevo SMTP is configured correctly and the backend can send mail."
    )


def _build_order_status_email(order: dict, status: str, note: Optional[str], actor_name: str) -> tuple[str, str]:
    """(subject, body) telling the sales rep their order was shipped/rejected."""
    shipped = status == OrderStatus.SHIPPED.value
    customer = order.get("customer_name", "")
    invoice = order.get("invoice_number")
    if shipped:
        subject = f"Easy Order - porudžbina poslata / order shipped: {customer}"
        sr = f"Porudžbina za kupca {customer} je poslata."
        en = f"The order for {customer} has been shipped."
    else:
        subject = f"Easy Order - porudžbina odbijena / order rejected: {customer}"
        sr = f"Porudžbina za kupca {customer} je odbijena."
        en = f"The order for {customer} was rejected."
    lines = [f"Zdravo {order.get('created_by_name') or ''},".replace("  ", " "), "", sr, en, ""]
    if shipped and invoice:
        lines.append(f"Broj fakture / Invoice number: {invoice}")
    if not shipped and note:
        lines.append(f"Razlog / Reason: {note}")
    lines.append(f"Obradio / Handled by: {actor_name}")
    return subject, "\n".join(lines)


async def _send_order_status_email(order: dict, status: str, note: Optional[str], actor: "User") -> None:
    """Best effort, never raises: a failed email must not undo the status
    change. Skipped when the actor placed the order themselves, when the
    creator is unknown (old orders) or has been deleted."""
    creator_id = order.get("created_by_user_id")
    if not creator_id or creator_id == actor.id:
        return
    if not _smtp_is_configured():
        logger.warning("SMTP not configured - order status email for order %s not sent", order.get("id"))
        return
    try:
        creator = await db.users.find_one({"id": creator_id}, {"_id": 0, "email": 1})
        if not creator or not creator.get("email"):
            return
        subject, body = _build_order_status_email(order, status, note, actor.name)
        await asyncio.to_thread(_send_smtp_email_sync, creator["email"], subject, body)
    except Exception as exc:
        logger.exception("Failed to send order status email for order %s: %s", order.get("id"), exc)


async def _send_reset_email(to_email: str, user_name: str, reset_link: str) -> None:
    if not _smtp_is_configured():
        logger.warning("SMTP not configured. Password reset link for %s: %s", to_email, reset_link)
        return

    body = _build_reset_email_body(user_name, reset_link)
    await asyncio.to_thread(_send_smtp_email_sync, to_email, "Easy Order - promena lozinke / password reset", body)


async def _send_test_email(to_email: str, subject: str) -> None:
    if not _smtp_is_configured():
        raise HTTPException(
            status_code=500,
            detail="SMTP is not configured. Check SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD and SMTP_FROM.",
        )

    await asyncio.to_thread(_send_smtp_email_sync, to_email, subject, _build_test_email_body())


async def _create_password_reset_token_record(user: dict, channel: ForgotPasswordChannel) -> str:
    raw_token = secrets.token_urlsafe(48)
    token_hash = _hash_reset_token(raw_token)
    now = datetime.now(timezone.utc)
    expires_at = now + timedelta(minutes=PASSWORD_RESET_EXPIRE_MINUTES)
    await db.password_reset_tokens.insert_one(
        {
            "_id": str(uuid.uuid4()),
            "token_hash": token_hash,
            "user_id": user["id"],
            "client_id": user.get("client_id"),
            "channel": channel.value,
            "created_at": now_iso(),
            "expires_at": expires_at,
            "used_at": None,
        }
    )
    return raw_token


async def _get_valid_reset_record(token: str) -> Optional[dict]:
    record = await db.password_reset_tokens.find_one(
        {"token_hash": _hash_reset_token(token)},
        {"_id": 0},
    )
    if not record:
        return None
    if record.get("used_at"):
        return None
    expires_at = record.get("expires_at")
    if not expires_at or _to_aware_utc(expires_at) <= datetime.now(timezone.utc):
        return None
    return record


def cloudinary_available() -> bool:
    cfg = cloudinary.config()
    return bool(cfg.cloud_name and cfg.api_key and cfg.api_secret)


_CLOUDINARY_PRODUCT_URL = re.compile(
    rf"res\.cloudinary\.com/[^/]+/image/upload/(?:.+/)?(?:v\d+/)?((?:{re.escape(PRODUCT_IMAGE_FOLDER)}|{re.escape(CLIENT_LOGO_FOLDER)})/[^/.]+)\.\w+$"
)


def _product_image_public_id(url: Optional[str]) -> Optional[str]:
    """public_id of an image this app uploaded, or None for anything else
    (external URLs, data URIs) - those are never deleted."""
    if not url:
        return None
    match = _CLOUDINARY_PRODUCT_URL.search(url)
    return match.group(1) if match else None


async def _delete_product_image(url: Optional[str]) -> None:
    """Best effort: frees space on the Cloudinary plan when a product image is
    replaced or the product is deleted. Order snapshots keep the old URL, but
    no screen displays order item images, so nothing visible breaks."""
    public_id = _product_image_public_id(url)
    if not public_id or not cloudinary_available():
        return
    try:
        await asyncio.to_thread(cloudinary.uploader.destroy, public_id, resource_type="image")
    except Exception as exc:
        logger.warning("Could not delete Cloudinary image %s: %s", public_id, exc)


async def find_user_by_email(email: str) -> Optional[dict]:
    # Emails are stored lowercase (see _migrate_user_emails_to_lowercase), so
    # this is an exact, index-backed lookup.
    return await db.users.find_one({"email": normalize_email(email)}, {"_id": 0})


async def _migrate_user_emails_to_lowercase() -> None:
    """One-off for accounts created before emails were normalized. Idempotent."""
    async for doc in db.users.find({}, {"_id": 1, "email": 1}):
        email = doc.get("email") or ""
        if email != normalize_email(email):
            await db.users.update_one({"_id": doc["_id"]}, {"$set": {"email": normalize_email(email)}})


async def client_is_active(client_id: Optional[str]) -> bool:
    """Superadmins have no client. For everyone else, a soft-deleted
    (active=False) or missing client locks out all of its users."""
    return (await _client_gate(client_id))["active"]


async def _client_gate(client_id: Optional[str]) -> dict:
    """active flag, modules and subscription state of a client in one lookup.
    Superadmins have no client: always active, every module, no subscription."""
    await refresh_settings()
    if not client_id:
        return {"active": True, "modules": list(ALL_MODULES), "subscription": {"status": "none"}}
    client = await db.clients.find_one({"id": client_id}, {"_id": 0, "active": 1, "modules": 1, "subscription": 1})
    if not client:
        return {"active": False, "modules": [], "subscription": {"status": "none"}}
    return {
        "active": client.get("active", True),
        "modules": effective_modules(client),
        "subscription": subscription_state(client.get("subscription")),
    }


security = HTTPBearer(auto_error=False)


async def user_out(raw: dict) -> UserOut:
    """User + the modules of their client (login / me responses)."""
    gate = await _client_gate(raw.get("client_id"))
    sub = gate["subscription"] if gate["subscription"]["status"] != "none" else None
    billing = await _billing_notice(raw) if sub and sub.get("billing_cycle") else None
    return UserOut(
        **{k: v for k, v in raw.items() if k not in ("modules", "subscription", "announcements", "billing")},
        modules=gate["modules"], subscription=sub, announcements=await _announcements_for(raw.get("client_id")),
        billing=billing,
    )


async def _touch_last_seen(raw: dict) -> None:
    """Remembers when a user was last active - at most one write per hour."""
    seen = raw.get("last_seen_at")
    if seen:
        try:
            if datetime.now(timezone.utc) - datetime.fromisoformat(seen) < timedelta(hours=1):
                return
        except ValueError:
            pass
    await db.users.update_one({"id": raw["id"]}, {"$set": {"last_seen_at": now_iso()}})


# A blocked client on calendar billing: its admin still gets in, but only to see
# what to pay (and to log out / change the password). Everyone else is shut out.
BLOCKED_ADMIN_PATHS = ("/api/auth/me", "/api/auth/logout", "/api/auth/change-password", "/api/billing/me")


def _blocked_admin_may(raw: dict, gate: dict, path: Optional[str]) -> bool:
    if raw.get("role") != Role.ADMIN.value or not gate["subscription"].get("billing_cycle"):
        return False
    return path is None or any(path == p or path.startswith(p + "/") for p in BLOCKED_ADMIN_PATHS)


async def get_authenticated_user(
    request: Request,
    creds: Optional[HTTPAuthorizationCredentials] = Depends(security),
) -> User:
    """Valid token + active user/client. Does NOT enforce the first-login
    password change - only for the few routes that must work before it."""
    await refresh_settings()
    if creds is None:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(creds.credentials, JWT_SECRET, algorithms=["HS256"])
    except jwt.PyJWTError:
        raise HTTPException(status_code=401, detail="Invalid or expired token")
    raw = await db.users.find_one({"id": payload.get("sub")}, {"_id": 0})
    if not raw or not raw.get("active", True):
        raise HTTPException(status_code=401, detail="User not found or inactive")
    if payload.get("tv", 0) != raw.get("token_version", 0):
        # Password was changed/reset (or role changed) after this token was issued.
        raise HTTPException(status_code=401, detail="Session expired, please log in again")
    # Checked on every request (not only at login) so deactivating a client
    # also cuts off tokens that were issued before the deactivation.
    gate = await _client_gate(raw.get("client_id"))
    if not gate["active"]:
        raise HTTPException(status_code=401, detail="Client account is disabled")
    if gate["subscription"]["status"] == "locked":
        if _blocked_admin_may(raw, gate, request.url.path):
            pass
        elif _blocked_admin_may(raw, gate, None):
            # 403, not 401: the session stays, the app sends the admin to the billing page.
            raise HTTPException(status_code=403, detail="Subscription blocked")
        else:
            raise HTTPException(status_code=401, detail="Subscription expired")
    await _touch_last_seen(raw)
    user = User(**raw)
    user._modules = gate["modules"]
    user._subscription = gate["subscription"]
    return user


async def get_current_user(user: User = Depends(get_authenticated_user)) -> User:
    """Default dependency for every business route."""
    if user.must_change_password:
        # 403, not 401: the session is valid, clients should send the user to
        # the change-password screen instead of logging them out.
        raise HTTPException(status_code=403, detail="Password change required")
    return user


def has_module(user: User, module: Module) -> bool:
    return user.role == Role.SUPERADMIN or module.value in (user._modules or [])


def require_module(module: Module, *roles: Role):
    """Role check (if roles are given) plus: the caller's client must have the module."""
    base = require_roles(*roles) if roles else get_current_user

    async def _dep(user: User = Depends(base)) -> User:
        if not has_module(user, module):
            raise HTTPException(status_code=403, detail=f"Module not enabled: {module.value}")
        return user
    return _dep


def require_roles(*roles: Role):
    async def _dep(user: User = Depends(get_current_user)) -> User:
        if user.role not in roles:
            raise HTTPException(status_code=403, detail="Forbidden")
        return user
    return _dep


# Admin-level writes (catalog, customers, deleting orders, user management).
# Operators only read the catalog and create orders; warehouse only reads.
require_manager = require_roles(Role.SUPERADMIN, Role.ADMIN)
require_stock_writer = require_roles(Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE)
# Roles that may place orders (warehouse processes them, never creates).
require_order_creator = require_roles(Role.SUPERADMIN, Role.ADMIN, Role.OPERATOR)


# ---------------- Tenant scoping helpers ----------------
async def resolve_write_client_id(user: User, payload_client_id: Optional[str]) -> str:
    if user.role == Role.SUPERADMIN:
        if not payload_client_id:
            raise HTTPException(status_code=400, detail="client_id is required for superadmin writes")
        if not await db.clients.find_one({"id": payload_client_id}, {"_id": 0}):
            raise HTTPException(status_code=404, detail="Client not found")
        return payload_client_id
    return user.client_id


def normalize_additional_discounts(values: Optional[List[int]]) -> List[int]:
    if not values:
        return [0]

    allowed = set(range(0, 16))
    normalized = sorted({int(v) for v in values if int(v) in allowed})
    return normalized or [0]


def normalize_discounts(values: Optional[List[float]], default_discount: float) -> List[float]:
    base = values if values else [default_discount]
    normalized = sorted({max(0.0, min(100.0, float(v))) for v in base})
    default_normalized = max(0.0, min(100.0, float(default_discount)))
    if default_normalized not in normalized:
        normalized.append(default_normalized)
        normalized.sort()
    return normalized or [0.0]


def _scope_query(user: User) -> dict:
    if user.role == Role.SUPERADMIN:
        return {}
    return {"client_id": user.client_id}


async def get_scoped_or_404(collection_name: str, item_id: str, user: User) -> dict:
    item = await db[collection_name].find_one({"id": item_id}, {"_id": 0})
    if not item:
        raise HTTPException(status_code=404, detail="Not found")
    if user.role != Role.SUPERADMIN and item.get("client_id") != user.client_id:
        # 404, not 403 - avoids leaking cross-tenant existence via status code.
        raise HTTPException(status_code=404, detail="Not found")
    return item


# ---------------- Auth routes ----------------
api_router = APIRouter(prefix="/api")


@api_router.post("/auth/login", response_model=TokenResponse)
async def login(inp: LoginInput):
    key = _login_key(inp.email)
    window = timedelta(minutes=LOGIN_LOCKOUT_MINUTES)
    # Counted per email (not IP): admin-web proxies logins server-to-server,
    # so every portal login would share the same IP.
    if await _recent_attempts(key, window) >= LOGIN_MAX_FAILURES:
        raise HTTPException(
            status_code=429,
            detail=f"Too many failed login attempts. Try again in {LOGIN_LOCKOUT_MINUTES} minutes.",
        )
    raw = await find_user_by_email(inp.email)
    if not raw or not raw.get("active", True) or not verify_password(inp.password, raw["password_hash"]):
        await _record_attempt(key)
        raise HTTPException(status_code=401, detail="Invalid email or password")
    await db.auth_attempts.delete_many({"key": key})
    # After the password check, so this can't be used to probe which emails exist.
    gate = await _client_gate(raw.get("client_id"))
    if not gate["active"]:
        raise HTTPException(status_code=403, detail="Client account is disabled")
    if gate["subscription"]["status"] == "locked" and not _blocked_admin_may(raw, gate, None):
        raise HTTPException(status_code=403, detail="Subscription expired")
    stamp = now_iso()
    await db.users.update_one({"id": raw["id"]}, {"$set": {"last_login_at": stamp, "last_seen_at": stamp}})
    token = create_access_token(raw)
    return TokenResponse(access_token=token, user=await user_out(raw))


@api_router.post("/auth/forgot-password", response_model=ForgotPasswordResponse)
async def forgot_password(inp: ForgotPasswordInput):
    debug_token: Optional[str] = None
    key = f"forgot:{normalize_email(inp.email)}"
    if await _recent_attempts(key, timedelta(hours=1)) >= FORGOT_PASSWORD_MAX_PER_HOUR:
        # Same generic response as always - a 429 here would reveal nothing
        # useful to the user and would let an attacker probe the limiter.
        logger.warning("Forgot-password rate limit hit for %s", inp.email)
        return ForgotPasswordResponse()
    await _record_attempt(key)
    user = await find_user_by_email(inp.email)

    if user and user.get("active", True) and await client_is_active(user.get("client_id")):
        raw_token = await _create_password_reset_token_record(user, inp.channel)
        reset_link = _build_reset_link(raw_token, inp.channel)
        try:
            await _send_reset_email(user["email"], user.get("name", ""), reset_link)
        except Exception as exc:
            logger.exception("Failed to send reset email for %s: %s", user["email"], exc)
        if PASSWORD_RESET_DEBUG_TOKEN_IN_RESPONSE:
            debug_token = raw_token

    return ForgotPasswordResponse(reset_token=debug_token)


@api_router.post("/auth/reset-password", response_model=GenericOkResponse)
async def reset_password(inp: ResetPasswordInput):
    ensure_strong_password(inp.new_password)

    record = await _get_valid_reset_record(inp.token)
    if not record:
        raise HTTPException(status_code=400, detail="Invalid or expired reset token")

    user = await db.users.find_one({"id": record["user_id"]}, {"_id": 0})
    if not user or not user.get("active", True):
        raise HTTPException(status_code=400, detail="Invalid or expired reset token")

    await _set_password(user["id"], inp.new_password)
    # Invalidate all outstanding reset tokens for this user after a successful reset.
    await db.password_reset_tokens.update_many(
        {"user_id": user["id"], "used_at": None},
        {"$set": {"used_at": now_iso()}},
    )

    return GenericOkResponse()


@api_router.post("/auth/test-email", response_model=TestEmailResponse)
async def test_email(inp: TestEmailInput, current_user: User = Depends(require_roles(Role.SUPERADMIN, Role.ADMIN))):
    try:
        await _send_test_email(inp.to_email, inp.subject)
    except Exception as exc:
        # Details stay in the server log - SMTP errors can include host/user info.
        logger.exception("SMTP test email failed for %s: %s", inp.to_email, exc)
        raise HTTPException(status_code=500, detail="SMTP test failed, check server logs") from exc

    return TestEmailResponse(sent_to=inp.to_email, subject=inp.subject)


async def _set_password(user_id: str, new_password: str) -> dict:
    """Sets a password the user chose themselves: clears the first-login flag
    and bumps token_version so every older session is logged out."""
    updated = await db.users.find_one_and_update(
        {"id": user_id},
        {
            "$set": {"password_hash": hash_password(new_password), "must_change_password": False},
            "$inc": {"token_version": 1},
        },
        projection={"_id": 0},
        return_document=True,
    )
    return updated


@api_router.post("/auth/change-password", response_model=TokenResponse)
async def change_password(inp: ChangePasswordInput, current_user: User = Depends(get_authenticated_user)):
    raw = await db.users.find_one({"id": current_user.id}, {"_id": 0})
    if not verify_password(inp.current_password, raw["password_hash"]):
        raise HTTPException(status_code=400, detail="Current password is incorrect")
    if inp.new_password == inp.current_password:
        raise HTTPException(status_code=400, detail="New password must be different from the current one")
    ensure_strong_password(inp.new_password)
    updated = await _set_password(current_user.id, inp.new_password)
    # The old token is now invalid (token_version bumped) - hand back a fresh
    # one so the user stays logged in on this device.
    return TokenResponse(access_token=create_access_token(updated), user=await user_out(updated))


@api_router.get("/auth/me", response_model=UserOut)
async def me(current_user: User = Depends(get_authenticated_user)):
    modules = list(ALL_MODULES) if current_user.role == Role.SUPERADMIN else (current_user._modules or [])
    sub = current_user._subscription if (current_user._subscription or {}).get("status") not in (None, "none") else None
    billing = await _billing_notice(current_user.dict()) if sub and sub.get("billing_cycle") else None
    return UserOut(
        **current_user.dict(), modules=modules, subscription=sub,
        announcements=await _announcements_for(current_user.client_id), billing=billing,
    )


@api_router.post("/auth/logout")
async def logout(current_user: User = Depends(get_authenticated_user)):
    return {"ok": True}


@api_router.post("/upload-image")
async def upload_product_image(
    file: UploadFile = File(...),
    kind: str = Query("product", pattern="^(product|client_logo)$"),
    current_user: User = Depends(require_roles(Role.SUPERADMIN, Role.ADMIN)),
):
    if kind == "client_logo" and current_user.role != Role.SUPERADMIN:
        raise HTTPException(status_code=403, detail="Forbidden")
    if (file.content_type or "").lower() not in IMAGE_ALLOWED_TYPES:
        raise HTTPException(status_code=415, detail="Only JPEG, PNG, WEBP or HEIC images are allowed")
    data = await file.read(IMAGE_MAX_BYTES + 1)
    if len(data) > IMAGE_MAX_BYTES:
        raise HTTPException(
            status_code=413, detail=f"Image is too large (max {IMAGE_MAX_BYTES // (1024 * 1024)} MB)"
        )

    if not cloudinary_available():
        raise HTTPException(
            status_code=500,
            detail="Cloudinary is not configured. Set CLOUDINARY_URL or CLOUDINARY_CLOUD_NAME/CLOUDINARY_API_KEY/CLOUDINARY_API_SECRET.",
        )

    try:
        # Upload-time ("incoming") transformation: Cloudinary stores only the
        # shrunk version - max 1000x1000, WEBP, auto quality - typically well
        # under 200 KB, which keeps the free-plan storage from filling up.
        result = await asyncio.to_thread(
            cloudinary.uploader.upload,
            data,
            folder=CLIENT_LOGO_FOLDER if kind == "client_logo" else PRODUCT_IMAGE_FOLDER,
            resource_type="image",
            format="webp",
            transformation=[{"width": 1000, "height": 1000, "crop": "limit"}, {"quality": "auto:good"}],
        )
    except Exception as exc:
        logger.exception("Cloudinary upload failed: %s", exc)
        raise HTTPException(status_code=502, detail="Image upload failed") from exc

    return {
        "url": result.get("secure_url") or result.get("url"),
        "public_id": result.get("public_id"),
    }


@api_router.delete("/upload-image")
async def delete_uploaded_image(
    url: str = Query(..., max_length=2048),
    current_user: User = Depends(require_manager),
):
    """Cleanup for an image uploaded on 'Save' whose product write then failed.
    Only our own folder is touched, and never an image a product still uses."""
    if not _product_image_public_id(url):
        raise HTTPException(status_code=400, detail="Not an image uploaded by this app")
    if await db.products.find_one({"image": url}, {"_id": 1}) or await db.clients.find_one({"logo": url}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Image is in use by a product or client")
    await _delete_product_image(url)
    return {"ok": True}


# ---------------- Clients (tenant companies) ----------------
_INVOICE_PREFIX_RE = re.compile(r"^[A-Z0-9]{1,10}$")


def _invoice_counter_id(client_id: str) -> str:
    return f"invoice:{client_id}"


_CLIENT_COMPUTED = {"invoice_next_seq", "subscription_state", "user_count", "last_activity_at", "last_order_at", "orders_30d"}


async def _client_out(doc: dict) -> Client:
    counter = await db.counters.find_one({"_id": _invoice_counter_id(doc["id"])})
    state = subscription_state(doc.get("subscription"))
    return Client(
        **{
            **doc,
            "invoice_next_seq": (counter or {}).get("seq", 0) + 1,
            "subscription_state": None if state["status"] == "none" else state,
        }
    )


def _validated_invoice_settings(inp: ClientInput) -> dict:
    prefix = (inp.invoice_prefix or "").strip().upper()
    if prefix and not _INVOICE_PREFIX_RE.match(prefix):
        raise HTTPException(status_code=400, detail="Invoice prefix must be 1-10 letters/digits")
    if inp.invoice_numbering not in ("auto", "manual"):
        raise HTTPException(status_code=400, detail="invoice_numbering must be 'auto' or 'manual'")
    return {"invoice_prefix": prefix}


def _validated_alert_days(days: Optional[List[int]]) -> Optional[List[int]]:
    """None = not sent (keep). Otherwise 1-6 distinct day counts, 1..365, ascending."""
    if days is None:
        return None
    unique = sorted(set(days))
    if not 1 <= len(unique) <= 6 or unique[0] < 1 or unique[-1] > 365:
        raise HTTPException(status_code=400, detail="expiry_alert_days: 1-6 different values between 1 and 365")
    return unique


async def _raise_invoice_counter(client_id: str, next_seq: Optional[int]) -> None:
    """Continue a series from an old system: the next number becomes
    `next_seq`. Only ever raises the counter - lowering could reuse a number."""
    if next_seq is None:
        return
    cid = _invoice_counter_id(client_id)
    current = ((await db.counters.find_one({"_id": cid})) or {}).get("seq", 0)
    if next_seq - 1 < current:
        raise HTTPException(status_code=400, detail=f"Next invoice number must be at least {current + 1}")
    await db.counters.update_one({"_id": cid}, {"$set": {"seq": next_seq - 1}}, upsert=True)


@api_router.get("/clients", response_model=List[Client])
async def list_clients(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    docs = await db.clients.find({}, {"_id": 0}).sort("name", 1).to_list(None)
    users = {
        row["_id"]: row
        async for row in db.users.aggregate([
            {"$match": {"client_id": {"$ne": None}}},
            {"$group": {
                "_id": "$client_id",
                "n": {"$sum": {"$cond": [{"$ne": ["$active", False]}, 1, 0]}},
                "seen": {"$max": "$last_seen_at"},
            }},
        ])
    }
    cutoff = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
    orders = {
        row["_id"]: row
        async for row in db.orders.aggregate([
            {"$group": {
                "_id": "$client_id",
                "last": {"$max": "$created_at"},
                "recent": {"$sum": {"$cond": [{"$gte": ["$created_at", cutoff]}, 1, 0]}},
            }},
        ])
    }
    out = []
    for doc in docs:
        client = await _client_out(doc)
        u, o = users.get(doc["id"], {}), orders.get(doc["id"], {})
        client.user_count = u.get("n", 0)
        client.last_order_at = o.get("last")
        client.orders_30d = o.get("recent", 0)
        stamps = [x for x in (u.get("seen"), o.get("last")) if x]
        client.last_activity_at = max(stamps) if stamps else None
        out.append(client)
    return out


@api_router.post("/clients", response_model=ClientCreateResponse)
async def create_client(inp: ClientCreateInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    if await find_user_by_email(inp.admin_email):
        raise HTTPException(status_code=400, detail="Email already in use")
    settings = _validated_invoice_settings(inp)
    alert_days = _validated_alert_days(inp.expiry_alert_days)
    if alert_days:
        settings["expiry_alert_days"] = alert_days
    plan = None
    if inp.plan_id:
        plan = await _get_plan_or_400(inp.plan_id)
        settings["modules"] = list(plan["modules"])
        settings["subscription"] = _new_subscription(plan, inp.subscription_ends_at or _today().isoformat(), None, "manual")
    else:
        settings["modules"] = validated_modules(inp.modules if inp.modules is not None else list(Module))

    client_obj = Client(
        **{
            **inp.dict(
                exclude={
                    "admin_name", "admin_email", "invoice_next_seq", "expiry_alert_days", "modules",
                    "plan_id", "subscription_ends_at", "billing_cycle",
                }
            ),
            **settings,
        }
    )
    await db.clients.insert_one(
        {**client_obj.dict(exclude=_CLIENT_COMPUTED), "_id": client_obj.id}
    )
    await _raise_invoice_counter(client_obj.id, inp.invoice_next_seq)
    await _audit(current_user, "client.create", "client", client_obj.id, client_obj.name, client_obj.id)

    if plan:
        await _log_subscription_event(client_obj.id, "assigned", current_user, plan_name=plan["name"], ends_at=settings["subscription"]["ends_at"])
        await _enroll_calendar(client_obj.id, inp.billing_cycle, current_user, from_today=False)
    admin_obj = User(
        email=normalize_email(inp.admin_email), name=inp.admin_name, role=Role.ADMIN, client_id=client_obj.id
    )
    invite = await _invite_user(admin_obj)
    return ClientCreateResponse(client=await _client_out(client_obj.dict()), admin_user=admin_obj, **invite.dict())


@api_router.get("/clients/me", response_model=Client)
async def get_my_client(current_user: User = Depends(get_current_user)):
    """The caller's own tenant (invoice header data) - read-only for admin,
    warehouse and operator."""
    if not current_user.client_id:
        raise HTTPException(status_code=404, detail="No client")
    doc = await db.clients.find_one({"id": current_user.client_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Client not found")
    return await _client_out(doc)


@api_router.get("/clients/{client_id}", response_model=Client)
async def get_client(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    doc = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Client not found")
    return await _client_out(doc)


@api_router.put("/clients/{client_id}", response_model=Client)
async def update_client(client_id: str, inp: ClientInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    existing = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Client not found")
    settings = _validated_invoice_settings(inp)
    await _raise_invoice_counter(client_id, inp.invoice_next_seq)
    alert_days = _validated_alert_days(inp.expiry_alert_days)
    if alert_days:
        settings["expiry_alert_days"] = alert_days
    if inp.modules is not None:
        if (existing.get("subscription") or {}).get("plan_id"):
            raise HTTPException(status_code=400, detail="Modules come from the client's plan - change the plan instead")
        settings["modules"] = validated_modules(inp.modules)
    updated = {**existing, **inp.dict(exclude={"invoice_next_seq", "expiry_alert_days", "modules"}), **settings}
    await db.clients.replace_one({"id": client_id}, {**updated, "_id": client_id})
    await _audit(current_user, "client.update", "client", client_id, updated.get("name"), client_id)
    if existing.get("logo") != updated.get("logo"):
        await _delete_product_image(existing.get("logo"))
    return await _client_out(updated)


@api_router.delete("/clients/{client_id}")
async def delete_client(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    existing = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Client not found")
    # Soft delete: hard-deleting would orphan this client's users/products/
    # customers/orders and break historical stats.
    await db.clients.update_one({"id": client_id}, {"$set": {"active": False}})
    await _audit(current_user, "client.deactivate", "client", client_id, existing.get("name"), client_id)
    return {"ok": True}


@api_router.post("/clients/{client_id}/activate")
async def activate_client(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    result = await db.clients.update_one({"id": client_id}, {"$set": {"active": True}})
    if result.matched_count == 0:
        raise HTTPException(status_code=404, detail="Client not found")
    named = await db.clients.find_one({"id": client_id}, {"_id": 0, "name": 1})
    await _audit(current_user, "client.activate", "client", client_id, (named or {}).get("name"), client_id)
    return {"ok": True}


# ---------------- Plans & subscriptions (PLAN_PRETPLATE.md) ----------------
class Plan(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    description: Optional[str] = ""
    modules: List[str] = Field(default_factory=list)
    active: bool = True  # inactive plans can't be assigned any more
    # List price per period: months ("1", "3", "6", "12", ...) -> amount. Optional.
    prices: Dict[str, float] = Field(default_factory=dict)
    currency: str = "RSD"
    # Accounts the package includes per role (admin/warehouse/operator); an empty map
    # or a missing role = no limit. Price per extra account per month.
    included_seats: Dict[str, int] = Field(default_factory=dict)
    seat_prices: Dict[str, float] = Field(default_factory=dict)
    created_at: str = Field(default_factory=now_iso)


class PlanInput(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    description: Optional[str] = Field(default="", max_length=500)
    modules: List[Module] = Field(default_factory=list)
    active: bool = True
    # months -> amount; None = keep the stored prices (update) / no prices (create).
    prices: Optional[Dict[str, float]] = None
    # role -> count / monthly price; None = keep (update) / none (create).
    included_seats: Optional[Dict[str, int]] = None
    seat_prices: Optional[Dict[str, float]] = None
    # On update: also rewrite the modules of every client currently on this plan.
    apply_to_clients: bool = False


class SubscriptionAssignInput(BaseModel):
    plan_id: str
    ends_at: Optional[str] = None  # YYYY-MM-DD; moved to the end of its month (calendar billing); None = this month
    billing_cycle: Literal["monthly", "annual"] = "monthly"  # for a client that had no plan yet
    note: Optional[str] = Field(default=None, max_length=500)


class SubscriptionExtendInput(BaseModel):
    months: Optional[int] = Field(default=None, ge=1, le=60)
    ends_at: Optional[str] = None
    note: Optional[str] = Field(default=None, max_length=500)


class SubscriptionNoteInput(BaseModel):
    note: Optional[str] = Field(default=None, max_length=500)


class PurgePauseInput(BaseModel):
    paused: bool
    note: Optional[str] = Field(default=None, max_length=500)


class SubscriptionRow(BaseModel):
    client_id: str
    client_name: str
    client_active: bool
    plan_id: Optional[str] = None
    plan_name: Optional[str] = None
    modules: List[str] = Field(default_factory=list)
    note: Optional[str] = None
    purge_paused: bool = False
    purged_at: Optional[str] = None
    state: Dict[str, Any] = Field(default_factory=lambda: {"status": "none"})


class SubscriptionEvent(BaseModel):
    id: str
    client_id: str
    type: str  # assigned | extended | canceled | purge_paused | purge_resumed | purged | payment_received | charge_created | payment_canceled
    plan_name: Optional[str] = None
    ends_at: Optional[str] = None
    note: Optional[str] = None
    actor_name: Optional[str] = None
    source: str = "manual"  # later e.g. a billing provider
    data: Optional[Dict[str, Any]] = None
    created_at: str


class SubscriptionDetail(BaseModel):
    row: SubscriptionRow
    events: List[SubscriptionEvent]


def _valid_day(value: Optional[str], name: str) -> str:
    try:
        parsed = datetime.strptime((value or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{name} must be YYYY-MM-DD")
    if parsed.year < 2000 or parsed.year > _today().year + 20:
        raise HTTPException(status_code=400, detail=f"{name} is out of range")
    return parsed.isoformat()


def _validated_prices(prices: Optional[Dict[str, float]]) -> Dict[str, float]:
    out: Dict[str, float] = {}
    for key, amount in (prices or {}).items():
        if not str(key).isdigit() or not 1 <= int(key) <= 60:
            raise HTTPException(status_code=400, detail="Price period must be a number of months (1-60)")
        if not 0 < float(amount) <= 1_000_000_000:
            raise HTTPException(status_code=400, detail="Price must be greater than 0")
        out[str(int(key))] = round(float(amount), 2)
    return dict(sorted(out.items(), key=lambda kv: int(kv[0])))


SEAT_ROLES = ("admin", "warehouse", "operator")


def _validated_seat_map(values: Optional[Dict[str, float]], name: str, integer: bool) -> dict:
    out: dict = {}
    for role, amount in (values or {}).items():
        if role not in SEAT_ROLES:
            raise HTTPException(status_code=400, detail=f"{name}: role must be one of {', '.join(SEAT_ROLES)}")
        if float(amount) < 0 or float(amount) > (10_000 if integer else 1_000_000_000):
            raise HTTPException(status_code=400, detail=f"{name}: invalid value for {role}")
        out[role] = int(amount) if integer else round(float(amount), 2)
    return out


def _seat_figures(plan: Optional[dict], sub: Optional[dict]) -> dict:
    """Per role: included (plan) + extra (contract) = limit; None = unlimited."""
    included = (plan or {}).get("included_seats") or {}
    extra = (sub or {}).get("extra_seats") or {}
    out = {}
    for role in SEAT_ROLES:
        inc = included.get(role)
        ext = int(extra.get(role, 0) or 0)
        out[role] = {
            "included": inc, "extra": ext,
            "limit": None if inc is None else int(inc) + ext,
            "price": ((plan or {}).get("seat_prices") or {}).get(role),
        }
    return out


async def _check_seat_limit(client_id: Optional[str], role, exclude_user_id: Optional[str] = None) -> None:
    """403 when one more active account of this role would exceed the client's limit."""
    role_v = getattr(role, "value", role)
    if not client_id or role_v not in SEAT_ROLES:
        return
    client = await db.clients.find_one({"id": client_id}, {"_id": 0, "subscription": 1})
    sub = (client or {}).get("subscription")
    if not sub or not sub.get("plan_id"):
        return
    plan = await db.plans.find_one({"id": sub["plan_id"]}, {"_id": 0})
    limit = _seat_figures(plan, sub)[role_v]["limit"]
    if limit is None:
        return
    query: dict = {"client_id": client_id, "role": role_v, "active": {"$ne": False}}
    if exclude_user_id:
        query["id"] = {"$ne": exclude_user_id}
    if await db.users.count_documents(query) >= limit:
        raise HTTPException(status_code=403, detail="Seat limit reached for this role")


async def _get_plan_or_400(plan_id: str, must_be_active: bool = True) -> dict:
    plan = await db.plans.find_one({"id": plan_id}, {"_id": 0})
    if not plan:
        raise HTTPException(status_code=400, detail="Plan not found")
    if must_be_active and not plan.get("active", True):
        raise HTTPException(status_code=400, detail="Plan is not active")
    return plan


def _new_subscription(plan: dict, ends_at: Optional[str], note: Optional[str], source: str) -> dict:
    return {
        "plan_id": plan["id"],
        "plan_name": plan["name"],
        "starts_at": _today().isoformat(),
        "ends_at": _valid_day(ends_at, "ends_at"),
        "note": note,
        "source": source,
        "canceled_at": None,
        "purge_paused": False,
        "purged_at": None,
    }


async def _audit(
    user: Optional[User], action: str, target_type: Optional[str] = None, target_id: Optional[str] = None,
    target_name: Optional[str] = None, client_id: Optional[str] = None, data: Optional[dict] = None,
) -> None:
    """Superadmin action log (who did what to whom). Nothing is logged for the system itself."""
    if user is None:
        return
    entry_id = str(uuid.uuid4())
    await db.audit_log.insert_one({
        "_id": entry_id, "id": entry_id, "created_at": now_iso(), "actor_id": user.id, "actor_name": user.name,
        "action": action, "target_type": target_type, "target_id": target_id, "target_name": target_name,
        "client_id": client_id, "data": data,
    })


async def _log_subscription_event(
    client_id: str, type_: str, user: Optional[User], note: Optional[str] = None,
    plan_name: Optional[str] = None, ends_at: Optional[str] = None, data: Optional[dict] = None,
    source: str = "manual",
) -> None:
    if user is not None:
        client = await db.clients.find_one({"id": client_id}, {"_id": 0, "name": 1})
        detail = {"plan": plan_name, "ends_at": ends_at, "note": note, **(data or {})}
        await _audit(
            user, f"subscription.{type_}", "client", client_id, (client or {}).get("name"), client_id,
            {k: v for k, v in detail.items() if v is not None},
        )
    await db.subscription_events.insert_one({
        "_id": str(uuid.uuid4()), "id": str(uuid.uuid4()), "client_id": client_id, "type": type_,
        "plan_name": plan_name, "ends_at": ends_at, "note": note,
        "actor_name": user.name if user else "system", "source": source, "data": data, "created_at": now_iso(),
    })


async def _subscription_row(client: dict) -> SubscriptionRow:
    sub = client.get("subscription") or {}
    return SubscriptionRow(
        client_id=client["id"],
        client_name=client["name"],
        client_active=client.get("active", True),
        plan_id=sub.get("plan_id"),
        plan_name=sub.get("plan_name"),
        modules=effective_modules(client),
        note=sub.get("note"),
        purge_paused=bool(sub.get("purge_paused")),
        purged_at=sub.get("purged_at"),
        state=subscription_state(client.get("subscription")),
    )


async def _client_or_404(client_id: str) -> dict:
    client = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")
    return client


@api_router.get("/plans", response_model=List[Plan])
async def list_plans(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    return [Plan(**p) for p in await db.plans.find({}, {"_id": 0}).sort("name", 1).to_list(None)]


@api_router.post("/plans", response_model=Plan)
async def create_plan(inp: PlanInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    if await db.plans.find_one({"name": inp.name.strip()}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="A plan with that name already exists")
    plan = Plan(
        name=inp.name.strip(), description=inp.description or "", modules=validated_modules(inp.modules), active=inp.active,
        prices=_validated_prices(inp.prices),
        included_seats=_validated_seat_map(inp.included_seats, "included_seats", integer=True),
        seat_prices=_validated_seat_map(inp.seat_prices, "seat_prices", integer=False),
    )
    await db.plans.insert_one({**plan.dict(), "_id": plan.id})
    await _audit(current_user, "plan.create", "plan", plan.id, plan.name)
    return plan


@api_router.put("/plans/{plan_id}", response_model=Plan)
async def update_plan(plan_id: str, inp: PlanInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    existing = await db.plans.find_one({"id": plan_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Plan not found")
    clash = await db.plans.find_one({"name": inp.name.strip(), "id": {"$ne": plan_id}}, {"_id": 1})
    if clash:
        raise HTTPException(status_code=409, detail="A plan with that name already exists")
    updated = {
        **existing, "name": inp.name.strip(), "description": inp.description or "",
        "modules": validated_modules(inp.modules), "active": inp.active,
    }
    if inp.prices is not None:
        updated["prices"] = _validated_prices(inp.prices)
    if inp.included_seats is not None:
        updated["included_seats"] = _validated_seat_map(inp.included_seats, "included_seats", integer=True)
    if inp.seat_prices is not None:
        updated["seat_prices"] = _validated_seat_map(inp.seat_prices, "seat_prices", integer=False)
    await db.plans.replace_one({"id": plan_id}, {**updated, "_id": plan_id})
    await _audit(current_user, "plan.update", "plan", plan_id, updated["name"], data={"apply_to_clients": inp.apply_to_clients})
    # Clients keep their own copy of the plan name; refresh it, and optionally the modules.
    await db.clients.update_many({"subscription.plan_id": plan_id}, {"$set": {"subscription.plan_name": updated["name"]}})
    if inp.apply_to_clients:
        async for client in db.clients.find({"subscription.plan_id": plan_id}, {"_id": 0, "id": 1}):
            await db.clients.update_one({"id": client["id"]}, {"$set": {"modules": updated["modules"]}})
    return Plan(**updated)


@api_router.delete("/plans/{plan_id}")
async def delete_plan(plan_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    if not await db.plans.find_one({"id": plan_id}, {"_id": 1}):
        raise HTTPException(status_code=404, detail="Plan not found")
    if await db.clients.find_one({"subscription.plan_id": plan_id}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="Clients are on this plan - deactivate it instead")
    named = await db.plans.find_one({"id": plan_id}, {"_id": 0, "name": 1})
    await db.plans.delete_one({"id": plan_id})
    await _audit(current_user, "plan.delete", "plan", plan_id, (named or {}).get("name"))
    return {"ok": True}


@api_router.get("/subscriptions", response_model=List[SubscriptionRow])
async def list_subscriptions(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    clients = await db.clients.find({}, {"_id": 0}).sort("name", 1).to_list(None)
    return [await _subscription_row(c) for c in clients]


@api_router.get("/subscriptions/{client_id}", response_model=SubscriptionDetail)
async def get_subscription(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    client = await _client_or_404(client_id)
    events = await db.subscription_events.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(200)
    return SubscriptionDetail(row=await _subscription_row(client), events=[SubscriptionEvent(**e) for e in events])


@api_router.post("/subscriptions/{client_id}/assign", response_model=SubscriptionRow)
async def assign_subscription(
    client_id: str, inp: SubscriptionAssignInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """Gives the client a package until `ends_at` (replaces any earlier subscription,
    unlocks a locked client) and sets its modules to the plan's."""
    client = await _client_or_404(client_id)
    plan = await _get_plan_or_400(inp.plan_id)
    subscription = _new_subscription(plan, inp.ends_at or _today().isoformat(), inp.note, "manual")
    for kept in ("extra_seats", "discount_percent", *_BILLING_KEPT_FIELDS):
        if (client.get("subscription") or {}).get(kept) is not None:
            subscription[kept] = client["subscription"][kept]
    await db.clients.update_one({"id": client_id}, {"$set": {"subscription": subscription, "modules": plan["modules"]}})
    await _log_subscription_event(
        client_id, "assigned", current_user, note=inp.note, plan_name=plan["name"], ends_at=subscription["ends_at"]
    )
    # Every client with a plan is on calendar billing: the end date goes to the end of
    # its month (a date in the past stays in the past - an assigned end date is kept).
    if subscription.get("billing_cycle"):
        aligned = _month_end(date.fromisoformat(subscription["ends_at"])).isoformat()
        await db.clients.update_one({"id": client_id}, {"$set": {"subscription.ends_at": aligned}})
        await _apply_ledger(client_id, current_user)
    else:
        await _enroll_calendar(client_id, inp.billing_cycle, current_user, from_today=False)
    return await _subscription_row(await _client_or_404(client_id))


class SeatsInput(BaseModel):
    # Contracted accounts on top of the package, per role. None = keep.
    extra_seats: Optional[Dict[str, int]] = None
    # Percent off the monthly total (package + extra accounts). None = keep.
    discount_percent: Optional[float] = Field(default=None, ge=0, le=100)


class SeatRow(BaseModel):
    role: str
    used: int
    included: Optional[int] = None
    extra: int = 0
    limit: Optional[int] = None  # None = unlimited
    price: Optional[float] = None  # monthly price of one extra account


class SeatsOut(BaseModel):
    plan_name: Optional[str] = None
    currency: str = "RSD"
    seats: List[SeatRow]
    package_price: Optional[float] = None  # monthly list price of the package
    extra_amount: float = 0
    discount_percent: float = 0
    monthly_total: Optional[float] = None


async def _seats_out(client_id: str, include_prices: bool = True) -> SeatsOut:
    client = await _client_or_404(client_id)
    sub = client.get("subscription")
    plan = await db.plans.find_one({"id": sub["plan_id"]}, {"_id": 0}) if sub and sub.get("plan_id") else None
    figures = _seat_figures(plan, sub)
    modules = effective_modules(client)
    rows, extra_amount = [], 0.0
    for role in SEAT_ROLES:
        if role == "warehouse" and Module.WAREHOUSE.value not in modules:
            continue
        used = await db.users.count_documents({"client_id": client_id, "role": role, "active": {"$ne": False}})
        f = figures[role]
        rows.append(SeatRow(role=role, used=used, included=f["included"], extra=f["extra"], limit=f["limit"], price=f["price"]))
        extra_amount += f["extra"] * (f["price"] or 0)
    package_price = ((plan or {}).get("prices") or {}).get("1")
    discount = float((sub or {}).get("discount_percent") or 0)
    total = None
    if plan and (package_price is not None or extra_amount):
        total = round(((package_price or 0) + extra_amount) * (1 - discount / 100), 2)
    return SeatsOut(
        plan_name=(sub or {}).get("plan_name"), currency=(plan or {}).get("currency", "RSD"), seats=rows,
        package_price=package_price, extra_amount=round(extra_amount, 2), discount_percent=discount, monthly_total=total,
    )


@api_router.get("/clients/me/seats", response_model=SeatsOut)
async def my_seats(current_user: User = Depends(require_roles(Role.ADMIN))):
    """The admin's own client: accounts used / allowed and the monthly amount."""
    return await _seats_out(current_user.client_id)


@api_router.get("/clients/{client_id}/seats", response_model=SeatsOut)
async def client_seats(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    return await _seats_out(client_id)


@api_router.put("/subscriptions/{client_id}/seats", response_model=SeatsOut)
async def set_client_seats(client_id: str, inp: SeatsInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Sets the contracted extra accounts / discount. Takes effect at once (limit);
    billed with the next charge. Lowering below the accounts in use is allowed -
    existing accounts keep working, new ones are refused."""
    client = await _client_or_404(client_id)
    sub = client.get("subscription")
    if not sub or not sub.get("plan_id"):
        raise HTTPException(status_code=400, detail="Client has no subscription - assign a plan first")
    changes = {}
    if inp.extra_seats is not None:
        changes["subscription.extra_seats"] = _validated_seat_map(inp.extra_seats, "extra_seats", integer=True)
    if inp.discount_percent is not None:
        changes["subscription.discount_percent"] = round(inp.discount_percent, 2)
    if changes:
        await db.clients.update_one({"id": client_id}, {"$set": changes})
        await _log_subscription_event(
            client_id, "seats", current_user, data={k.split(".")[1]: v for k, v in changes.items()}
        )
        await _apply_ledger(client_id, current_user)
    return await _seats_out(client_id)


async def _extend_subscription(
    client: dict, months: Optional[int], ends_at: Optional[str], note: Optional[str], user: User
) -> str:
    """Moves `ends_at` forward by whole months (from the later of today and the
    current end) or to an exact date; lifts a lock/cancel. Returns the new end."""
    sub = client["subscription"]
    if months is not None:
        new_end = _add_months(max(date.fromisoformat(sub["ends_at"]), _today()), months).isoformat()
    else:
        new_end = _valid_day(ends_at, "ends_at")
    sub = {**sub, "ends_at": new_end, "canceled_at": None, "purge_paused": False, "reminder_for": None, "purge_warned_for": None}
    if note:
        sub["note"] = note
    await db.clients.update_one({"id": client["id"]}, {"$set": {"subscription": sub}})
    await _log_subscription_event(client["id"], "extended", user, note=note, plan_name=sub.get("plan_name"), ends_at=new_end)
    return new_end


@api_router.post("/subscriptions/{client_id}/extend", response_model=SubscriptionRow)
async def extend_subscription(
    client_id: str, inp: SubscriptionExtendInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """Pushes `ends_at` forward by whole months (from the later of today and the
    current end) or to an exact date. Also lifts a lock and a cancel."""
    client = await _client_or_404(client_id)
    if not client.get("subscription"):
        raise HTTPException(status_code=400, detail="Client has no subscription - assign a plan first")
    if (inp.months is None) == (inp.ends_at is None):
        raise HTTPException(status_code=400, detail="Send either months or ends_at")
    await _extend_subscription(client, inp.months, inp.ends_at, inp.note, current_user)
    return await _subscription_row(await _client_or_404(client_id))


@api_router.post("/subscriptions/{client_id}/cancel", response_model=SubscriptionRow)
async def cancel_subscription(
    client_id: str, inp: SubscriptionNoteInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """Locks the client right away (no grace period); data stays until the purge."""
    client = await _client_or_404(client_id)
    sub = client.get("subscription")
    if not sub:
        raise HTTPException(status_code=400, detail="Client has no subscription")
    sub = {**sub, "canceled_at": now_iso()}
    await db.clients.update_one({"id": client_id}, {"$set": {"subscription": sub}})
    await _log_subscription_event(client_id, "canceled", current_user, note=inp.note, plan_name=sub.get("plan_name"), ends_at=sub["ends_at"])
    return await _subscription_row(await _client_or_404(client_id))


@api_router.post("/subscriptions/{client_id}/purge-pause", response_model=SubscriptionRow)
async def pause_purge(
    client_id: str, inp: PurgePauseInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """Stops (or resumes) the scheduled deletion of this client's module data."""
    client = await _client_or_404(client_id)
    sub = client.get("subscription")
    if not sub:
        raise HTTPException(status_code=400, detail="Client has no subscription")
    await db.clients.update_one({"id": client_id}, {"$set": {"subscription.purge_paused": inp.paused}})
    await _log_subscription_event(
        client_id, "purge_paused" if inp.paused else "purge_resumed", current_user, note=inp.note, plan_name=sub.get("plan_name")
    )
    return await _subscription_row(await _client_or_404(client_id))


@api_router.get("/subscriptions/{client_id}/export")
async def export_purgeable_data(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Everything the purge would delete, as JSON - to hand over before it runs."""
    client = await _client_or_404(client_id)
    products = await db.products.find({"client_id": client_id}, {"_id": 0}).to_list(None)
    return {
        "exported_at": now_iso(),
        "client": {"id": client_id, "name": client["name"]},
        "products_stock": [
            {"id": p["id"], "name": p["name"], "barcode": p.get("barcode"), "stock_qty": p.get("stock_qty"),
             "track_expiry": bool(p.get("track_expiry"))}
            for p in products
        ],
        "stock_batches": await db.stock_batches.find({"client_id": client_id}, {"_id": 0}).to_list(None),
        "stock_movements": await db.stock_movements.find({"client_id": client_id}, {"_id": 0}).to_list(None),
    }


async def _purge_module_data(client_id: str) -> Dict[str, int]:
    """What the scheduled purge deletes for a client that stayed locked: the data
    of the optional modules (stock levels, movements, batches). Orders,
    customers, products and users are never touched."""
    movements = await db.stock_movements.delete_many({"client_id": client_id})
    batches = await db.stock_batches.delete_many({"client_id": client_id})
    products = await db.products.update_many(
        {"client_id": client_id}, {"$unset": {"stock_qty": ""}, "$set": {"track_expiry": False}}
    )
    return {
        "stock_movements": movements.deleted_count,
        "stock_batches": batches.deleted_count,
        "products_reset": products.modified_count,
    }


async def _count_purgeable(client_id: str) -> Dict[str, int]:
    return {
        "stock_movements": await db.stock_movements.count_documents({"client_id": client_id}),
        "stock_batches": await db.stock_batches.count_documents({"client_id": client_id}),
        "products_reset": await db.products.count_documents({"client_id": client_id, "stock_qty": {"$exists": True}}),
    }


async def _email_client_admins(client_id: str, subject: str, body: str) -> int:
    """Best effort: skipped (0) when SMTP isn't configured."""
    if not _smtp_is_configured():
        return 0
    sent = 0
    async for admin in db.users.find({"client_id": client_id, "role": Role.ADMIN.value, "active": {"$ne": False}}, {"email": 1}):
        try:
            await asyncio.to_thread(_send_smtp_email_sync, admin["email"], subject, body)
            sent += 1
        except Exception as exc:
            logger.warning("Subscription email to %s failed: %s", admin.get("email"), exc)
    return sent


@api_router.get("/internal/cron/subscriptions")
async def cron_subscriptions(
    dry_run: bool = Query(False), as_of: Optional[str] = Query(None), authorization: Optional[str] = Header(None),
):
    """Daily job (Vercel Cron sends `Authorization: Bearer $CRON_SECRET`):
    expiry reminders, purge warnings and the purge itself, calendar billing drafts
    and billing emails. The purge deletes only when AUTO_PURGE_ENABLED=true and
    dry_run is not set; otherwise it reports what it would delete. `as_of` (not in
    production) runs the billing emails as if it were that day - for tests."""
    if not CRON_SECRET:
        raise HTTPException(status_code=503, detail="CRON_SECRET is not configured")
    if not hmac.compare_digest(authorization or "", f"Bearer {CRON_SECRET}"):
        raise HTTPException(status_code=401, detail="Unauthorized")
    await refresh_settings(force=True)
    live = AUTO_PURGE_ENABLED and not dry_run
    today = _today()
    notify_day = today
    if as_of:
        if IS_PRODUCTION:
            raise HTTPException(status_code=400, detail="as_of is not allowed in production")
        notify_day = date.fromisoformat(_valid_day(as_of, "as_of"))
    result: Dict[str, Any] = {"live": live, "reminders": [], "purge_warnings": [], "purged": [], "would_purge": []}
    async for client in db.clients.find({"subscription": {"$ne": None}}, {"_id": 0}):
        sub, cid = client["subscription"], client["id"]
        state = subscription_state(sub, today)
        # Monthly calendar billing gets its charge every month; the end-date reminder is for the rest.
        if sub.get("billing_cycle") != "monthly" and state["status"] == "active" and state["days_left"] <= setting("reminder_days") and sub.get("reminder_for") != sub["ends_at"]:
            if sub.get("billing_cycle") == "annual":
                await _notify_annual_expiry(client, state)
            else:
                await _email_client_admins(
                    cid, "Easy Order - pretplata uskoro ističe / subscription ends soon",
                    f"Pretplata za {client['name']} ističe {sub['ends_at']} (za {state['days_left']} dana).\n\n"
                    f"Your subscription for {client['name']} ends on {sub['ends_at']}.",
                )
            await db.clients.update_one({"id": cid}, {"$set": {"subscription.reminder_for": sub["ends_at"]}})
            result["reminders"].append(client["name"])
        if state["status"] != "locked" or not state.get("purge_at") or sub.get("purged_at") or sub.get("purge_paused"):
            continue
        purge_at = date.fromisoformat(state["purge_at"])
        if purge_at > today:
            if (purge_at - today).days <= setting("reminder_days") and sub.get("purge_warned_for") != state["purge_at"]:
                await _email_client_admins(
                    cid, "Easy Order - podaci će biti obrisani / data will be deleted",
                    f"Nalog {client['name']} je zaključan. Podaci o zalihama biće obrisani {state['purge_at']}. "
                    f"Obratite se podršci ako želite da produžite pretplatu.\n\n"
                    f"The account {client['name']} is locked. Stock data will be deleted on {state['purge_at']}.",
                )
                await db.clients.update_one({"id": cid}, {"$set": {"subscription.purge_warned_for": state["purge_at"]}})
                result["purge_warnings"].append(client["name"])
            continue
        if live:
            counts = await _purge_module_data(cid)
            await db.clients.update_one({"id": cid}, {"$set": {"subscription.purged_at": now_iso()}})
            await _log_subscription_event(cid, "purged", None, plan_name=sub.get("plan_name"), data=counts, source="cron")
            result["purged"].append({"client": client["name"], **counts})
        else:
            counts = await _count_purgeable(cid)
            result["would_purge"].append({"client": client["name"], **counts})
    # Calendar billing: drafts for next month from the draft day on (one per client and period).
    result["drafts"] = []
    if today.day >= int(setting("draft_day")):
        made = await _make_drafts(today)
        result["drafts"] = [d["client_id"] for d in made]
        await _notify_drafts_ready(made)
    result["billing_emails"] = await _billing_notifications(notify_day)
    await db.system_runs.replace_one(
        {"_id": "subscriptions_cron"},
        {
            "_id": "subscriptions_cron", "finished_at": now_iso(), "live": live,
            "summary": {k: len(result[k]) for k in ("reminders", "purge_warnings", "purged", "would_purge", "drafts")},
        },
        upsert=True,
    )
    return result


# ---------------- Payments & debts (PLAN_SUPERADMIN.md, phase C) ----------------
# Kept by hand by the superadmin. `expected` = a debt (due_date), `received` =
# money in (paid_at), `canceled` stays on record but is never counted.
PAYMENT_CURRENCY = "RSD"


class Payment(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_id: str
    client_name: Optional[str] = None  # filled in lists
    status: Literal["draft", "expected", "received", "canceled"]
    # Calendar billing (ledger) entries: charges stay `expected`; `open` is what is
    # still unpaid after the client's payments were applied oldest first.
    ledger: bool = False
    number: Optional[str] = None
    open: Optional[float] = None
    amount: float
    currency: str = PAYMENT_CURRENCY
    due_date: Optional[str] = None
    paid_at: Optional[str] = None
    method: Optional[Literal["bank", "card", "cash", "other"]] = None
    note: Optional[str] = None
    plan_name: Optional[str] = None
    period_months: Optional[int] = None
    # subscription = package (+ extra accounts) for a period, setup = one-time integration, other.
    kind: Literal["subscription", "setup", "other"] = "subscription"
    period: Optional[str] = None  # YYYY-MM of a monthly charge
    breakdown: Optional[Dict[str, Any]] = None  # package / extra accounts / discount of a monthly charge
    created_by: Optional[str] = None
    source: str = "manual"
    created_at: str = Field(default_factory=now_iso)
    canceled_at: Optional[str] = None
    overdue: bool = False  # computed: expected and past its due date


class PaymentInput(BaseModel):
    client_id: str
    status: Literal["received", "expected"] = "received"
    amount: float = Field(gt=0, le=1_000_000_000)
    paid_at: Optional[str] = None  # received: defaults to today
    due_date: Optional[str] = None  # expected: required
    method: Optional[Literal["bank", "card", "cash", "other"]] = None
    note: Optional[str] = Field(default=None, max_length=500)
    period_months: Optional[int] = Field(default=None, ge=1, le=60)
    plan_name: Optional[str] = Field(default=None, max_length=80)
    kind: Literal["subscription", "setup", "other"] = "subscription"


class PaymentReceiveInput(BaseModel):
    paid_at: Optional[str] = None
    amount: Optional[float] = Field(default=None, gt=0, le=1_000_000_000)
    method: Optional[Literal["bank", "card", "cash", "other"]] = None
    note: Optional[str] = Field(default=None, max_length=500)


class PaymentTotals(BaseModel):
    received: float = 0
    expected: float = 0
    overdue_total: float = 0
    overdue_count: int = 0


class PaymentsList(BaseModel):
    items: List[Payment]
    totals: PaymentTotals


def _valid_month(value: str) -> str:
    try:
        parsed = datetime.strptime(value, "%Y-%m")
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="month must be YYYY-MM")
    if parsed.year < 2000 or parsed.year > _today().year + 5:
        raise HTTPException(status_code=400, detail="month is out of range")
    return parsed.strftime("%Y-%m")


def _payment_out(doc: dict, names: Optional[Dict[str, str]] = None, today: Optional[date] = None) -> Payment:
    today = today or _today()
    still_owed = doc.get("open", doc.get("amount", 0)) > 0 if doc.get("ledger") else True
    overdue = doc.get("status") == "expected" and still_owed and bool(doc.get("due_date")) and date.fromisoformat(doc["due_date"]) < today
    return Payment(**{**doc, "client_name": (names or {}).get(doc["client_id"]), "overdue": overdue})


def _paid_day(value: Optional[str]) -> str:
    """Day the money arrived: today by default, never in the future."""
    if not value:
        return _today().isoformat()
    day = _valid_day(value, "paid_at")
    if date.fromisoformat(day) > _today():
        raise HTTPException(status_code=400, detail="paid_at cannot be in the future")
    return day


async def _payment_or_404(payment_id: str) -> dict:
    doc = await db.payments.find_one({"id": payment_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Payment not found")
    return doc


async def _client_names() -> Dict[str, str]:
    return {c["id"]: c["name"] async for c in db.clients.find({}, {"_id": 0, "id": 1, "name": 1})}


@api_router.get("/payments", response_model=PaymentsList)
async def list_payments(
    client_id: Optional[str] = None,
    status: Optional[Literal["expected", "received", "canceled"]] = None,
    from_date: Optional[str] = Query(None, description="YYYY-MM-DD, by paid_at (received) or due_date (expected)"),
    to_date: Optional[str] = Query(None, description="YYYY-MM-DD, inclusive"),
    current_user: User = Depends(require_roles(Role.SUPERADMIN)),
):
    query: Dict[str, Any] = {}
    if client_id:
        query["client_id"] = client_id
    if status:
        query["status"] = status
    lo = _valid_day(from_date, "from_date") if from_date else None
    hi = _valid_day(to_date, "to_date") if to_date else None
    names, today = await _client_names(), _today()
    items = []
    docs = [d for d in await db.payments.find(query, {"_id": 0}).to_list(None) if d.get("status") != "draft"]
    open_by_id = await _ledger_open_map({d["client_id"] for d in docs if d.get("ledger")})
    for doc in docs:
        if doc["id"] in open_by_id:
            doc = {**doc, "open": open_by_id[doc["id"]]}
        day = doc.get("paid_at") or doc.get("due_date") or doc["created_at"][:10]
        if (lo and day < lo) or (hi and day > hi):
            continue
        items.append((day, _payment_out(doc, names, today)))
    items.sort(key=lambda pair: (pair[0], pair[1].created_at), reverse=True)
    out = [p for _, p in items]
    totals = PaymentTotals()
    for p in out:
        if p.status == "received":
            totals.received += p.amount
        elif p.status == "expected":
            owed = p.open if p.ledger and p.open is not None else p.amount
            totals.expected += owed
            if p.overdue:
                totals.overdue_total += owed
                totals.overdue_count += 1
    totals.received, totals.expected, totals.overdue_total = (
        round(totals.received, 2), round(totals.expected, 2), round(totals.overdue_total, 2)
    )
    return PaymentsList(items=out, totals=totals)


async def _ledger_entry(client: dict, inp: "PaymentInput", user: User) -> "Payment":
    if inp.status == "received":
        doc = {
            "id": str(uuid.uuid4()), "client_id": client["id"], "status": "received", "ledger": True, "kind": inp.kind,
            "amount": round(inp.amount, 2), "currency": PAYMENT_CURRENCY, "paid_at": _paid_day(inp.paid_at),
            "due_date": None, "method": inp.method or "bank", "note": inp.note,
            "plan_name": client["subscription"].get("plan_name"), "created_by": user.name, "source": "manual",
            "created_at": now_iso(), "canceled_at": None,
        }
        event = "payment_received"
    else:
        if not inp.due_date:
            raise HTTPException(status_code=400, detail="due_date is required for a debt")
        label = inp.note or {"setup": "Uvođenje", "other": "Ostalo"}.get(inp.kind, "Pretplata")
        doc = {
            "id": str(uuid.uuid4()), "client_id": client["id"], "status": "expected", "ledger": True, "kind": inp.kind,
            "amount": round(inp.amount, 2), "currency": PAYMENT_CURRENCY, "due_date": _valid_day(inp.due_date, "due_date"),
            "paid_at": None, "method": None, "note": inp.note, "lines": [{"kind": "manual", "label": label, "amount": round(inp.amount, 2)}],
            "number": await _next_proforma_number(_today()), "approved_at": now_iso(), "approved_by": user.name,
            "period_start": None, "period_end": None, "plan_name": client["subscription"].get("plan_name"),
            "created_by": user.name, "source": "manual", "created_at": now_iso(), "canceled_at": None,
        }
        event = "charge_created"
    await db.payments.insert_one({**doc, "_id": doc["id"]})
    await _log_subscription_event(
        client["id"], event, user, note=inp.note, plan_name=doc["plan_name"],
        data={"amount": doc["amount"], "currency": PAYMENT_CURRENCY, "kind": inp.kind},
    )
    await _apply_ledger(client["id"], user)
    if doc["status"] == "expected":
        await _notify_issued(client["id"], [doc])
    return _payment_out(doc, {client["id"]: client["name"]})


@api_router.post("/payments", response_model=Payment)
async def create_payment(inp: PaymentInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Records money received (paid_at) or a debt to be paid (due_date). For a client
    on calendar billing both go into its balance: a debt gets a pro-forma number and
    is shown to the client (e.g. the setup fee), money in is applied oldest first."""
    client = await _client_or_404(inp.client_id)
    if (client.get("subscription") or {}).get("billing_cycle"):
        return await _ledger_entry(client, inp, current_user)
    doc = {
        "id": str(uuid.uuid4()), "client_id": client["id"], "status": inp.status, "amount": round(inp.amount, 2),
        "currency": PAYMENT_CURRENCY, "due_date": None, "paid_at": None, "method": inp.method, "note": inp.note,
        "plan_name": inp.plan_name or (client.get("subscription") or {}).get("plan_name"),
        "period_months": inp.period_months, "kind": inp.kind, "period": None, "breakdown": None,
        "created_by": current_user.name, "source": "manual", "created_at": now_iso(), "canceled_at": None,
    }
    if inp.status == "received":
        doc["paid_at"] = _paid_day(inp.paid_at)
        doc["method"] = inp.method or "bank"
    else:
        if not inp.due_date:
            raise HTTPException(status_code=400, detail="due_date is required for a debt")
        doc["due_date"] = _valid_day(inp.due_date, "due_date")
    await db.payments.insert_one({**doc, "_id": doc["id"]})
    await _log_subscription_event(
        client["id"], "payment_received" if inp.status == "received" else "charge_created", current_user,
        note=inp.note, plan_name=doc["plan_name"], data={"amount": doc["amount"], "currency": PAYMENT_CURRENCY},
    )
    return _payment_out(doc, {client["id"]: client["name"]})


async def _settle_payment(doc: dict, inp: PaymentReceiveInput, user: User, period_months: Optional[int] = None) -> dict:
    if doc.get("ledger"):
        raise HTTPException(status_code=400, detail="Record payments for calendar billing with POST /billing/payments")
    if doc["status"] != "expected":
        raise HTTPException(status_code=409, detail="Only an expected payment can be marked as received")
    update = {
        "status": "received", "paid_at": _paid_day(inp.paid_at), "method": inp.method or doc.get("method") or "bank",
        "amount": round(inp.amount, 2) if inp.amount else doc["amount"], "note": inp.note or doc.get("note"),
    }
    if period_months:
        update["period_months"] = period_months
    await db.payments.update_one({"id": doc["id"]}, {"$set": update})
    await _log_subscription_event(
        doc["client_id"], "payment_received", user, note=update["note"], plan_name=doc.get("plan_name"),
        data={"amount": update["amount"], "currency": PAYMENT_CURRENCY},
    )
    return await _payment_or_404(doc["id"])


@api_router.post("/payments/{payment_id}/receive", response_model=Payment)
async def receive_payment(
    payment_id: str, inp: PaymentReceiveInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """Turns an expected payment (debt) into a received one."""
    doc = await _settle_payment(await _payment_or_404(payment_id), inp, current_user)
    return _payment_out(doc, await _client_names())


@api_router.post("/payments/{payment_id}/cancel", response_model=Payment)
async def cancel_payment(
    payment_id: str, inp: SubscriptionNoteInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    """A wrong entry stays on record as canceled and is not counted."""
    doc = await _payment_or_404(payment_id)
    if doc["status"] == "canceled":
        raise HTTPException(status_code=409, detail="Payment is already canceled")
    note = inp.note or doc.get("note")
    await db.payments.update_one({"id": payment_id}, {"$set": {"status": "canceled", "canceled_at": now_iso(), "note": note}})
    await _log_subscription_event(
        doc["client_id"], "payment_canceled", current_user, note=inp.note, plan_name=doc.get("plan_name"),
        data={"amount": doc["amount"], "currency": PAYMENT_CURRENCY, "was": doc["status"]},
    )
    if doc.get("ledger"):
        await _apply_ledger(doc["client_id"], current_user)
    return _payment_out(await _payment_or_404(payment_id), await _client_names())


async def _payment_figures(today: date) -> Dict[str, Any]:
    """Money figures for the dashboard (canceled entries are never counted)."""
    month, year = today.strftime("%Y-%m"), today.strftime("%Y")
    out: Dict[str, Any] = {
        "received_month": 0.0, "received_year": 0.0, "expected_total": 0.0, "overdue_total": 0.0, "overdue_count": 0,
        "overdue_by_client": {},
    }
    ledger_clients = set()
    async for doc in db.payments.find({"status": {"$in": ["received", "expected"]}}, {"_id": 0}):
        if doc.get("ledger") and doc["status"] == "expected":
            ledger_clients.add(doc["client_id"])
            continue  # counted below by what is still open
        if doc["status"] == "received":
            paid = doc.get("paid_at") or ""
            if paid.startswith(month):
                out["received_month"] += doc["amount"]
            if paid.startswith(year):
                out["received_year"] += doc["amount"]
        else:
            out["expected_total"] += doc["amount"]
            if doc.get("due_date") and date.fromisoformat(doc["due_date"]) < today:
                out["overdue_total"] += doc["amount"]
                out["overdue_count"] += 1
                oldest = out["overdue_by_client"].get(doc["client_id"])
                out["overdue_by_client"][doc["client_id"]] = min(oldest, doc["due_date"]) if oldest else doc["due_date"]
    for client_id in ledger_clients:
        charges, payments = await _ledger(client_id)
        allocated, _, _ = _allocate(charges, payments)
        for c in allocated:
            if c["open"] <= 0:
                continue
            out["expected_total"] += c["open"]
            if c.get("due_date") and date.fromisoformat(c["due_date"]) < today:
                out["overdue_total"] += c["open"]
                out["overdue_count"] += 1
                oldest = out["overdue_by_client"].get(client_id)
                out["overdue_by_client"][client_id] = min(oldest, c["due_date"]) if oldest else c["due_date"]
    for key in ("received_month", "received_year", "expected_total", "overdue_total"):
        out[key] = round(out[key], 2)
    return out


async def _ledger_open_map(client_ids) -> Dict[str, float]:
    """charge id -> still unpaid, for the ledger charges of these clients."""
    out: Dict[str, float] = {}
    for client_id in client_ids:
        charges, payments = await _ledger(client_id)
        for c in _allocate(charges, payments)[0]:
            out[c["id"]] = c["open"]
    return out


# ---------------- Calendar billing: ledger, drafts, blocking (PLAN_NAPLATA_KALENDARSKA.md, N2) ----------------
# Only entries with `ledger: true` count toward a client's balance: charges made
# from approved drafts (status `expected`, never flipped) and payments recorded
# through POST /billing/payments (status `received`). Older manual entries keep
# working as before but stay out of the balance.
_BILLING_KEPT_FIELDS = (
    "billing_cycle", "balance_due", "credit", "open_charges", "monthly_amount", "extra_seats_billed", "manual_block",
    "due_extension_until", "billing_emails", "block_notified",
)


def _month_end(day: date) -> date:
    return _add_months(date(day.year, day.month, 1), 1) - timedelta(days=1)


def _month_label(day: date) -> str:
    return day.strftime("%m.%Y")


class ChargeLine(BaseModel):
    kind: str  # package | seats | seats_prorated | discount | manual | free
    label: str
    amount: float
    role: Optional[str] = None
    qty: Optional[int] = None
    unit_price: Optional[float] = None
    months: Optional[int] = None


class LedgerCharge(BaseModel):
    id: str
    status: str  # draft | expected | canceled
    number: Optional[str] = None
    amount: float
    paid: float = 0
    open: float = 0
    due_date: Optional[str] = None
    period_start: Optional[str] = None
    period_end: Optional[str] = None
    cycle: Optional[str] = None
    kind: str = "subscription"
    lines: List[ChargeLine] = Field(default_factory=list)
    note: Optional[str] = None
    calculated_total: Optional[float] = None
    approved_at: Optional[str] = None
    created_at: str


class LedgerPayment(BaseModel):
    id: str
    amount: float
    paid_at: Optional[str] = None
    method: Optional[str] = None
    note: Optional[str] = None
    created_at: str


class Balance(BaseModel):
    client_id: str
    balance_due: float  # what is owed now (charges - payments), 0 when nothing is owed
    credit: float  # overpayment carried to the next charge
    monthly_amount: float
    charges: List[LedgerCharge]
    payments: List[LedgerPayment]


async def _ledger(client_id: str) -> Tuple[List[dict], List[dict]]:
    charges = await db.payments.find(
        {"client_id": client_id, "ledger": True, "status": "expected"}, {"_id": 0}
    ).to_list(None)
    payments = await db.payments.find(
        {"client_id": client_id, "ledger": True, "status": "received"}, {"_id": 0}
    ).to_list(None)
    charges.sort(key=lambda c: (c.get("due_date") or "", c.get("approved_at") or c["created_at"]))
    payments.sort(key=lambda p: (p.get("paid_at") or "", p["created_at"]))
    return charges, payments


def _allocate(charges: List[dict], payments: List[dict]) -> Tuple[List[dict], float, float]:
    """FIFO: money in closes the oldest charges first. Returns charges with
    paid/open, the amount owed and the credit left over."""
    pool = round(sum(p["amount"] for p in payments), 2)
    out = []
    for c in charges:
        paid = round(min(pool, c["amount"]), 2)
        pool = round(pool - paid, 2)
        out.append({**c, "paid": paid, "open": round(c["amount"] - paid, 2)})
    owed = round(sum(c["open"] for c in out), 2)
    return out, owed, max(0.0, pool)


async def _monthly_amount(client_id: str) -> float:
    """What one month costs this client (package + extra accounts - discount); the debt
    threshold for blocking. Annual-only plans: the annual price / 12."""
    try:
        seats = await _seats_out(client_id)
    except HTTPException:
        return 0.0
    if seats.monthly_total:
        return round(seats.monthly_total, 2)
    client = await db.clients.find_one({"id": client_id}, {"_id": 0, "subscription": 1})
    sub = (client or {}).get("subscription") or {}
    plan = await db.plans.find_one({"id": sub.get("plan_id")}, {"_id": 0}) if sub.get("plan_id") else None
    yearly = ((plan or {}).get("prices") or {}).get("12")
    return round(float(yearly) / 12, 2) if yearly else 0.0


async def _apply_ledger(client_id: str, user: Optional[User] = None) -> dict:
    """Recomputes the balance after any change, stores it on the subscription (the
    lock check reads it on every request) and moves the end date to the end of the
    latest period that has money (or a zero amount) on it. Never shortens it."""
    client = await db.clients.find_one({"id": client_id}, {"_id": 0})
    sub = (client or {}).get("subscription")
    charges, payments = await _ledger(client_id)
    allocated, owed, credit = _allocate(charges, payments)
    if not sub:
        return {"balance_due": owed, "credit": credit}
    updates: Dict[str, Any] = {
        "subscription.balance_due": owed,
        "subscription.credit": round(credit, 2),
        "subscription.open_charges": [{"due_date": c.get("due_date"), "open": c["open"]} for c in allocated if c["open"] > 0],
        "subscription.monthly_amount": await _monthly_amount(client_id),
    }
    covered = [c["period_end"] for c in allocated if c.get("period_end") and (c["paid"] > 0 or c["amount"] == 0)]
    if covered and sub.get("billing_cycle"):
        new_end = max(covered)
        if new_end > sub["ends_at"]:
            updates.update({
                "subscription.ends_at": new_end, "subscription.canceled_at": None,
                "subscription.reminder_for": None, "subscription.purge_warned_for": None,
            })
            await _log_subscription_event(
                client_id, "extended", user, plan_name=sub.get("plan_name"), ends_at=new_end, source="billing",
            )
    await db.clients.update_one({"id": client_id}, {"$set": updates})
    return {"balance_due": owed, "credit": round(credit, 2)}


async def _next_proforma_number(day: date) -> str:
    counter = await db.counters.find_one_and_update(
        {"_id": f"proforma:{day.year}"}, {"$inc": {"seq": 1}}, upsert=True, return_document=ReturnDocument.AFTER,
    )
    return f"PR-{day.year}-{counter['seq']:04d}"


def _charge_out(doc: dict) -> LedgerCharge:
    return LedgerCharge(**{
        **doc,
        "lines": doc.get("lines") or [],
        "paid": doc.get("paid", 0),
        "open": doc.get("open", doc["amount"]),
        "calculated_total": (doc.get("breakdown") or {}).get("calculated_total"),
    })


async def _balance_out(client_id: str, include_drafts: bool = False) -> Balance:
    client = await _client_or_404(client_id)
    charges, payments = await _ledger(client_id)
    allocated, owed, credit = _allocate(charges, payments)
    rows = [_charge_out(c) for c in allocated]
    if include_drafts:
        drafts = await db.payments.find({"client_id": client_id, "ledger": True, "status": "draft"}, {"_id": 0}).to_list(None)
        rows += [_charge_out(d) for d in drafts]
    rows.sort(key=lambda c: (c.period_start or "", c.created_at), reverse=True)
    return Balance(
        client_id=client_id, balance_due=owed, credit=round(credit, 2),
        monthly_amount=float(((client.get("subscription") or {}).get("monthly_amount")) or 0),
        charges=rows,
        payments=[LedgerPayment(**p) for p in sorted(payments, key=lambda p: (p.get("paid_at") or "", p["created_at"]), reverse=True)],
    )


async def _draft_lines(client: dict, cycle: str, period_start: date, period_end: date) -> Tuple[List[dict], float]:
    """Lines of the next charge: package, extra accounts, discount."""
    sub = client["subscription"]
    plan = await db.plans.find_one({"id": sub["plan_id"]}, {"_id": 0}) or {}
    prices = plan.get("prices") or {}
    figures = _seat_figures(plan, sub)
    modules = effective_modules(client)
    lines: List[dict] = []
    if cycle == "annual":
        package = prices.get("12")
        if package is None and prices.get("1") is not None:
            package = prices["1"] * 12
        months = 12
    else:
        package = prices.get("1")
        months = 1
    if package:
        lines.append({"kind": "package", "label": f"{sub.get('plan_name') or plan.get('name')} ({_month_label(period_start)} - {_month_label(period_end)})", "amount": round(float(package), 2)})
    for role in SEAT_ROLES:
        if role == "warehouse" and Module.WAREHOUSE.value not in modules:
            continue
        f = figures[role]
        if f["extra"] and f["price"]:
            lines.append({
                "kind": "seats", "label": f"extra:{role}", "role": role, "qty": f["extra"],
                "unit_price": f["price"], "months": months, "amount": round(f["extra"] * f["price"] * months, 2),
            })
    discount = float(sub.get("discount_percent") or 0)
    subtotal = round(sum(line["amount"] for line in lines), 2)
    if discount and subtotal:
        lines.append({"kind": "discount", "label": f"{discount:g}%", "amount": round(-subtotal * discount / 100, 2)})
    return lines, round(sum(line["amount"] for line in lines), 2)


async def _prorated_seat_lines(client: dict, today: date) -> Tuple[List[dict], Dict[str, int]]:
    """Annual clients that added accounts since the year was paid: those accounts
    to the end of the paid year, whole months from next month's 1st."""
    sub = client["subscription"]
    plan = await db.plans.find_one({"id": sub["plan_id"]}, {"_id": 0}) or {}
    figures = _seat_figures(plan, sub)
    billed = sub.get("extra_seats_billed") or {}
    start = _add_months(date(today.year, today.month, 1), 1)
    ends = date.fromisoformat(sub["ends_at"])
    months = (ends.year - start.year) * 12 + ends.month - start.month + 1
    lines: List[dict] = []
    if months <= 0:
        return lines, {}
    for role in SEAT_ROLES:
        added = figures[role]["extra"] - int(billed.get(role, 0))
        if added > 0 and figures[role]["price"]:
            lines.append({
                "kind": "seats_prorated", "label": f"extra:{role}", "role": role, "qty": added,
                "unit_price": figures[role]["price"], "months": months,
                "amount": round(added * figures[role]["price"] * months, 2),
            })
    return lines, {r: figures[r]["extra"] for r in SEAT_ROLES}


async def _make_drafts(today: date, client_ids: Optional[List[str]] = None, force: bool = False) -> List[dict]:
    """Drafts for the period that starts next month: monthly clients every month,
    annual clients when their year ends this month (and prorated extra accounts).
    One per client and period; `force` ignores canceled ones (a fresh draft after a
    wrong one was dropped)."""
    start = _add_months(date(today.year, today.month, 1), 1)
    made: List[dict] = []
    query: Dict[str, Any] = {"subscription.billing_cycle": {"$in": ["monthly", "annual"]}, "active": {"$ne": False}}
    if client_ids:
        query["id"] = {"$in": client_ids}
    async for client in db.clients.find(query, {"_id": 0}):
        sub = client["subscription"]
        if sub.get("canceled_at") or not sub.get("plan_id"):
            continue
        ends = date.fromisoformat(sub["ends_at"])
        cycle = sub["billing_cycle"]
        taken_query = {"client_id": client["id"], "ledger": True, "period_start": start.isoformat(), "kind": "subscription"}
        if force:
            taken_query["status"] = {"$ne": "canceled"}
        if ends >= start:  # the next period is already paid (annual mid-year)
            if cycle == "annual":
                lines, billed = await _prorated_seat_lines(client, today)
                addon_query = {**taken_query, "cycle": "annual_addon"}
                if lines and not await db.payments.find_one(addon_query, {"_id": 1}):
                    made.append(await _insert_draft(client, "annual_addon", start, ends, lines, billed))
            continue
        if await db.payments.find_one({**taken_query, "cycle": {"$ne": "annual_addon"}}, {"_id": 1}):
            continue
        period_end = _month_end(_add_months(start, 11)) if cycle == "annual" else _month_end(start)
        lines, _ = await _draft_lines(client, cycle, start, period_end)
        if not lines:
            continue  # nothing priced (free plan)
        billed = (sub.get("extra_seats") or {}) if cycle == "annual" else None
        made.append(await _insert_draft(client, cycle, start, period_end, lines, billed))
    return made


async def _insert_draft(client: dict, cycle: str, start: date, end: date, lines: List[dict], billed: Optional[dict]) -> dict:
    total = round(sum(line["amount"] for line in lines), 2)
    doc = {
        "id": str(uuid.uuid4()), "client_id": client["id"], "status": "draft", "ledger": True, "kind": "subscription",
        "cycle": cycle, "amount": max(total, 0.0), "currency": PAYMENT_CURRENCY,
        "period_start": start.isoformat(), "period_end": end.isoformat(), "period": start.strftime("%Y-%m"),
        "period_months": 12 if cycle == "annual" else 1, "lines": lines,
        "breakdown": {"calculated_total": max(total, 0.0)}, "seats_billed": billed,
        "plan_name": client["subscription"].get("plan_name"), "due_date": None, "paid_at": None, "method": None,
        "note": None, "created_by": "system", "source": "billing", "created_at": now_iso(), "canceled_at": None,
    }
    await db.payments.insert_one({**doc, "_id": doc["id"]})
    return doc


class DraftRow(BaseModel):
    charge: LedgerCharge
    client_id: str
    client_name: str
    balance_due: float  # owed before this draft
    credit: float
    to_pay: float  # balance after approving it


class DraftUpdate(BaseModel):
    amount: Optional[float] = Field(default=None, ge=0, le=1_000_000_000)
    note: Optional[str] = Field(default=None, max_length=500)
    free: Optional[bool] = None  # True = 0 din, note "Besplatan mesec"


class DraftApproveInput(BaseModel):
    ids: Optional[List[str]] = None
    all: bool = False


class DraftGenerateInput(BaseModel):
    client_ids: Optional[List[str]] = None
    month: Optional[str] = None  # YYYY-MM of the period to draft (default: next month)


async def _draft_or_404(draft_id: str) -> dict:
    doc = await db.payments.find_one({"id": draft_id, "ledger": True, "status": "draft"}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Draft not found")
    return doc


async def _draft_row(doc: dict, names: Dict[str, str]) -> DraftRow:
    charges, payments = await _ledger(doc["client_id"])
    _, owed, credit = _allocate(charges, payments)
    return DraftRow(
        charge=_charge_out(doc), client_id=doc["client_id"], client_name=names.get(doc["client_id"], ""),
        balance_due=owed, credit=credit, to_pay=round(max(0.0, owed + doc["amount"] - credit), 2),
    )


@api_router.get("/billing/drafts", response_model=List[DraftRow])
async def list_drafts(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    names = await _client_names()
    docs = await db.payments.find({"ledger": True, "status": "draft"}, {"_id": 0}).sort("created_at", 1).to_list(None)
    return [await _draft_row(d, names) for d in docs]


@api_router.post("/billing/drafts/generate", response_model=List[DraftRow])
async def generate_drafts(inp: DraftGenerateInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Drafts now, without waiting for the draft day (for one client or all)."""
    base = _today()
    if inp.month:
        month = date.fromisoformat(_valid_month(inp.month) + "-01")
        base = _add_months(month, -1)
    made = await _make_drafts(base, inp.client_ids, force=True)
    await _audit(current_user, "billing.drafts_generated", None, None, None, None, {"count": len(made)})
    names = await _client_names()
    return [await _draft_row(d, names) for d in made]


@api_router.put("/billing/drafts/{draft_id}", response_model=DraftRow)
async def update_draft(draft_id: str, inp: DraftUpdate, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    doc = await _draft_or_404(draft_id)
    changes: Dict[str, Any] = {}
    if inp.free:
        changes.update(amount=0.0, note=inp.note or "Besplatan mesec")
    else:
        if inp.amount is not None:
            changes["amount"] = round(inp.amount, 2)
        if inp.note is not None:
            changes["note"] = inp.note.strip() or None
    if changes:
        await db.payments.update_one({"id": draft_id}, {"$set": changes})
    fresh = await _draft_or_404(draft_id)
    return await _draft_row(fresh, await _client_names())


@api_router.post("/billing/drafts/approve", response_model=List[LedgerCharge])
async def approve_drafts(inp: DraftApproveInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Drafts become charges the client sees: a pro-forma number, a due date
    (settings.due_day of the period's month) and the balance is recomputed."""
    if not inp.all and not inp.ids:
        raise HTTPException(status_code=400, detail="Send ids or all=true")
    query: Dict[str, Any] = {"ledger": True, "status": "draft"}
    if not inp.all:
        query["id"] = {"$in": inp.ids}
    docs = await db.payments.find(query, {"_id": 0}).to_list(None)
    approved = []
    for doc in docs:
        start = date.fromisoformat(doc["period_start"])
        update = {
            "status": "expected", "approved_at": now_iso(), "approved_by": current_user.name,
            "number": await _next_proforma_number(_today()),
            "due_date": date(start.year, start.month, int(setting("due_day"))).isoformat(),
        }
        result = await db.payments.update_one({"id": doc["id"], "status": "draft"}, {"$set": update})
        if result.modified_count == 0:
            continue  # approved meanwhile
        if doc.get("seats_billed") is not None:
            await db.clients.update_one({"id": doc["client_id"]}, {"$set": {"subscription.extra_seats_billed": doc["seats_billed"]}})
        await _log_subscription_event(
            doc["client_id"], "charge_created", current_user, plan_name=doc.get("plan_name"),
            data={"amount": doc["amount"], "currency": PAYMENT_CURRENCY, "number": update["number"], "period": doc.get("period")},
        )
        await _apply_ledger(doc["client_id"], current_user)
        approved.append({**doc, **update})
    by_client: Dict[str, List[dict]] = {}
    for doc in approved:
        by_client.setdefault(doc["client_id"], []).append(doc)
    for client_id, charges in by_client.items():
        await _notify_issued(client_id, charges)
    return [_charge_out(doc) for doc in approved]


class LedgerPaymentInput(BaseModel):
    client_id: str
    amount: float = Field(gt=0, le=1_000_000_000)
    paid_at: Optional[str] = None
    method: Literal["bank", "card", "cash", "other"] = "bank"
    note: Optional[str] = Field(default=None, max_length=500)
    send_receipt: bool = False  # email a payment confirmation to the client
    billing_emails: Optional[bool] = None  # also switch the client's billing emails on/off


class LedgerPaymentResult(BaseModel):
    payment: LedgerPayment
    balance: Balance
    subscription: SubscriptionRow


@api_router.post("/billing/payments", response_model=LedgerPaymentResult)
async def record_ledger_payment(inp: LedgerPaymentInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Money in for a client on calendar billing: closes the oldest charges first,
    a surplus stays as credit, the end date moves to the paid period and a block
    caused by the debt lifts."""
    client = await _client_or_404(inp.client_id)
    if not (client.get("subscription") or {}).get("billing_cycle"):
        raise HTTPException(status_code=400, detail="Client is not on calendar billing")
    doc = {
        "id": str(uuid.uuid4()), "client_id": client["id"], "status": "received", "ledger": True, "kind": "subscription",
        "amount": round(inp.amount, 2), "currency": PAYMENT_CURRENCY, "paid_at": _paid_day(inp.paid_at),
        "due_date": None, "method": inp.method, "note": inp.note, "plan_name": client["subscription"].get("plan_name"),
        "created_by": current_user.name, "source": "manual", "created_at": now_iso(), "canceled_at": None,
    }
    await db.payments.insert_one({**doc, "_id": doc["id"]})
    await _log_subscription_event(
        client["id"], "payment_received", current_user, note=inp.note, plan_name=doc["plan_name"],
        data={"amount": doc["amount"], "currency": PAYMENT_CURRENCY},
    )
    await _apply_ledger(client["id"], current_user)
    if inp.billing_emails is not None:
        await db.clients.update_one({"id": client["id"]}, {"$set": {"subscription.billing_emails": inp.billing_emails}})
    if inp.send_receipt:
        await _notify_paid(client["id"], doc)
    return LedgerPaymentResult(
        payment=LedgerPayment(**doc), balance=await _balance_out(client["id"]),
        subscription=await _subscription_row(await _client_or_404(client["id"])),
    )


@api_router.get("/billing/clients/{client_id}", response_model=Balance)
async def client_balance(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    return await _balance_out(client_id, include_drafts=True)


class EnrollInput(BaseModel):
    cycle: Literal["monthly", "annual"] = "monthly"


async def _enroll_calendar(
    client_id: str, cycle: str, user: Optional[User], *, from_today: bool = True, free_month: bool = True,
    source: str = "manual",
) -> None:
    """Puts a client with a subscription on calendar billing (or changes its cycle).
    The end date goes to the end of its month - from today on (`from_today`), or of the
    stored end date as it is (the startup migration: nobody gets unlocked or shortened).
    A new client's first partial month is a 0 charge "Besplatan mesec"."""
    client = await _client_or_404(client_id)
    sub = client.get("subscription")
    if not sub or not sub.get("plan_id"):
        raise HTTPException(status_code=400, detail="Client has no subscription - assign a plan first")
    today = _today()
    ends = date.fromisoformat(sub["ends_at"])
    aligned = _month_end(max(ends, today) if from_today else ends)
    updates: Dict[str, Any] = {"subscription.billing_cycle": cycle, "subscription.ends_at": aligned.isoformat()}
    if cycle == "annual" and sub.get("extra_seats_billed") is None:
        updates["subscription.extra_seats_billed"] = sub.get("extra_seats") or {}
    await db.clients.update_one({"id": client_id}, {"$set": updates})
    if free_month and not sub.get("billing_cycle") and aligned == _month_end(today) and setting("first_month_free"):
        month_start = date(today.year, today.month, 1)
        doc = {
            "id": str(uuid.uuid4()), "client_id": client_id, "status": "expected", "ledger": True, "kind": "subscription",
            "cycle": "monthly", "amount": 0.0, "currency": PAYMENT_CURRENCY, "period_start": month_start.isoformat(),
            "period_end": aligned.isoformat(), "period": month_start.strftime("%Y-%m"), "period_months": 1,
            "lines": [{"kind": "free", "label": "Besplatan mesec", "amount": 0.0}], "note": "Besplatan mesec",
            "number": await _next_proforma_number(today), "due_date": today.isoformat(), "approved_at": now_iso(),
            "approved_by": user.name if user else "system", "plan_name": sub.get("plan_name"), "paid_at": None, "method": None,
            "created_by": user.name if user else "system", "source": "billing", "created_at": now_iso(), "canceled_at": None,
        }
        await db.payments.insert_one({**doc, "_id": doc["id"]})
    await _log_subscription_event(
        client_id, "billing_enrolled", user, plan_name=sub.get("plan_name"), ends_at=aligned.isoformat(),
        data={"cycle": cycle}, source=source,
    )
    await _apply_ledger(client_id, user)


async def _migrate_calendar_billing() -> None:
    """One-off at startup: every client with a plan that is still on the old model
    goes to monthly calendar billing (annual when its plan has only an annual price)."""
    async for client in db.clients.find(
        {"subscription.plan_id": {"$ne": None}, "subscription.billing_cycle": {"$exists": False}}, {"_id": 0}
    ):
        try:
            plan = await db.plans.find_one({"id": client["subscription"]["plan_id"]}, {"_id": 0}) or {}
            prices = plan.get("prices") or {}
            cycle = "annual" if prices.get("12") is not None and prices.get("1") is None else "monthly"
            await _enroll_calendar(client["id"], cycle, None, from_today=False, free_month=False, source="migration")
            logger.info("Calendar billing: moved client %s (%s)", client["id"], cycle)
        except Exception as exc:
            logger.error("Calendar billing migration failed for %s: %s", client.get("id"), exc)


@api_router.post("/billing/clients/{client_id}/enroll", response_model=SubscriptionRow)
async def enroll_calendar_billing(client_id: str, inp: EnrollInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Calendar billing cycle of a client: monthly or annual (every client with a plan
    is on calendar billing; this changes the cycle or re-enrolls a client)."""
    await _enroll_calendar(client_id, inp.cycle, current_user)
    return await _subscription_row(await _client_or_404(client_id))


class BlockInput(BaseModel):
    note: Optional[str] = Field(default=None, max_length=500)


class DueExtensionInput(BaseModel):
    until: Optional[str] = None  # YYYY-MM-DD; None = remove the extension
    note: Optional[str] = Field(default=None, max_length=500)


@api_router.post("/subscriptions/{client_id}/block", response_model=SubscriptionRow)
async def block_client(client_id: str, inp: BlockInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    client = await _client_or_404(client_id)
    if not (client.get("subscription") or {}).get("billing_cycle"):
        raise HTTPException(status_code=400, detail="Client is not on calendar billing")
    await db.clients.update_one(
        {"id": client_id}, {"$set": {"subscription.manual_block": {"since": _today().isoformat(), "by": current_user.name, "note": inp.note}}}
    )
    await _log_subscription_event(client_id, "blocked", current_user, note=inp.note, plan_name=client["subscription"].get("plan_name"))
    return await _subscription_row(await _client_or_404(client_id))


@api_router.post("/subscriptions/{client_id}/unblock", response_model=SubscriptionRow)
async def unblock_client(client_id: str, inp: BlockInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    client = await _client_or_404(client_id)
    await db.clients.update_one({"id": client_id}, {"$set": {"subscription.manual_block": None}})
    await _log_subscription_event(client_id, "unblocked", current_user, note=inp.note, plan_name=(client.get("subscription") or {}).get("plan_name"))
    return await _subscription_row(await _client_or_404(client_id))


@api_router.post("/subscriptions/{client_id}/extend-due", response_model=SubscriptionRow)
async def extend_due(client_id: str, inp: DueExtensionInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Exception for one client: no block until this date (e.g. paid, statement not in yet)."""
    client = await _client_or_404(client_id)
    if not (client.get("subscription") or {}).get("billing_cycle"):
        raise HTTPException(status_code=400, detail="Client is not on calendar billing")
    until = _valid_day(inp.until, "until") if inp.until else None
    await db.clients.update_one({"id": client_id}, {"$set": {"subscription.due_extension_until": until}})
    await _log_subscription_event(client_id, "due_extended", current_user, note=inp.note, plan_name=client["subscription"].get("plan_name"), ends_at=until)
    return await _subscription_row(await _client_or_404(client_id))


async def _billing_attention(today: date, active_clients: Dict[str, dict]) -> List["AttentionItem"]:
    """Drafts waiting for approval and clients with debt in two or more periods."""
    items: List[AttentionItem] = []
    drafts: Dict[str, int] = {}
    async for d in db.payments.find({"ledger": True, "status": "draft"}, {"_id": 0, "client_id": 1}):
        drafts[d["client_id"]] = drafts.get(d["client_id"], 0) + 1
    for client_id, count in drafts.items():
        c = active_clients.get(client_id)
        if c:
            items.append(AttentionItem(client_id=client_id, client_name=c["name"], reason="drafts_pending", plan_name=(c.get("subscription") or {}).get("plan_name")))
    for client_id, c in active_clients.items():
        if not (c.get("subscription") or {}).get("billing_cycle") or not (c["subscription"].get("balance_due") or 0):
            continue
        charges, payments = await _ledger(client_id)
        allocated, _, _ = _allocate(charges, payments)
        late = [x for x in allocated if x["open"] > 0 and x.get("due_date") and x["due_date"] < today.isoformat()]
        if len(late) >= 2:
            items.append(AttentionItem(client_id=client_id, client_name=c["name"], reason="debt_repeated", date=late[0]["due_date"], plan_name=c["subscription"].get("plan_name")))
    return items


# ---------------- What the client sees: charges, payment details, IPS QR (N3) ----------------
def _ips_reference(number: Optional[str]) -> Optional[str]:
    """Pro-forma PR-2026-0001 -> reference 97 + control + 20260001 (model 97, ISO 7064 mod 97-10)."""
    if not number:
        return None
    digits = re.sub(r"\D", "", number)
    if not digits:
        return None
    control = 98 - (int(digits) * 100) % 97
    return f"97{control:02d}{digits}"


def _ips_string(amount: float, number: Optional[str]) -> Optional[str]:
    """NBS IPS QR payload (K:PR) for paying `amount`; None when the account or the
    recipient isn't set in the settings. Empty optional tags are left out."""
    account, name = setting("ips_account"), setting("ips_recipient_name")
    if not account or not name or amount <= 0:
        return None
    parts = [
        "K:PR", "V:01", "C:1", f"R:{account}", f"N:{name[:70]}",
        "I:RSD" + f"{amount:.2f}".replace(".", ","),
    ]
    code = setting("ips_payment_code")
    if code:
        parts.append(f"SF:{code}")
    purpose = setting("ips_purpose")
    if purpose:
        parts.append(f"S:{purpose[:35]}")
    reference = _ips_reference(number) if setting("ips_reference_from_number") else None
    if reference:
        parts.append(f"RO:{reference}")
    return "|".join(parts)


def _qr_png(text: str) -> bytes:
    import io
    import segno

    buf = io.BytesIO()
    segno.make(text, error="m").save(buf, kind="png", scale=6, border=3)
    return buf.getvalue()


async def _billing_notice(raw: dict) -> Optional[dict]:
    """Banner data for /auth/me: the earliest unpaid due date, whether it is late,
    the day the account will be blocked; amounts only for the client's admin."""
    client = await db.clients.find_one({"id": raw.get("client_id")}, {"_id": 0, "subscription": 1})
    sub = (client or {}).get("subscription") or {}
    if not sub.get("billing_cycle"):
        return None
    today = _today()
    state = subscription_state(sub, today)
    dues = sorted(c["due_date"] for c in (sub.get("open_charges") or []) if c.get("due_date") and c.get("open", 0) > 0)
    next_due = dues[0] if dues else None
    blocks_on = None
    if next_due:
        due = date.fromisoformat(next_due)
        blocks_on = (date(due.year, due.month, int(setting("block_after_day"))) + timedelta(days=1)).isoformat()
    out: Dict[str, Any] = {
        "status": state["status"], "lock_reason": state.get("lock_reason"), "next_due_date": next_due,
        "late": bool(next_due and next_due <= today.isoformat()), "blocks_on": blocks_on,
    }
    if raw.get("role") == Role.ADMIN.value:
        out.update(to_pay=round(float(sub.get("balance_due") or 0), 2), credit=round(float(sub.get("credit") or 0), 2))
    return out


class IpsInfo(BaseModel):
    available: bool
    recipient: Optional[str] = None
    account: Optional[str] = None
    payment_code: Optional[str] = None
    purpose: Optional[str] = None
    reference: Optional[str] = None
    amount: float = 0


class BillingMe(BaseModel):
    client_name: str
    state: Dict[str, Any]
    to_pay: float  # owed now, all open charges
    credit: float
    next_due_date: Optional[str] = None
    monthly_estimate: float  # what one month costs now (an estimate, the charge is what counts)
    charges: List[LedgerCharge]  # approved charges, newest first, with what is still open
    payments: List[LedgerPayment]
    instructions_sr: str = ""
    instructions_en: str = ""
    ips: IpsInfo


async def _billing_for_client(client_id: str) -> Tuple[BillingMe, Optional[str]]:
    client = await _client_or_404(client_id)
    balance = await _balance_out(client_id)
    open_rows = sorted((c for c in balance.charges if c.open > 0), key=lambda c: (c.due_date or "", c.created_at))
    reference_number = open_rows[0].number if open_rows else None
    payload = _ips_string(balance.balance_due, reference_number)
    sub = client.get("subscription") or {}
    me = BillingMe(
        client_name=client["name"], state=subscription_state(sub) if sub else {"status": "none"},
        to_pay=balance.balance_due, credit=balance.credit,
        next_due_date=open_rows[0].due_date if open_rows else None,
        monthly_estimate=await _monthly_amount(client_id) if sub else 0,
        charges=balance.charges, payments=balance.payments,
        instructions_sr=setting("payment_instructions_sr") or "", instructions_en=setting("payment_instructions_en") or "",
        ips=IpsInfo(
            available=payload is not None, recipient=setting("ips_recipient_name") or None,
            account=setting("ips_account") or None, payment_code=setting("ips_payment_code") or None,
            purpose=setting("ips_purpose") or None,
            reference=_ips_reference(reference_number) if setting("ips_reference_from_number") else None,
            amount=balance.balance_due,
        ),
    )
    return me, payload


@api_router.get("/billing/me", response_model=BillingMe)
async def my_billing(current_user: User = Depends(require_roles(Role.ADMIN))):
    """The client's admin: charges with lines and numbers, what is owed, payment
    details. Works while the client is blocked (see BLOCKED_ADMIN_PATHS)."""
    me, _ = await _billing_for_client(current_user.client_id)
    return me


def _png_response(data: bytes):
    from fastapi.responses import Response

    return Response(content=data, media_type="image/png", headers={"Cache-Control": "no-store"})


@api_router.get("/billing/me/qr.png")
async def my_billing_qr(current_user: User = Depends(require_roles(Role.ADMIN))):
    _, payload = await _billing_for_client(current_user.client_id)
    if not payload:
        raise HTTPException(status_code=404, detail="Nothing to pay or payment details are not set")
    return _png_response(_qr_png(payload))


@api_router.get("/billing/clients/{client_id}/qr.png")
async def client_billing_qr(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    _, payload = await _billing_for_client(client_id)
    if not payload:
        raise HTTPException(status_code=404, detail="Nothing to pay or payment details are not set")
    return _png_response(_qr_png(payload))


# ---------------- Billing emails (PLAN_NAPLATA_KALENDARSKA.md, N4) ----------------
# Sent to the client's active admins + its accounting address (one email per
# address). Every email is logged in `billing_emails` (also when it isn't sent:
# `skipped` = disabled / no_recipients / smtp_off) and charge emails are marked
# on the charge (`notified.<event>`), so the daily cron never sends one twice.
_ROLE_NAMES = {
    "sr": {"admin": "administrator", "operator": "komercijalista", "warehouse": "magacioner"},
    "en": {"admin": "admin", "operator": "sales rep", "warehouse": "warehouse"},
}
_MONTHS = {
    "sr": ["januar", "februar", "mart", "april", "maj", "jun", "jul", "avgust", "septembar", "oktobar", "novembar", "decembar"],
    "en": ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"],
}


def _money(amount: float, lang: str = "sr") -> str:
    """10800 -> 10.800,00 RSD (sr) / 10,800.00 RSD (en)."""
    text = f"{float(amount):,.2f}"
    if lang == "sr":
        text = text.replace(",", " ").replace(".", ",").replace(" ", ".")
    return f"{text} {PAYMENT_CURRENCY}"


def _day_text(value: Optional[str]) -> str:
    return date.fromisoformat(value).strftime("%d.%m.%Y") if value else ""


def _period_text(charge: dict, lang: str) -> str:
    if not charge.get("period_start"):
        return ""
    start, end = date.fromisoformat(charge["period_start"]), date.fromisoformat(charge["period_end"])
    if (start.year, start.month) == (end.year, end.month):
        return f"{_MONTHS[lang][start.month - 1]} {start.year}"
    return f"{_day_text(charge['period_start'])} - {_day_text(charge['period_end'])}"


def _line_text(line: dict, lang: str) -> str:
    kind = line.get("kind")
    if kind in ("seats", "seats_prorated"):
        role = _ROLE_NAMES[lang].get(line.get("role") or "", line.get("role") or "")
        months = line.get("months") or 1
        head = "Dodatni nalozi" if lang == "sr" else "Extra accounts"
        text = f"{head}: {role} × {line.get('qty')}"
        if months > 1:
            text += f" × {months} {'mes.' if lang == 'sr' else 'mo.'}"
        return text
    if kind == "discount":
        return f"{'Popust' if lang == 'sr' else 'Discount'} {line.get('label')}"
    return line.get("label") or ""


def _blocks_on(due_date: str) -> date:
    due = date.fromisoformat(due_date)
    return date(due.year, due.month, int(setting("block_after_day"))) + timedelta(days=1)


class _Mail:
    """One email in one language: paragraphs, an optional table and payment details."""

    def __init__(self, lang: str):
        self.lang = lang
        self.text: List[str] = []
        self.html: List[str] = []

    def p(self, value: str, bold: bool = False):
        self.text.append(value)
        body = html_lib.escape(value)
        self.html.append(f"<p style=\"margin:0 0 12px\">{'<b>' + body + '</b>' if bold else body}</p>")

    def table(self, rows: List[Tuple[str, str]], total: Optional[Tuple[str, str]] = None):
        width = max([len(r[0]) for r in rows + ([total] if total else [])] + [10])
        self.text.append("\n".join(f"{label.ljust(width)}  {value}" for label, value in rows + ([total] if total else [])))
        cells = "".join(
            f"<tr><td style=\"padding:4px 12px 4px 0\">{html_lib.escape(a)}</td><td style=\"padding:4px 0;text-align:right;white-space:nowrap\">{html_lib.escape(b)}</td></tr>"
            for a, b in rows
        )
        if total:
            cells += (
                f"<tr><td style=\"padding:6px 12px 4px 0;border-top:1px solid #ccc\"><b>{html_lib.escape(total[0])}</b></td>"
                f"<td style=\"padding:6px 0 4px;border-top:1px solid #ccc;text-align:right;white-space:nowrap\"><b>{html_lib.escape(total[1])}</b></td></tr>"
            )
        self.html.append(f"<table style=\"border-collapse:collapse;margin:0 0 12px\">{cells}</table>")

    def details(self, me: "BillingMe", with_qr: bool):
        sr = self.lang == "sr"
        ips = me.ips
        rows = [
            ("Primalac" if sr else "Recipient", ips.recipient or ""),
            ("Račun" if sr else "Account", ips.account or ""),
            ("Šifra plaćanja" if sr else "Payment code", ips.payment_code or ""),
            ("Svrha" if sr else "Purpose", ips.purpose or ""),
            ("Poziv na broj" if sr else "Reference", ips.reference or ""),
        ]
        rows = [r for r in rows if r[1]]
        if rows:
            self.p("Podaci za uplatu" if sr else "Payment details", bold=True)
            self.table(rows)
        extra = me.instructions_sr if sr else (me.instructions_en or me.instructions_sr)
        if extra:
            self.p(extra)
        if with_qr:
            self.p(
                "Možete platiti i skeniranjem IPS QR koda u aplikaciji vaše banke." if sr
                else "You can also pay by scanning the IPS QR code in your banking app."
            )
            self.html.append("<p style=\"margin:0 0 12px\"><img src=\"cid:ips-qr\" alt=\"IPS QR\" width=\"180\" height=\"180\"></p>")

    def link(self, label: str, url: str):
        self.text.append(f"{label}: {url}")
        self.html.append(f"<p style=\"margin:0 0 12px\">{html_lib.escape(label)}: <a href=\"{html_lib.escape(url)}\">{html_lib.escape(url)}</a></p>")

    def footer(self):
        sr = self.lang == "sr"
        self.p("Srdačan pozdrav,\nEasy Order" if sr else "Kind regards,\nEasy Order")
        note = (
            "Ovo je automatska poruka. Za pitanja odgovorite na ovaj email." if sr
            else "This is an automated message. Reply to this email with any questions."
        )
        self.text.append(note)
        self.html.append(f"<p style=\"margin:0 0 12px;color:#777;font-size:12px\">{html_lib.escape(note)}</p>")


def _compose(build, subjects: Dict[str, str]) -> Tuple[str, str, str]:
    """Serbian, plus English under it when `bilingual_emails` is on."""
    langs = ["sr", "en"] if setting("bilingual_emails") else ["sr"]
    mails = []
    for lang in langs:
        mail = _Mail(lang)
        build(mail)
        mail.footer()
        mails.append(mail)
    subject = " / ".join(subjects[lang] for lang in langs)
    text = "\n\n----------\n\n".join("\n\n".join(m.text) for m in mails)
    body = "<hr style=\"border:0;border-top:1px solid #ddd;margin:24px 0\">".join("".join(m.html) for m in mails)
    html_out = (
        "<div style=\"font-family:Arial,Helvetica,sans-serif;font-size:14px;line-height:1.45;color:#222;max-width:600px\">"
        + body.replace("\n", "<br>") + "</div>"
    )
    return subject, text, html_out


def _send_billing_email_sync(
    to_email: str, sender: Tuple[str, str], reply_to: Optional[str], subject: str, text: str, html_body: str,
    qr_png: Optional[bytes], qr_name: str,
) -> None:
    msg = EmailMessage()
    msg["From"] = formataddr(sender)
    msg["To"] = to_email
    msg["Subject"] = subject
    if reply_to:
        msg["Reply-To"] = reply_to
    msg.set_content(text)
    msg.add_alternative(html_body, subtype="html")
    if qr_png:
        msg.get_payload()[1].add_related(qr_png, maintype="image", subtype="png", cid="<ips-qr>", filename=qr_name)
    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=20) as smtp:
        if SMTP_USE_TLS:
            smtp.starttls()
        smtp.login(SMTP_USER, SMTP_PASSWORD)
        smtp.send_message(msg)


async def _billing_recipients(client: dict) -> List[str]:
    """Active admins of the client + the accounting address; each address once."""
    out: List[str] = []
    async for user in db.users.find(
        {"client_id": client["id"], "role": Role.ADMIN.value, "active": {"$ne": False}}, {"_id": 0, "email": 1}
    ).sort("created_at", 1):
        email = normalize_email(user.get("email") or "")
        if email and email not in out:
            out.append(email)
    accounting = normalize_email(client.get("accounting_email") or "")
    if accounting and accounting not in out:
        out.append(accounting)
    return out


async def _send_billing_email(
    event: str, subject: str, text: str, html_body: str, *, client: Optional[dict] = None,
    recipients: Optional[List[str]] = None, charge_ids: Optional[List[str]] = None, qr_payload: Optional[str] = None,
    qr_name: str = "ips-qr.png", ignore_opt_out: bool = False,
) -> dict:
    """Sends (best effort) and logs one billing email; marks `notified.<event>` on the charges."""
    skipped = None
    if recipients is None:
        recipients = await _billing_recipients(client) if client else []
        if client and not ignore_opt_out and ((client.get("subscription") or {}).get("billing_emails") is False):
            skipped = "disabled"
    if not skipped and not recipients:
        skipped = "no_recipients"
    if not skipped and not _smtp_is_configured():
        skipped = "smtp_off"
    sent_to: List[str] = []
    errors: List[str] = []
    if not skipped:
        sender = (setting("billing_from_name") or "Easy Order", setting("billing_from_email") or SMTP_FROM)
        qr_png = _qr_png(qr_payload) if qr_payload else None
        for email in recipients:
            try:
                await asyncio.to_thread(
                    _send_billing_email_sync, email, sender, setting("billing_reply_to") or None, subject, text, html_body, qr_png, qr_name,
                )
                sent_to.append(email)
            except Exception as exc:
                logger.warning("Billing email %s to %s failed: %s", event, email, exc)
                errors.append(f"{email}: {exc}"[:300])
    log = {
        "id": str(uuid.uuid4()), "client_id": (client or {}).get("id"), "event": event, "subject": subject,
        "text": text, "recipients": recipients, "sent_to": sent_to, "sent": bool(sent_to), "skipped": skipped,
        "errors": errors, "charge_ids": charge_ids or [], "with_qr": bool(qr_payload), "created_at": now_iso(),
    }
    await db.billing_emails.insert_one({**log, "_id": log["id"]})
    if charge_ids:
        await db.payments.update_many({"id": {"$in": charge_ids}}, {"$set": {f"notified.{event}": log["created_at"]}})
    return log


async def _charge_email(client: dict, event: str, charges: List[dict], today: date) -> dict:
    """issued / reminder / warning: the charge(s), what is owed in total, payment details and the IPS QR."""
    me, payload = await _billing_for_client(client["id"])
    charge = charges[0]
    number = charge.get("number") or ""
    due = charge.get("due_date")
    blocks_on = _blocks_on(due).isoformat() if due else None
    this_amount = round(sum(float(c["amount"]) for c in charges), 2) if event == "issued" else None

    def build(m: _Mail):
        sr = m.lang == "sr"
        name = client["name"]
        if event == "issued":
            period = _period_text(charge, m.lang)
            m.p("Poštovani," if sr else "Hello,")
            if period:
                m.p(
                    f"izdat je predračun {number} za korišćenje aplikacije Easy Order ({name}) za period {period}." if sr
                    else f"pro-forma invoice {number} has been issued for using Easy Order ({name}) for {period}."
                )
            else:
                m.p(f"izdat je predračun {number} ({name})." if sr else f"pro-forma invoice {number} has been issued ({name}).")
            rows = [(_line_text(line, m.lang), _money(line["amount"], m.lang)) for c in charges for line in (c.get("lines") or [])]
            if not rows:
                rows = [(number, _money(this_amount, m.lang))]
            if len(charges) == 1 and charge.get("note") and abs(sum(line["amount"] for line in charge.get("lines") or []) - charge["amount"]) > 0.005:
                rows.append((charge["note"], _money(charge["amount"] - sum(line["amount"] for line in charge.get("lines") or []), m.lang)))
            previous = round(me.to_pay - this_amount, 2)
            if previous > 0.005:
                rows.append(("Dug iz prethodnog perioda" if sr else "Unpaid from earlier", _money(previous, m.lang)))
            elif previous < -0.005:
                rows.append(("Preplata iz prethodnog perioda" if sr else "Credit from earlier", _money(previous, m.lang)))
            m.table(rows, ("Ukupno za uplatu" if sr else "Total to pay", _money(me.to_pay, m.lang)))
            if me.to_pay > 0 and due:
                m.p((f"Rok plaćanja: {_day_text(due)}." if sr else f"Due date: {_day_text(due)}."), bold=True)
        elif event == "reminder":
            m.p("Poštovani," if sr else "Hello,")
            m.p(
                f"podsećamo da je rok za uplatu po predračunu {number} ({name}) {_day_text(due)}, a uplata još nije evidentirana." if sr
                else f"a reminder that payment for pro-forma invoice {number} ({name}) was due on {_day_text(due)} and has not been recorded yet."
            )
            m.p((f"Za uplatu: {_money(me.to_pay, m.lang)}" if sr else f"To pay: {_money(me.to_pay, m.lang)}"), bold=True)
        else:  # warning
            m.p("Poštovani," if sr else "Hello,")
            m.p(
                f"uplata po predračunu {number} ({name}) još nije evidentirana. Ako uplata ne stigne, pristup aplikaciji biće ograničen od {_day_text(blocks_on)}." if sr
                else f"payment for pro-forma invoice {number} ({name}) has not been recorded yet. Without it, access to the app will be limited from {_day_text(blocks_on)}."
            )
            m.p((f"Za uplatu: {_money(me.to_pay, m.lang)}" if sr else f"To pay: {_money(me.to_pay, m.lang)}"), bold=True)
        if me.to_pay > 0:
            m.details(me, with_qr=bool(payload))
        if event == "issued" and me.to_pay > 0 and blocks_on:
            m.p(
                f"Ako uplata ne stigne do {_day_text((date.fromisoformat(blocks_on) - timedelta(days=1)).isoformat())}, pristup aplikaciji može biti privremeno ograničen." if sr
                else f"If payment does not arrive by {_day_text((date.fromisoformat(blocks_on) - timedelta(days=1)).isoformat())}, access to the app may be limited."
            )
        if event != "issued":
            m.p("Ako ste već platili, zanemarite ovu poruku." if sr else "If you have already paid, please ignore this message.")
        m.link("Pregled zaduženja i uplata" if sr else "Charges and payments", f"{ADMIN_WEB_URL}/billing")

    period_sr, period_en = _period_text(charge, "sr"), _period_text(charge, "en")
    subjects = {
        "issued": {
            "sr": f"Predračun {number}" + (f" za {period_sr}" if period_sr else "") + " - Easy Order",
            "en": f"Pro-forma invoice {number}" + (f" for {period_en}" if period_en else ""),
        },
        "reminder": {"sr": f"Podsetnik za uplatu - predračun {number}", "en": f"Payment reminder - {number}"},
        "warning": {
            "sr": f"Uplata nije evidentirana - pristup se ograničava {_day_text(blocks_on)}",
            "en": f"Payment not recorded - access will be limited on {_day_text(blocks_on)}",
        },
    }[event]
    subject, text, html_body = _compose(build, subjects)
    return await _send_billing_email(
        event, subject, text, html_body, client=client, charge_ids=[c["id"] for c in charges],
        qr_payload=payload if me.to_pay > 0 else None, qr_name=f"ips-qr-{number or 'uplata'}.png",
    )


async def _notify_issued(client_id: str, charges: List[dict]) -> Optional[dict]:
    """After approval: one email per client with its new charges (free months are skipped)."""
    priced = [c for c in charges if float(c.get("amount") or 0) > 0]
    if not priced:
        return None
    client = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not client:
        return None
    try:
        return await _charge_email(client, "issued", priced, _today())
    except Exception as exc:  # never fail the approval because of an email
        logger.warning("Billing email (issued) for %s failed: %s", client_id, exc)
        return None


async def _notify_paid(client_id: str, payment: dict) -> Optional[dict]:
    client = await db.clients.find_one({"id": client_id}, {"_id": 0})
    if not client:
        return None
    me, _ = await _billing_for_client(client_id)

    def build(m: _Mail):
        sr = m.lang == "sr"
        m.p("Poštovani," if sr else "Hello,")
        m.p(
            f"primili smo uplatu od {_money(payment['amount'], m.lang)} ({_day_text(payment.get('paid_at'))}) za {client['name']}. Hvala!" if sr
            else f"we have received your payment of {_money(payment['amount'], m.lang)} ({_day_text(payment.get('paid_at'))}) for {client['name']}. Thank you!"
        )
        if me.to_pay > 0:
            m.p((f"Preostalo za uplatu: {_money(me.to_pay, m.lang)}" if sr else f"Still to pay: {_money(me.to_pay, m.lang)}"), bold=True)
        elif me.credit > 0:
            m.p(
                f"Preplata od {_money(me.credit, m.lang)} biće uračunata u sledeće zaduženje." if sr
                else f"The overpayment of {_money(me.credit, m.lang)} will be applied to the next charge."
            )
        else:
            m.p("Trenutno nemate dugovanja." if sr else "You have nothing left to pay.")
        m.link("Pregled zaduženja i uplata" if sr else "Charges and payments", f"{ADMIN_WEB_URL}/billing")

    subject, text, html_body = _compose(build, {"sr": "Potvrda uplate - Easy Order", "en": "Payment received"})
    return await _send_billing_email("paid", subject, text, html_body, client=client)


async def _notify_blocked(client: dict, state: dict) -> dict:
    me, payload = await _billing_for_client(client["id"])

    def build(m: _Mail):
        sr = m.lang == "sr"
        m.p("Poštovani," if sr else "Hello,")
        m.p(
            f"pristup aplikaciji Easy Order za {client['name']} je privremeno ograničen jer uplata nije evidentirana." if sr
            else f"access to Easy Order for {client['name']} has been limited because payment has not been recorded."
        )
        m.p(
            "Administrator i dalje može da se prijavi na portal i vidi zaduženja i podatke za uplatu. Pristup se vraća čim se uplata evidentira." if sr
            else "Your admin can still sign in to the portal to see the charges and payment details. Access returns as soon as the payment is recorded."
        )
        if me.to_pay > 0:
            m.p((f"Za uplatu: {_money(me.to_pay, m.lang)}" if sr else f"To pay: {_money(me.to_pay, m.lang)}"), bold=True)
            m.details(me, with_qr=bool(payload))
        m.link("Zaduženja" if sr else "Charges", f"{ADMIN_WEB_URL}/billing")

    subject, text, html_body = _compose(
        build, {"sr": "Pristup aplikaciji je ograničen - Easy Order", "en": "Access to the app is limited"}
    )
    return await _send_billing_email(
        "blocked", subject, text, html_body, client=client,
        qr_payload=payload if me.to_pay > 0 else None, qr_name="ips-qr.png",
    )


async def _notify_annual_expiry(client: dict, state: dict) -> dict:
    sub = client["subscription"]

    def build(m: _Mail):
        sr = m.lang == "sr"
        m.p("Poštovani," if sr else "Hello,")
        m.p(
            f"godišnja pretplata za {client['name']} ističe {_day_text(sub['ends_at'])} (za {state['days_left']} dana). Predračun za sledeću godinu stiže krajem meseca." if sr
            else f"the annual subscription for {client['name']} ends on {_day_text(sub['ends_at'])} (in {state['days_left']} days). The pro-forma invoice for the next year will follow at the end of the month."
        )
        m.link("Zaduženja" if sr else "Charges", f"{ADMIN_WEB_URL}/billing")

    subject, text, html_body = _compose(
        build, {"sr": "Godišnja pretplata uskoro ističe - Easy Order", "en": "Annual subscription ends soon"}
    )
    return await _send_billing_email("annual_expiry", subject, text, html_body, client=client, ignore_opt_out=True)


async def _notify_drafts_ready(drafts: List[dict]) -> Optional[dict]:
    """To every active superadmin: how many drafts the cron made and the total."""
    if not drafts:
        return None
    recipients = [
        normalize_email(u["email"]) async for u in db.users.find(
            {"role": Role.SUPERADMIN.value, "active": {"$ne": False}}, {"_id": 0, "email": 1}
        ) if u.get("email")
    ]
    total = round(sum(float(d["amount"]) for d in drafts), 2)
    month = date.fromisoformat(drafts[0]["period_start"])

    def build(m: _Mail):
        sr = m.lang == "sr"
        label = f"{_MONTHS[m.lang][month.month - 1]} {month.year}"
        m.p(
            f"Napravljeno je {len(drafts)} nacrta zaduženja za {label}, ukupno {_money(total, m.lang)}. Pregledaj ih i odobri." if sr
            else f"{len(drafts)} charge drafts were made for {label}, {_money(total, m.lang)} in total. Review and approve them."
        )
        m.link("Nacrti zaduženja" if sr else "Charge drafts", f"{ADMIN_WEB_URL}/billing-drafts")

    subject, text, html_body = _compose(
        build, {"sr": f"Nacrti zaduženja: {len(drafts)}", "en": f"Charge drafts: {len(drafts)}"}
    )
    return await _send_billing_email("drafts_ready", subject, text, html_body, recipients=sorted(set(recipients)))


async def _billing_notifications(today: date) -> Dict[str, List[str]]:
    """Daily: reminder on/after the due date, warning two days before the block
    (only when the client would really be blocked), blocked notice once per block.
    Each one once per charge (`notified`), one email per client and event."""
    out: Dict[str, List[str]] = {"reminder": [], "warning": [], "blocked": []}
    query = {"subscription.billing_cycle": {"$in": ["monthly", "annual"]}, "active": {"$ne": False}}
    async for client in db.clients.find(query, {"_id": 0}):
        try:
            await _client_notifications(client, today, out)
        except Exception as exc:  # one client's problem must not stop the rest of the cron
            logger.warning("Billing notifications for %s failed: %s", client.get("id"), exc)
    return out


async def _client_notifications(client: dict, today: date, out: Dict[str, List[str]]) -> None:
    sub = client["subscription"]
    state = subscription_state(sub, today)
    blocked = state["status"] == "locked" and state.get("lock_reason") in ("debt", "expired")
    if blocked:
        if not sub.get("block_notified"):
            await _notify_blocked(client, state)
            await db.clients.update_one({"id": client["id"]}, {"$set": {"subscription.block_notified": today.isoformat()}})
            out["blocked"].append(client["name"])
        return
    if sub.get("block_notified"):
        await db.clients.update_one({"id": client["id"]}, {"$set": {"subscription.block_notified": None}})
    if state["status"] == "locked" or not float(sub.get("balance_due") or 0):
        return
    charges, payments = await _ledger(client["id"])
    allocated, _, _ = _allocate(charges, payments)
    due_now: Dict[str, List[dict]] = {"reminder": [], "warning": []}
    for charge in allocated:
        if charge["open"] <= 0 or not charge.get("due_date"):
            continue
        due = date.fromisoformat(charge["due_date"])
        blocks_on = _blocks_on(charge["due_date"])
        if today < due or today >= blocks_on:
            continue
        notified = charge.get("notified") or {}
        if today >= blocks_on - timedelta(days=2):
            if not notified.get("warning") and subscription_state(sub, blocks_on)["status"] == "locked":
                due_now["warning"].append(charge)
        elif not notified.get("reminder"):
            due_now["reminder"].append(charge)
    for event in ("warning", "reminder"):
        if due_now[event]:
            await _charge_email(client, event, due_now[event], today)
            out[event].append(client["name"])
            return  # one email a day is enough


class BillingEmailLog(BaseModel):
    id: str
    client_id: Optional[str] = None
    event: str
    subject: str
    text: str = ""
    recipients: List[str] = Field(default_factory=list)
    sent_to: List[str] = Field(default_factory=list)
    sent: bool = False
    skipped: Optional[str] = None
    errors: List[str] = Field(default_factory=list)
    charge_ids: List[str] = Field(default_factory=list)
    with_qr: bool = False
    created_at: str


class BillingEmailSettings(BaseModel):
    billing_emails: bool
    accounting_email: Optional[str] = None
    recipients: List[str]
    log: List[BillingEmailLog]


class BillingEmailSettingsInput(BaseModel):
    billing_emails: Optional[bool] = None
    accounting_email: Optional[str] = Field(default=None, max_length=254)  # "" = remove


async def _email_settings_out(client_id: str) -> BillingEmailSettings:
    client = await _client_or_404(client_id)
    log = await db.billing_emails.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(50)
    return BillingEmailSettings(
        billing_emails=(client.get("subscription") or {}).get("billing_emails") is not False,
        accounting_email=client.get("accounting_email") or None,
        recipients=await _billing_recipients(client),
        log=[BillingEmailLog(**x) for x in log],
    )


@api_router.get("/billing/clients/{client_id}/emails", response_model=BillingEmailSettings)
async def client_billing_emails(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    return await _email_settings_out(client_id)


@api_router.put("/billing/clients/{client_id}/emails", response_model=BillingEmailSettings)
async def update_client_billing_emails(
    client_id: str, inp: BillingEmailSettingsInput, current_user: User = Depends(require_roles(Role.SUPERADMIN)),
):
    """Billing emails on/off for the client (the banners stay) and its accounting address."""
    client = await _client_or_404(client_id)
    changes: Dict[str, Any] = {}
    if inp.billing_emails is not None:
        if not client.get("subscription"):
            raise HTTPException(status_code=400, detail="Client has no subscription")
        changes["subscription.billing_emails"] = inp.billing_emails
    if inp.accounting_email is not None:
        value = normalize_email(inp.accounting_email)
        if value and not _EMAIL_RE.match(value):
            raise HTTPException(status_code=400, detail="accounting_email: Not a valid email address")
        changes["accounting_email"] = value or None
    if changes:
        await db.clients.update_one({"id": client_id}, {"$set": changes})
        await _audit(current_user, "billing.emails_updated", "client", client_id, client["name"], client_id, {
            k.split(".")[-1]: v for k, v in changes.items()
        })
    return await _email_settings_out(client_id)


# ---------------- Superadmin tools (PLAN_SUPERADMIN.md, phase D) ----------------
# --- client notes: internal, append-only ---
class ClientNote(BaseModel):
    id: str
    client_id: str
    text: str
    author_name: Optional[str] = None
    created_at: str


class ClientNoteInput(BaseModel):
    text: str = Field(min_length=1, max_length=2000)


@api_router.get("/clients/{client_id}/notes", response_model=List[ClientNote])
async def list_client_notes(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    await _client_or_404(client_id)
    docs = await db.client_notes.find({"client_id": client_id}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return [ClientNote(**d) for d in docs]


@api_router.post("/clients/{client_id}/notes", response_model=ClientNote)
async def add_client_note(
    client_id: str, inp: ClientNoteInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    await _client_or_404(client_id)
    doc = {
        "id": str(uuid.uuid4()), "client_id": client_id, "text": inp.text.strip(),
        "author_name": current_user.name, "created_at": now_iso(),
    }
    if not doc["text"]:
        raise HTTPException(status_code=400, detail="Note is empty")
    await db.client_notes.insert_one({**doc, "_id": doc["id"]})
    return ClientNote(**doc)


# --- activity of a client ---
class UserActivity(BaseModel):
    id: str
    name: str
    role: str
    active: bool = True
    last_login_at: Optional[str] = None
    last_seen_at: Optional[str] = None


class ClientActivity(BaseModel):
    last_activity_at: Optional[str] = None
    last_order_at: Optional[str] = None
    orders_30d: int = 0
    orders_total: int = 0
    users: List[UserActivity]


@api_router.get("/clients/{client_id}/activity", response_model=ClientActivity)
async def client_activity(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    await _client_or_404(client_id)
    users = await db.users.find({"client_id": client_id}, {"_id": 0}).sort("name", 1).to_list(None)
    cutoff = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
    orders_total = await db.orders.count_documents({"client_id": client_id})
    orders_30d = await db.orders.count_documents({"client_id": client_id, "created_at": {"$gte": cutoff}})
    last = await db.orders.find({"client_id": client_id}, {"_id": 0, "created_at": 1}).sort("created_at", -1).limit(1).to_list(1)
    last_order = last[0]["created_at"] if last else None
    seen = [u["last_seen_at"] for u in users if u.get("last_seen_at")]
    return ClientActivity(
        last_activity_at=max(seen + ([last_order] if last_order else [])) if (seen or last_order) else None,
        last_order_at=last_order,
        orders_30d=orders_30d,
        orders_total=orders_total,
        users=[UserActivity(**{**u, "role": str(getattr(u["role"], "value", u["role"]))}) for u in users],
    )


# --- system status ---
class SystemCron(BaseModel):
    configured: bool
    status: str  # ok | late | never | disabled
    last_run_at: Optional[str] = None
    live: Optional[bool] = None
    summary: Dict[str, int] = Field(default_factory=dict)


class SystemStatus(BaseModel):
    environment: str
    server_time: str
    smtp_configured: bool
    cloudinary_configured: bool
    auto_purge_enabled: bool
    app_version: Optional[str] = None
    cron: SystemCron


CRON_LATE_HOURS = 48


class SettingsInput(BaseModel):
    """Partial update: a field that isn't sent keeps its value; null resets it to the default."""
    draft_day: Optional[int] = Field(default=None, ge=1, le=28)
    due_day: Optional[int] = Field(default=None, ge=1, le=28)
    block_after_day: Optional[int] = Field(default=None, ge=1, le=28)
    first_month_free: Optional[bool] = None
    reminder_days: Optional[int] = Field(default=None, ge=1, le=120)
    purge_after_days: Optional[int] = Field(default=None, ge=7, le=3650)
    expiry_alert_days: Optional[List[int]] = None
    payment_instructions_sr: Optional[str] = Field(default=None, max_length=2000)
    payment_instructions_en: Optional[str] = Field(default=None, max_length=2000)
    ips_account: Optional[str] = Field(default=None, max_length=30)
    ips_recipient_name: Optional[str] = Field(default=None, max_length=70)
    ips_payment_code: Optional[str] = Field(default=None, max_length=3)
    ips_purpose: Optional[str] = Field(default=None, max_length=35)
    ips_reference_from_number: Optional[bool] = None
    issuer_name: Optional[str] = Field(default=None, max_length=120)
    issuer_address: Optional[str] = Field(default=None, max_length=200)
    issuer_pib: Optional[str] = Field(default=None, max_length=20)
    issuer_account: Optional[str] = Field(default=None, max_length=30)
    issuer_note: Optional[str] = Field(default=None, max_length=300)
    billing_from_email: Optional[str] = Field(default=None, max_length=254)
    billing_from_name: Optional[str] = Field(default=None, max_length=80)
    billing_reply_to: Optional[str] = Field(default=None, max_length=254)
    bilingual_emails: Optional[bool] = None


class SettingsOut(BaseModel):
    values: Dict[str, Any]  # what is in effect
    defaults: Dict[str, Any]  # what applies when a value is reset (env, else code)


_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def normalize_bank_account(value: str) -> str:
    """Serbian current account: 3-13-2 digits (dashes optional, the middle part may
    be shortened) -> 18 digits; the control number must be valid (ISO 7064 mod 97)."""
    raw = value.strip().replace(" ", "")
    if "-" in raw:
        parts = raw.split("-")
        if len(parts) != 3 or not all(p.isdigit() for p in parts) or len(parts[0]) != 3 or len(parts[2]) != 2 or not 1 <= len(parts[1]) <= 13:
            raise ValueError("Account must look like 160-0000000000000-00")
        raw = parts[0] + parts[1].zfill(13) + parts[2]
    if not (raw.isdigit() and len(raw) == 18):
        raise ValueError("Account must have 18 digits")
    if int(raw) % 97 != 1:
        raise ValueError("Account control number is not valid")
    return raw


def _validated_settings(inp: SettingsInput) -> Dict[str, Any]:
    changes: Dict[str, Any] = {}
    for field in inp.model_fields_set:
        value = getattr(inp, field)
        if isinstance(value, str):
            value = value.strip()
            if value == "":
                value = None  # empty text = back to the default
        if value is None:
            changes[field] = None
            continue
        try:
            if field in ("ips_account", "issuer_account"):
                value = normalize_bank_account(value)
            elif field == "ips_payment_code" and not (value.isdigit() and len(value) == 3):
                raise ValueError("Payment code must have 3 digits")
            elif field == "issuer_pib" and not (value.isdigit() and len(value) == 9):
                raise ValueError("PIB must have 9 digits")
            elif field in ("billing_from_email", "billing_reply_to"):
                value = value.lower()
                if not _EMAIL_RE.match(value):
                    raise ValueError("Not a valid email address")
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"{field}: {exc}")
        if field == "expiry_alert_days":
            value = _validated_alert_days(value)
        changes[field] = value
    merged = {**effective_settings(), **{k: (SETTINGS_DEFAULTS[k] if v is None else v) for k, v in changes.items()}}
    if merged["block_after_day"] < merged["due_day"]:
        raise HTTPException(status_code=400, detail="block_after_day must not be before due_day")
    return changes


@api_router.get("/settings", response_model=SettingsOut)
async def get_settings(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    await refresh_settings(force=True)
    return SettingsOut(values=effective_settings(), defaults=SETTINGS_DEFAULTS)


@api_router.put("/settings", response_model=SettingsOut)
async def update_settings(inp: SettingsInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    await refresh_settings(force=True)
    changes = _validated_settings(inp)
    if changes:
        to_set = {k: v for k, v in changes.items() if v is not None}
        to_unset = {k: "" for k, v in changes.items() if v is None}
        update: Dict[str, Any] = {"$set": {**to_set, "updated_at": now_iso(), "updated_by": current_user.name}}
        if to_unset:
            update["$unset"] = to_unset
        await db.app_settings.update_one({"_id": "global"}, update, upsert=True)
        await _audit(current_user, "settings.update", "settings", "global", None, None, {"fields": sorted(changes)})
        await refresh_settings(force=True)
    return SettingsOut(values=effective_settings(), defaults=SETTINGS_DEFAULTS)


@api_router.get("/superadmin/system", response_model=SystemStatus)
async def system_status(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Is the daily job running, is mail/image storage set up. Nothing secret is returned."""
    run = await db.system_runs.find_one({"_id": "subscriptions_cron"})
    if not CRON_SECRET:
        cron_status = "disabled"
    elif not run:
        cron_status = "never"
    else:
        age = datetime.now(timezone.utc) - datetime.fromisoformat(run["finished_at"])
        cron_status = "late" if age > timedelta(hours=CRON_LATE_HOURS) else "ok"
    version = None
    try:
        version = json.loads(APP_UPDATE_FILE.read_text(encoding="utf-8")).get("version")
    except (OSError, ValueError):
        pass
    return SystemStatus(
        environment="production" if IS_PRODUCTION else "development",
        server_time=now_iso(),
        smtp_configured=_smtp_is_configured(),
        cloudinary_configured=cloudinary_available(),
        auto_purge_enabled=AUTO_PURGE_ENABLED,
        app_version=version,
        cron=SystemCron(
            configured=bool(CRON_SECRET),
            status=cron_status,
            last_run_at=(run or {}).get("finished_at"),
            live=(run or {}).get("live"),
            summary=(run or {}).get("summary") or {},
        ),
    )


# --- audit log of superadmin actions ---
class AuditEntry(BaseModel):
    id: str
    created_at: str
    actor_name: Optional[str] = None
    action: str
    target_type: Optional[str] = None
    target_name: Optional[str] = None
    client_id: Optional[str] = None
    data: Optional[Dict[str, Any]] = None


@api_router.get("/audit", response_model=List[AuditEntry])
async def list_audit(
    limit: int = Query(50, ge=1, le=200),
    skip: int = Query(0, ge=0),
    action: Optional[str] = Query(None, description="prefix, e.g. 'client.' or 'subscription.payment'"),
    client_id: Optional[str] = None,
    current_user: User = Depends(require_roles(Role.SUPERADMIN)),
):
    query: Dict[str, Any] = {}
    if action:
        query["action"] = {"$regex": f"^{re.escape(action)}"}
    if client_id:
        query["client_id"] = client_id
    docs = await db.audit_log.find(query, {"_id": 0}).sort("created_at", -1).skip(skip).limit(limit).to_list(limit)
    return [AuditEntry(**d) for d in docs]


# --- announcements to clients ---
class AnnouncementInput(BaseModel):
    message_sr: str = Field(default="", max_length=500)
    message_en: str = Field(default="", max_length=500)
    level: Literal["info", "warning"] = "info"
    starts_at: Optional[str] = None  # YYYY-MM-DD; empty = from now on
    ends_at: Optional[str] = None  # YYYY-MM-DD, inclusive; empty = until switched off
    client_ids: List[str] = Field(default_factory=list)  # empty = every client
    active: bool = True


class Announcement(AnnouncementInput):
    id: str
    status: str = "live"  # computed: live | scheduled | ended | off
    created_by: Optional[str] = None
    created_at: str


def _announcement_status(doc: dict, today: date) -> str:
    if not doc.get("active", True):
        return "off"
    if doc.get("starts_at") and date.fromisoformat(doc["starts_at"]) > today:
        return "scheduled"
    if doc.get("ends_at") and date.fromisoformat(doc["ends_at"]) < today:
        return "ended"
    return "live"


async def _validated_announcement(inp: AnnouncementInput) -> dict:
    if not inp.message_sr.strip() and not inp.message_en.strip():
        raise HTTPException(status_code=400, detail="Write the message in at least one language")
    starts = _valid_day(inp.starts_at, "starts_at") if inp.starts_at else None
    ends = _valid_day(inp.ends_at, "ends_at") if inp.ends_at else None
    if starts and ends and ends < starts:
        raise HTTPException(status_code=400, detail="ends_at is before starts_at")
    ids = list(dict.fromkeys(inp.client_ids))
    if ids and await db.clients.count_documents({"id": {"$in": ids}}) != len(ids):
        raise HTTPException(status_code=400, detail="Unknown client in client_ids")
    return {
        "message_sr": inp.message_sr.strip(), "message_en": inp.message_en.strip(), "level": inp.level,
        "starts_at": starts, "ends_at": ends, "client_ids": ids, "active": inp.active,
    }


async def _announcements_for(client_id: Optional[str]) -> List[dict]:
    """Live announcements a client's users see (a missing language falls back to the other)."""
    if not client_id:
        return []
    today = _today()
    out = []
    async for doc in db.announcements.find({"active": {"$ne": False}}, {"_id": 0}).sort("created_at", -1):
        if _announcement_status(doc, today) != "live":
            continue
        if doc.get("client_ids") and client_id not in doc["client_ids"]:
            continue
        sr, en = doc.get("message_sr") or doc.get("message_en"), doc.get("message_en") or doc.get("message_sr")
        out.append({"id": doc["id"], "level": doc.get("level", "info"), "message_sr": sr, "message_en": en})
    return out


@api_router.get("/announcements", response_model=List[Announcement])
async def list_announcements(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    today = _today()
    docs = await db.announcements.find({}, {"_id": 0}).sort("created_at", -1).to_list(200)
    return [Announcement(**{**d, "status": _announcement_status(d, today)}) for d in docs]


@api_router.post("/announcements", response_model=Announcement)
async def create_announcement(inp: AnnouncementInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    doc = {
        "id": str(uuid.uuid4()), **await _validated_announcement(inp),
        "created_by": current_user.name, "created_at": now_iso(),
    }
    await db.announcements.insert_one({**doc, "_id": doc["id"]})
    await _audit(current_user, "announcement.create", "announcement", doc["id"], (doc["message_sr"] or doc["message_en"])[:60])
    return Announcement(**{**doc, "status": _announcement_status(doc, _today())})


@api_router.put("/announcements/{announcement_id}", response_model=Announcement)
async def update_announcement(
    announcement_id: str, inp: AnnouncementInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))
):
    existing = await db.announcements.find_one({"id": announcement_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Announcement not found")
    doc = {**existing, **await _validated_announcement(inp)}
    await db.announcements.replace_one({"id": announcement_id}, {**doc, "_id": announcement_id})
    await _audit(current_user, "announcement.update", "announcement", announcement_id, (doc["message_sr"] or doc["message_en"])[:60])
    return Announcement(**{**doc, "status": _announcement_status(doc, _today())})


@api_router.delete("/announcements/{announcement_id}")
async def delete_announcement(announcement_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    existing = await db.announcements.find_one({"id": announcement_id}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Announcement not found")
    await db.announcements.delete_one({"id": announcement_id})
    await _audit(current_user, "announcement.delete", "announcement", announcement_id, (existing.get("message_sr") or existing.get("message_en") or "")[:60])
    return {"ok": True}


# ---------------- Superadmin overview (PLAN_SUPERADMIN.md, phase A) ----------------
class OverviewClients(BaseModel):
    total: int = 0
    active: int = 0
    inactive: int = 0
    locked: int = 0  # active clients whose subscription is locked


class OverviewUsers(BaseModel):
    total: int = 0
    by_role: Dict[str, int] = Field(default_factory=dict)


class OverviewSubscriptions(BaseModel):
    active: int = 0
    ending_30d: int = 0  # active and ending within 30 days (also counted in active)
    grace: int = 0
    locked: int = 0
    none: int = 0


class AttentionItem(BaseModel):
    client_id: str
    client_name: str
    reason: str  # locked | blocked_debt | grace | purge_soon | ending_soon | payment_overdue | setup_unpaid | drafts_pending | debt_repeated
    date: Optional[str] = None  # the date that matters for the reason
    plan_name: Optional[str] = None


class OverviewPayments(BaseModel):
    currency: str = PAYMENT_CURRENCY
    received_month: float = 0
    received_year: float = 0
    expected_total: float = 0  # open debts (expected payments)
    overdue_total: float = 0
    overdue_count: int = 0


class SuperadminOverview(BaseModel):
    clients: OverviewClients
    users: OverviewUsers
    subscriptions: OverviewSubscriptions
    payments: OverviewPayments
    attention: List[AttentionItem]


_ATTENTION_ORDER = {
    "locked": 0, "blocked_debt": 0, "purge_soon": 1, "payment_overdue": 2, "debt_repeated": 2, "grace": 3,
    "drafts_pending": 4, "setup_unpaid": 4, "month_not_charged": 5, "ending_soon": 6,
}
# From settings.draft_day a client that is billed monthly should have next month charged.


@api_router.get("/superadmin/overview", response_model=SuperadminOverview)
async def superadmin_overview(current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    """Numbers for the superadmin dashboard. Deactivated clients are counted as
    such and left out of the subscription figures and the attention list."""
    today = _today()
    clients = OverviewClients()
    subs = OverviewSubscriptions()
    attention: List[AttentionItem] = []
    active_ids = set()
    active_clients: Dict[str, dict] = {}
    async for client in db.clients.find({}, {"_id": 0}):
        clients.total += 1
        if not client.get("active", True):
            clients.inactive += 1
            continue
        clients.active += 1
        active_ids.add(client["id"])
        active_clients[client["id"]] = client
        sub = client.get("subscription")
        state = subscription_state(sub, today)
        status = state["status"]
        setattr(subs, status, getattr(subs, status) + 1)
        item = {"client_id": client["id"], "client_name": client["name"], "plan_name": (sub or {}).get("plan_name")}
        if status == "locked":
            clients.locked += 1
            purge_at = date.fromisoformat(state["purge_at"]) if state.get("purge_at") else None
            if state.get("lock_reason") == "debt":
                attention.append(AttentionItem(**item, reason="blocked_debt", date=state["locked_since"]))
            elif purge_at and not sub.get("purged_at") and not sub.get("purge_paused") and (purge_at - today).days <= setting("reminder_days"):
                attention.append(AttentionItem(**item, reason="purge_soon", date=state["purge_at"]))
            else:
                attention.append(AttentionItem(**item, reason="locked", date=state["locked_since"]))
        elif status == "grace":
            attention.append(AttentionItem(**item, reason="grace", date=state["grace_ends_at"]))
        elif status == "active":
            if state["days_left"] <= 30:
                subs.ending_30d += 1
            if state["days_left"] <= setting("reminder_days") and (sub or {}).get("billing_cycle") != "monthly":
                attention.append(AttentionItem(**item, reason="ending_soon", date=state["ends_at"]))
    money = await _payment_figures(today)
    for client_id, oldest_due in money["overdue_by_client"].items():
        if client_id in active_clients:
            c = active_clients[client_id]
            attention.append(AttentionItem(
                client_id=client_id, client_name=c["name"], reason="payment_overdue", date=oldest_due,
                plan_name=(c.get("subscription") or {}).get("plan_name"),
            ))
    overdue_ids = set(money["overdue_by_client"])
    setup_docs = await db.payments.find({"kind": "setup", "status": "expected"}, {"_id": 0}).to_list(None)
    open_by_id = await _ledger_open_map({p["client_id"] for p in setup_docs if p.get("ledger")})
    for pay in setup_docs:
        if pay.get("ledger") and open_by_id.get(pay["id"], 0) <= 0:
            continue  # paid (calendar billing: charges stay "expected", what's open counts)
        c = active_clients.get(pay["client_id"])
        if c and pay["client_id"] not in overdue_ids:  # a late setup debt already shows as payment_overdue
            attention.append(AttentionItem(
                client_id=c["id"], client_name=c["name"], reason="setup_unpaid", date=pay.get("due_date"),
                plan_name=(c.get("subscription") or {}).get("plan_name"),
            ))
    attention.extend(await _billing_attention(today, active_clients))
    by_role: Dict[str, int] = {}
    async for row in db.users.aggregate([
        {"$match": {"client_id": {"$in": list(active_ids)}, "active": {"$ne": False}}},
        {"$group": {"_id": "$role", "n": {"$sum": 1}}},
    ]):
        by_role[row["_id"]] = row["n"]
    attention.sort(key=lambda a: (_ATTENTION_ORDER[a.reason], a.date or "", a.client_name))
    return SuperadminOverview(
        clients=clients,
        users=OverviewUsers(total=sum(by_role.values()), by_role=by_role),
        subscriptions=subs,
        payments=OverviewPayments(**{k: v for k, v in money.items() if k != "overdue_by_client"}),
        attention=attention[:50],
    )


# ---------------- Users ----------------
@api_router.get("/users", response_model=List[UserRow])
async def list_users(client_id: Optional[str] = None, current_user: User = Depends(require_manager)):
    if current_user.role == Role.SUPERADMIN:
        query = {"client_id": client_id} if client_id else {}
    else:
        query = {"client_id": current_user.client_id}
    docs = await db.users.find(query, {"_id": 0}).sort("name", 1).to_list(None)
    names: Dict[str, str] = {}
    if current_user.role == Role.SUPERADMIN:
        ids = list({d["client_id"] for d in docs if d.get("client_id")})
        if ids:
            names = {c["id"]: c["name"] async for c in db.clients.find({"id": {"$in": ids}}, {"_id": 0, "id": 1, "name": 1})}
    return [UserRow(**{**v, "client_name": names.get(v.get("client_id"))}) for v in docs]


@api_router.post("/users", response_model=UserInviteResponse)
async def create_user(inp: UserInput, current_user: User = Depends(require_manager)):
    if await find_user_by_email(inp.email):
        raise HTTPException(status_code=400, detail="Email already in use")

    if current_user.role == Role.ADMIN:
        # Admins may only create operators/warehouse for their own client -
        # the payload client_id is ignored, never trusted.
        if inp.role not in ADMIN_MANAGEABLE_ROLES:
            raise HTTPException(status_code=403, detail="Admins cannot assign this role")
        role = inp.role
        client_id = current_user.client_id
    else:  # SUPERADMIN
        role = inp.role
        if role == Role.SUPERADMIN:
            client_id = None
        else:
            if not inp.client_id or not await db.clients.find_one({"id": inp.client_id}, {"_id": 0}):
                raise HTTPException(status_code=400, detail="Valid client_id is required for this role")
            client_id = inp.client_id

    if role == Role.WAREHOUSE and not await client_has_module(client_id, Module.WAREHOUSE):
        raise HTTPException(status_code=403, detail=f"Module not enabled: {Module.WAREHOUSE.value}")
    await _check_seat_limit(client_id, role)
    obj = User(email=normalize_email(inp.email), name=inp.name, phone=inp.phone, role=role, client_id=client_id)
    invite = await _invite_user(obj)
    if current_user.role == Role.SUPERADMIN:
        await _audit(current_user, "user.create", "user", obj.id, obj.email, client_id, {"role": role.value})
    return UserInviteResponse(**obj.dict(), **invite.dict())


@api_router.put("/users/{user_id}", response_model=User)
async def update_user(user_id: str, inp: UserUpdateInput, current_user: User = Depends(require_manager)):
    target = await get_scoped_or_404("users", user_id, current_user)
    if current_user.role == Role.ADMIN:
        is_self = target["id"] == current_user.id
        if not is_self and target.get("role") not in ADMIN_MANAGEABLE_ROLES:
            # Otherwise one admin could reset another admin's password and
            # take over their account, or deactivate them.
            raise HTTPException(status_code=403, detail="Admins can only manage operators and warehouse users")
        # Compare against the stored values: the admin-web edit form always
        # sends `active` (and may send the current role), so only an actual
        # change is rejected - re-sending the same value is fine.
        role_changes = inp.role is not None and inp.role != target.get("role")
        active_changes = inp.active is not None and inp.active != target.get("active", True)
        if is_self and (role_changes or active_changes):
            raise HTTPException(status_code=403, detail="Cannot change your own role or status")
        if inp.role is not None and inp.role not in ADMIN_MANAGEABLE_ROLES:
            raise HTTPException(status_code=403, detail="Admins cannot assign this role")

    updated = {**target}
    if inp.name is not None:
        updated["name"] = inp.name
    if inp.phone is not None:
        updated["phone"] = inp.phone
    if inp.active is not None:
        updated["active"] = inp.active
    role_changed = inp.role is not None and inp.role != target.get("role")
    if role_changed and inp.role == Role.WAREHOUSE and not await client_has_module(
        target.get("client_id"), Module.WAREHOUSE
    ):
        raise HTTPException(status_code=403, detail=f"Module not enabled: {Module.WAREHOUSE.value}")
    # Becoming (or staying) an active account of a role that is full is refused.
    becomes_active = inp.active is True and target.get("active", True) is False
    if (role_changed and updated.get("active", True)) or becomes_active:
        await _check_seat_limit(target.get("client_id"), inp.role or target.get("role"), exclude_user_id=user_id)
    if inp.role is not None:
        updated["role"] = inp.role
    if inp.password:
        ensure_strong_password(inp.password)
        updated["password_hash"] = hash_password(inp.password)
        # Set by someone else (admin/superadmin) = temporary, the user must
        # replace it. Changing your own password here doesn't force that.
        updated["must_change_password"] = user_id != current_user.id
    if inp.password or role_changed:
        # Old tokens carry the old role/password - log them out everywhere.
        updated["token_version"] = target.get("token_version", 0) + 1
    await db.users.replace_one({"id": user_id}, {**updated, "_id": user_id})
    if current_user.role == Role.SUPERADMIN:
        changes = {k: v for k, v in {"role": inp.role.value if role_changed else None, "active": inp.active, "password_reset": True if inp.password else None}.items() if v is not None}
        await _audit(current_user, "user.update", "user", user_id, target.get("email"), target.get("client_id"), changes or None)
    return User(**updated)


@api_router.delete("/users/{user_id}")
async def delete_user(user_id: str, current_user: User = Depends(require_manager)):
    target = await get_scoped_or_404("users", user_id, current_user)
    if user_id == current_user.id:
        raise HTTPException(status_code=400, detail="Cannot delete your own account")
    if current_user.role == Role.ADMIN and target.get("role") not in ADMIN_MANAGEABLE_ROLES:
        raise HTTPException(status_code=403, detail="Admins can only manage operators and warehouse users")
    await db.users.delete_one({"id": user_id})
    if current_user.role == Role.SUPERADMIN:
        await _audit(current_user, "user.delete", "user", user_id, target.get("email"), target.get("client_id"))
    return {"ok": True}


# ---------------- Customers ----------------
@api_router.get("/customers", response_model=List[Customer])
async def list_customers(current_user: User = Depends(get_current_user)):
    docs = await db.customers.find(_scope_query(current_user), {"_id": 0}).sort("name", 1).to_list(None)
    return [Customer(**v) for v in docs]


@api_router.post("/customers", response_model=Customer)
async def create_customer(inp: CustomerInput, current_user: User = Depends(require_manager)):
    client_id = await resolve_write_client_id(current_user, inp.client_id)
    obj = Customer(**inp.dict(exclude={"client_id"}), client_id=client_id)
    await db.customers.insert_one({**obj.dict(), "_id": obj.id})
    return obj


@api_router.put("/customers/{customer_id}", response_model=Customer)
async def update_customer(customer_id: str, inp: CustomerInput, current_user: User = Depends(require_manager)):
    existing = await get_scoped_or_404("customers", customer_id, current_user)
    updated = {**existing, **inp.dict(exclude={"client_id"})}
    await db.customers.replace_one({"id": customer_id}, {**updated, "_id": customer_id})
    return Customer(**updated)


# ---------------- Customer import from Excel ----------------
_CUSTOMER_IMPORT_HEADERS = {
    "name": "name", "naziv": "name", "nazivkupca": "name", "kupac": "name", "ime": "name",
    "address": "address", "adresa": "address",
    "email": "email", "eposta": "email", "mail": "email",
    "phone": "phone", "telefon": "phone", "tel": "phone", "mobilni": "phone", "brojtelefona": "phone",
    "pib": "pib", "pibbroj": "pib", "taxid": "pib",
}
_CUSTOMER_EXAMPLE_ROWS = [
    ("Maxi Market d.o.o.", "Bulevar oslobođenja 1, Novi Sad", "nabavka@maxi.example", "+381 21 555 111", "101234567"),
    ("Delikates Prodavnica", "Knez Mihailova 10, Beograd", "info@delikates.example", "+381 11 222 333", "107654321"),
    ("Mini Market Zora", "Kralja Petra 5, Niš", "zora@example.com", "018 444 555", "102345678"),
    ("Pekara Klas", "Cara Dušana 22, Kragujevac", None, "034 111 222", "103456789"),
    ("Restoran Lipa", "Savska 3, Beograd", "lipa@example.com", None, None),
]


def _cell_text(value: Any) -> str:
    """Excel hands numbers back as floats; phone/PIB cells must stay whole."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


@api_router.get("/customers/import-template")
async def customer_import_template(current_user: User = Depends(require_manager)):
    from openpyxl import Workbook
    from fastapi.responses import Response
    import io

    wb = Workbook()
    ws = wb.active
    ws.title = "Kupci"
    ws.append(["Naziv", "Adresa", "Email", "Telefon", "PIB"])
    for row in _CUSTOMER_EXAMPLE_ROWS:
        ws.append(list(row))
    for col in ("D", "E"):  # phone and PIB as text keep leading zeros / plus
        for cell in ws[col][1:]:
            cell.number_format = "@"
    for col, width in zip("ABCDE", (28, 36, 28, 18, 14)):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="uvoz-kupaca.xlsx"'},
    )


@api_router.post("/customers/import")
async def import_customers(
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    client_id: Optional[str] = Query(None),
    current_user: User = Depends(require_manager),
):
    """Excel import of customers. Matched by PIB, else by exact name; empty
    cells keep the stored value. dry_run (default) only reports."""
    target_client = await resolve_write_client_id(current_user, client_id)
    data = await file.read(IMPORT_MAX_BYTES + 1)
    if len(data) > IMPORT_MAX_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 4 MB)")
    try:
        from openpyxl import load_workbook
        import io

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        raw_rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the file - upload an .xlsx workbook")
    if not raw_rows:
        raise HTTPException(status_code=400, detail="The file is empty")
    columns: Dict[int, str] = {}
    for idx, head in enumerate(raw_rows[0]):
        field = _CUSTOMER_IMPORT_HEADERS.get(_import_header_key(head))
        if field and field not in columns.values():
            columns[idx] = field
    if "name" not in columns.values() and "pib" not in columns.values():
        raise HTTPException(status_code=400, detail="Missing a 'Naziv' (name) or 'PIB' column")
    body = [(i + 2, r) for i, r in enumerate(raw_rows[1:]) if any(c not in (None, "") for c in r)]
    if len(body) > IMPORT_MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Too many rows (max {IMPORT_MAX_ROWS})")

    existing = await db.customers.find({"client_id": target_client}, {"_id": 0}).to_list(None)
    by_pib = {c["pib"]: c for c in existing if c.get("pib")}
    by_name = {c["name"].strip().lower(): c for c in existing if c.get("name")}
    seen_pibs: Dict[str, int] = {}
    seen_ids: Dict[str, int] = {}
    rows_out = []
    for row_no, row in body:
        values = {field: _cell_text(row[idx] if idx < len(row) else None) for idx, field in columns.items()}
        errors: List[str] = []
        name = values.get("name", "")
        pib = re.sub(r"\s+", "", values.get("pib", ""))
        email = values.get("email", "").lower()
        if pib and not re.fullmatch(r"\d{1,15}", pib):
            errors.append("PIB: digits only")
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            errors.append("Email: invalid address")

        customer = by_pib.get(pib) if pib else None
        if customer is None and name:
            customer = by_name.get(name.lower())
        if customer is not None and pib and customer.get("pib") not in (None, "", pib):
            errors.append("Customer already has a different PIB")
        if pib and pib in seen_pibs:
            errors.append(f"PIB repeated (row {seen_pibs[pib]})")
        if customer is not None and customer["id"] in seen_ids:
            errors.append(f"Same customer repeated (row {seen_ids[customer['id']]})")
        if customer is None and not name:
            errors.append("Name is required for a new customer")
        if pib:
            seen_pibs.setdefault(pib, row_no)
        if customer is not None:
            seen_ids.setdefault(customer["id"], row_no)
        entry: Dict[str, Any] = {
            "row": row_no,
            "name": name or (customer or {}).get("name", ""),
            "pib": pib or None,
            "action": "error" if errors else ("update" if customer else "create"),
            "errors": errors,
        }
        if not errors:
            vals = {"address": values.get("address", ""), "phone": values.get("phone", "")}
            vals = {k: v for k, v in vals.items() if v}
            if email:
                vals["email"] = email
            if pib:
                vals["pib"] = pib
            if name:
                vals["name"] = name
            entry["_customer"] = customer
            entry["_values"] = vals
        rows_out.append(entry)

    if not dry_run:
        for entry in rows_out:
            if entry["action"] == "error":
                continue
            customer = entry["_customer"]
            if customer is None:
                obj = Customer(client_id=target_client, **entry["_values"])
                await db.customers.insert_one({**obj.dict(), "_id": obj.id})
            elif entry["_values"]:
                await db.customers.update_one({"id": customer["id"]}, {"$set": entry["_values"]})
    for entry in rows_out:
        entry.pop("_customer", None)
        entry.pop("_values", None)
    summary = {k: sum(1 for r in rows_out if r["action"] == k) for k in ("create", "update", "error")}
    return {"dry_run": dry_run, "summary": summary, "rows": rows_out}


@api_router.delete("/customers/{customer_id}")
async def delete_customer(customer_id: str, current_user: User = Depends(require_manager)):
    await get_scoped_or_404("customers", customer_id, current_user)
    await db.customers.delete_one({"id": customer_id})
    return {"ok": True}


# ---------------- Products ----------------
async def _reserved_by_product(scope: dict) -> Dict[str, int]:
    """Pieces committed to orders that are not shipped yet (new + in_progress)."""
    reserved: Dict[str, int] = {}
    open_orders = db.orders.find(
        {**scope, "status": {"$in": [OrderStatus.NEW.value, OrderStatus.IN_PROGRESS.value]}},
        {"_id": 0, "items.product_id": 1, "items.ordered_qty": 1},
    )
    async for order in open_orders:
        for item in order.get("items", []):
            reserved[item["product_id"]] = reserved.get(item["product_id"], 0) + (item.get("ordered_qty") or 0)
    return reserved


def _product_out(doc: dict, reserved: int, expiry: Optional[dict], user: User) -> ProductOut:
    """Stock fields for one product. Expired pieces don't count as available.
    Sales reps get no expiry information at all (only the lower availability)."""
    # A client without a module sees none of its data (it stays stored).
    has_stock = has_module(user, Module.STOCK)
    has_expiry = has_module(user, Module.EXPIRY)
    if not has_module(user, Module.WAREHOUSE):
        doc = {**doc, "barcode": None, "package_barcode": None}
    if not has_stock:
        doc = {**doc, "stock_qty": None}
        reserved = 0
    stock = doc.get("stock_qty")
    expired = (expiry or {}).get("expired_qty", 0) if has_expiry else 0
    out = ProductOut(
        **{**doc, "track_expiry": bool(doc.get("track_expiry")) and has_expiry},
        reserved_qty=reserved,
        available_qty=None if stock is None else stock - reserved - expired,
    )
    if user.role != Role.OPERATOR and has_expiry:
        out.expired_qty = expired
        out.next_expiry = (expiry or {}).get("next_expiry")
    else:
        out.track_expiry = False
    return out


@api_router.get("/products", response_model=List[ProductOut])
async def list_products(current_user: User = Depends(get_current_user)):
    scope = _scope_query(current_user)
    # Sales reps only see what can be ordered; managers and the warehouse also
    # see delisted/draft products (to edit, activate or receive stock).
    visible = {**scope, "active": {"$ne": False}} if current_user.role == Role.OPERATOR else scope
    docs = await db.products.find(visible, {"_id": 0}).sort("created_at", 1).to_list(None)
    reserved = await _reserved_by_product(scope)
    expiry = await _expiry_info(scope)
    return [_product_out(v, reserved.get(v["id"], 0), expiry.get(v["id"]), current_user) for v in docs]


@api_router.post("/products", response_model=Product)
async def create_product(inp: ProductInput, current_user: User = Depends(require_manager)):
    client_id = await resolve_write_client_id(current_user, inp.client_id)
    payload = inp.dict(exclude={"client_id"})
    payload["discount"] = max(0.0, min(100.0, float(payload.get("discount") or 0)))
    payload["discounts"] = normalize_discounts(inp.discounts, payload["discount"])
    payload["additional_discounts"] = normalize_additional_discounts(inp.additional_discounts)
    # Fields of modules the client doesn't have are ignored, not rejected: the
    # forms of such a client simply don't show them.
    if await client_has_module(client_id, Module.WAREHOUSE):
        payload["barcode"] = await _checked_barcode(client_id, inp.barcode)
        payload["package_barcode"] = await _checked_barcode(client_id, inp.package_barcode, other=payload["barcode"])
        _require_package_size(payload["package_barcode"], payload.get("pieces_per_package"))
    else:
        payload["barcode"] = payload["package_barcode"] = None
    payload["active"] = True if inp.active is None else inp.active
    payload["track_expiry"] = bool(inp.track_expiry) and await client_has_module(client_id, Module.EXPIRY)
    unit = await _unit_or_400(inp.unit or "kom")
    _require_count_unit_for_packaging(unit, payload.get("pieces_per_package"), payload.get("boxes_per_transport"), payload.get("package_barcode"))
    payload["unit"] = unit["code"]
    obj = Product(**payload, client_id=client_id)
    await db.products.insert_one({**obj.dict(), "_id": obj.id})
    return obj


@api_router.put("/products/{product_id}", response_model=Product)
async def update_product(product_id: str, inp: ProductInput, current_user: User = Depends(require_manager)):
    existing = await get_scoped_or_404("products", product_id, current_user)
    incoming = inp.dict(exclude={"client_id"})
    # `is None`, not `or`: an explicit 0 must reset the discount, only a
    # missing field keeps the stored one. Same for the two option lists.
    raw_discount = inp.discount if inp.discount is not None else existing.get("discount")
    discount_value = max(0.0, min(100.0, float(raw_discount or 0)))
    discounts = inp.discounts if inp.discounts is not None else existing.get("discounts")
    additional = (
        inp.additional_discounts if inp.additional_discounts is not None else existing.get("additional_discounts")
    )
    old_unit_code = existing.get("unit") or "kom"
    updated = {
        **existing,
        **incoming,
        "id": product_id,
        "unit": old_unit_code,
        "discount": discount_value,
        "discounts": normalize_discounts(discounts, discount_value),
        "additional_discounts": normalize_additional_discounts(additional),
    }
    # Fields of modules the client doesn't have are ignored (stored values stay).
    in_barcode, in_package = inp.barcode, inp.package_barcode
    in_track = inp.track_expiry
    if not await client_has_module(existing["client_id"], Module.WAREHOUSE):
        in_barcode = in_package = None
    if not await client_has_module(existing["client_id"], Module.EXPIRY):
        in_track = None
    was_active = existing.get("active", True)
    updated["active"] = was_active if inp.active is None else inp.active
    was_tracking = bool(existing.get("track_expiry"))
    updated["track_expiry"] = was_tracking if in_track is None else in_track
    if updated["active"] and not was_active and not (updated.get("price_no_vat") or 0) > 0:
        raise HTTPException(status_code=400, detail="Set a price before activating the product")
    if in_barcode is None:
        updated["barcode"] = existing.get("barcode")
    else:
        updated["barcode"] = await _checked_barcode(existing["client_id"], in_barcode, exclude_id=product_id)
    if in_package is None:
        updated["package_barcode"] = existing.get("package_barcode")
    else:
        updated["package_barcode"] = await _checked_barcode(
            existing["client_id"], in_package, exclude_id=product_id, other=updated["barcode"]
        )
    _require_package_size(updated["package_barcode"], updated.get("pieces_per_package"))
    if updated["barcode"] and updated["barcode"] == updated["package_barcode"]:
        raise HTTPException(status_code=409, detail="Piece and box barcode must differ")
    if inp.unit and inp.unit != old_unit_code:
        new_unit = await _unit_or_400(inp.unit)
        await _check_unit_change(existing, await _unit_or_400(old_unit_code, must_be_active=False), new_unit)
        updated["unit"] = new_unit["code"]
    _require_count_unit_for_packaging(
        await _unit_or_400(updated["unit"], must_be_active=False),
        updated.get("pieces_per_package"), updated.get("boxes_per_transport"), updated.get("package_barcode"),
    )
    if updated["track_expiry"] != was_tracking:
        await _switch_expiry_tracking(existing, updated["track_expiry"])
    # $set instead of replace_one: a concurrent stock change ($inc) must not
    # be overwritten by this stale copy of the product.
    await db.products.update_one(
        {"id": product_id}, {"$set": {k: v for k, v in updated.items() if k not in ("_id", "stock_qty")}}
    )
    if existing.get("image") != updated.get("image"):
        await _delete_product_image(existing.get("image"))
    return Product(**updated)


@api_router.delete("/products/{product_id}")
async def delete_product(product_id: str, current_user: User = Depends(require_manager)):
    existing = await get_scoped_or_404("products", product_id, current_user)
    await db.products.delete_one({"id": product_id})
    await db.stock_batches.delete_many({"product_id": product_id})
    await _delete_product_image(existing.get("image"))
    return {"ok": True}


# ---------------- Units of measure (PLAN_JEDINICE_MERE.md) ----------------
# (code, name_sr, name_en, symbol, group, base_code, factor)
DEFAULT_UNITS = [
    ("kom", "komad", "piece", "kom", "count", "kom", 1),
    ("g", "gram", "gram", "g", "mass", "g", 1),
    ("kg", "kilogram", "kilogram", "kg", "mass", "g", 1000),
    ("t", "tona", "tonne", "t", "mass", "g", 1_000_000),
    ("ml", "mililitar", "millilitre", "ml", "volume", "ml", 1),
    ("l", "litar", "litre", "l", "volume", "ml", 1000),
    ("hl", "hektolitar", "hectolitre", "hl", "volume", "ml", 100_000),
    ("mm", "milimetar", "millimetre", "mm", "length", "mm", 1),
    ("cm", "centimetar", "centimetre", "cm", "length", "mm", 10),
    ("m", "metar", "metre", "m", "length", "mm", 1000),
]


class Unit(BaseModel):
    code: str  # fixed once created
    name_sr: str
    name_en: str
    symbol: str
    group: str  # count / mass / volume / length / ...
    base_code: str  # the unit quantities are stored in (itself for a base unit)
    factor: int  # base units in one of this unit (kg = 1000 g)
    active: bool = True
    created_at: str = Field(default_factory=now_iso)


class UnitInput(BaseModel):
    code: str = Field(pattern=r"^[a-z0-9]{1,10}$")
    name_sr: str = Field(min_length=1, max_length=40)
    name_en: str = Field(min_length=1, max_length=40)
    symbol: str = Field(min_length=1, max_length=10)
    group: str = Field(pattern=r"^[a-z]{1,20}$")
    base_code: str = Field(pattern=r"^[a-z0-9]{1,10}$")
    factor: int = Field(ge=1, le=1_000_000_000)
    active: bool = True


class UnitUpdate(BaseModel):
    name_sr: Optional[str] = Field(default=None, min_length=1, max_length=40)
    name_en: Optional[str] = Field(default=None, min_length=1, max_length=40)
    symbol: Optional[str] = Field(default=None, min_length=1, max_length=10)
    active: Optional[bool] = None
    # Only while no product uses the unit.
    group: Optional[str] = Field(default=None, pattern=r"^[a-z]{1,20}$")
    base_code: Optional[str] = Field(default=None, pattern=r"^[a-z0-9]{1,10}$")
    factor: Optional[int] = Field(default=None, ge=1, le=1_000_000_000)


async def _seed_units() -> None:
    """Adds the built-in units that are missing; never touches existing ones."""
    for code, name_sr, name_en, symbol, group, base_code, factor in DEFAULT_UNITS:
        await db.units.update_one(
            {"code": code},
            {"$setOnInsert": {
                "_id": code, "code": code, "name_sr": name_sr, "name_en": name_en, "symbol": symbol, "group": group,
                "base_code": base_code, "factor": factor, "active": True, "created_at": now_iso(),
            }},
            upsert=True,
        )


async def _unit_or_400(code: str, must_be_active: bool = True) -> dict:
    unit = await db.units.find_one({"code": code}, {"_id": 0})
    if not unit:
        raise HTTPException(status_code=400, detail=f"Unknown unit: {code}")
    if must_be_active and not unit.get("active", True):
        raise HTTPException(status_code=400, detail=f"Unit is not active: {code}")
    return unit


def _require_count_unit_for_packaging(unit: dict, pieces: Optional[int], boxes: Optional[int], package_barcode: Optional[str]) -> None:
    """Boxes and transport packaging only make sense for products counted in pieces."""
    if unit.get("group") != "count" and ((pieces or 0) > 0 or (boxes or 0) > 0 or package_barcode):
        raise HTTPException(status_code=400, detail="Packaging is only for products counted in pieces")


async def _check_unit_change(product: dict, old: dict, new: dict) -> None:
    """A unit with the same base unit (m <-> cm) stores the same numbers, so it is
    always fine. Any other change would reinterpret stored quantities."""
    if old["base_code"] == new["base_code"]:
        return
    in_use = (
        (product.get("stock_qty") or 0) != 0
        or await db.stock_movements.find_one({"product_id": product["id"]}, {"_id": 1})
        or await db.orders.find_one({"items.product_id": product["id"]}, {"_id": 1})
    )
    if in_use:
        raise HTTPException(status_code=409, detail="Unit cannot be changed: the product has stock or orders")


async def _check_unit_structure(code: str, group: str, base_code: str, factor: int) -> None:
    if base_code == code:
        if factor != 1:
            raise HTTPException(status_code=400, detail="A base unit must have factor 1")
        return
    base = await db.units.find_one({"code": base_code}, {"_id": 0})
    if not base or base["base_code"] != base_code:
        raise HTTPException(status_code=400, detail="base_code must be an existing base unit")
    if base["group"] != group:
        raise HTTPException(status_code=400, detail="A unit must be in the same group as its base unit")


@api_router.get("/units", response_model=List[Unit])
async def list_units(current_user: User = Depends(get_current_user)):
    """All units (inactive ones too, so existing products can still show theirs)."""
    docs = await db.units.find({}, {"_id": 0}).to_list(None)
    return sorted((Unit(**d) for d in docs), key=lambda u: (u.group, u.factor, u.code))


@api_router.post("/units", response_model=Unit)
async def create_unit(inp: UnitInput, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    if await db.units.find_one({"code": inp.code}, {"_id": 1}):
        raise HTTPException(status_code=409, detail="A unit with that code already exists")
    await _check_unit_structure(inp.code, inp.group, inp.base_code, inp.factor)
    unit = Unit(**inp.dict())
    await db.units.insert_one({**unit.dict(), "_id": unit.code})
    await _audit(current_user, "unit.create", "unit", unit.code, unit.name_sr)
    return unit


@api_router.put("/units/{code}", response_model=Unit)
async def update_unit(code: str, inp: UnitUpdate, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    existing = await db.units.find_one({"code": code}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Unit not found")
    updated = {**existing}
    for field in ("name_sr", "name_en", "symbol", "active"):
        value = getattr(inp, field)
        if value is not None:
            updated[field] = value
    structural = {f: getattr(inp, f) for f in ("group", "base_code", "factor") if getattr(inp, f) is not None}
    if any(structural.get(f) not in (None, existing[f]) for f in structural):
        in_use = (
            code == "kom"
            or await db.products.find_one({"unit": code}, {"_id": 1})
            or await db.units.find_one({"base_code": code, "code": {"$ne": code}}, {"_id": 1})
        )
        if in_use:
            raise HTTPException(status_code=409, detail="Group, base unit and factor cannot change while the unit is in use")
        updated.update(structural)
        await _check_unit_structure(code, updated["group"], updated["base_code"], updated["factor"])
    await db.units.replace_one({"code": code}, {**updated, "_id": code})
    await _audit(current_user, "unit.update", "unit", code, updated["name_sr"])
    return Unit(**updated)


# ---------------- Barcodes ----------------
def normalize_barcode(value: Any) -> Optional[str]:
    """Trimmed barcode, None when empty. Raises ValueError for odd characters."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores numeric cells as floats
    code = re.sub(r"\s+", "", str(value))
    if not code:
        return None
    if not re.fullmatch(r"[A-Za-z0-9\-]{1,32}", code):
        raise ValueError("Barcode may only contain letters, digits and '-' (max 32 characters)")
    return code


async def _checked_barcode(
    client_id: str, value: Optional[str], exclude_id: Optional[str] = None, other: Optional[str] = None
) -> Optional[str]:
    """Normalized code, 409 if another product already uses it as piece OR box
    barcode. `other` is the same product's other code (a code can't be both)."""
    try:
        code = normalize_barcode(value)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if code:
        if other and other == code:
            raise HTTPException(status_code=409, detail="Piece and box barcode must differ")
        clash = await db.products.find_one(
            {"client_id": client_id, "$or": [{"barcode": code}, {"package_barcode": code}]}, {"id": 1, "name": 1}
        )
        if clash and clash["id"] != exclude_id:
            raise HTTPException(status_code=409, detail=f"Barcode already used by product: {clash.get('name', '')}")
    return code


def _require_package_size(package_barcode: Optional[str], pieces_per_package: Optional[int]) -> None:
    """A box barcode only makes sense with a known number of pieces per box."""
    if package_barcode and not (pieces_per_package or 0) > 0:
        raise HTTPException(status_code=400, detail="Set pieces per package before adding a box barcode")


class BarcodeInput(BaseModel):
    barcode: str = Field(min_length=1, max_length=40)
    kind: Literal["piece", "package"] = "piece"


@api_router.get("/products/by-barcode/{code}", response_model=ProductOut)
async def product_by_barcode(code: str, current_user: User = Depends(require_module(Module.WAREHOUSE))):
    try:
        normalized = normalize_barcode(code)
    except ValueError:
        normalized = None
    scope = _scope_query(current_user)
    doc = (
        await db.products.find_one(
            {**scope, "$or": [{"barcode": normalized}, {"package_barcode": normalized}]}, {"_id": 0}
        )
        if normalized
        else None
    )
    if not doc:
        raise HTTPException(status_code=404, detail="Product not found")
    r = (await _reserved_by_product(scope)).get(doc["id"], 0)
    expiry = await _expiry_info(scope, product_id=doc["id"])
    out = _product_out(doc, r, expiry.get(doc["id"]), current_user)
    if doc.get("package_barcode") == normalized and (doc.get("pieces_per_package") or 0) > 0:
        out.scan_unit, out.scan_qty = "package", int(doc["pieces_per_package"])
    else:
        out.scan_unit, out.scan_qty = "piece", 1
    return out


@api_router.post("/products/{product_id}/barcode", response_model=Product)
async def set_product_barcode(
    product_id: str, inp: BarcodeInput, current_user: User = Depends(require_module(Module.WAREHOUSE, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))
):
    """Links a scanned code to a product (warehouse staff can't edit products)."""
    existing = await get_scoped_or_404("products", product_id, current_user)
    if inp.kind == "package":
        code = await _checked_barcode(
            existing["client_id"], inp.barcode, exclude_id=product_id, other=existing.get("barcode")
        )
        _require_package_size(code, existing.get("pieces_per_package"))
        await db.products.update_one({"id": product_id}, {"$set": {"package_barcode": code}})
        return Product(**{**existing, "package_barcode": code})
    code = await _checked_barcode(
        existing["client_id"], inp.barcode, exclude_id=product_id, other=existing.get("package_barcode")
    )
    await db.products.update_one({"id": product_id}, {"$set": {"barcode": code}})
    return Product(**{**existing, "barcode": code})


class QuickProductInput(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    manufacturer: Optional[str] = Field(default="", max_length=200)
    barcode: Optional[str] = None
    barcode_kind: Literal["piece", "package"] = "piece"  # what `barcode` is; a box needs pieces_per_package
    pieces_per_package: Optional[int] = Field(default=0, ge=0, le=100000)
    boxes_per_transport: Optional[int] = Field(default=0, ge=0, le=100000)
    client_id: Optional[str] = None  # only honored for SUPERADMIN writes


@api_router.post("/products/quick", response_model=Product)
async def quick_add_product(
    inp: QuickProductInput, current_user: User = Depends(require_module(Module.WAREHOUSE, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))
):
    """Warehouse adds a product that has just arrived but isn't in the system.
    It is created inactive with no price: an admin sets price/VAT and activates
    it before sales reps can see or order it."""
    client_id = await resolve_write_client_id(current_user, inp.client_id)
    code = await _checked_barcode(client_id, inp.barcode)
    if inp.barcode_kind == "package":
        _require_package_size(code, inp.pieces_per_package)
    obj = Product(
        client_id=client_id,
        name=inp.name.strip(),
        manufacturer=(inp.manufacturer or "").strip(),
        pieces_per_package=inp.pieces_per_package or 0,
        boxes_per_transport=inp.boxes_per_transport or 0,
        barcode=code if inp.barcode_kind == "piece" else None,
        package_barcode=code if inp.barcode_kind == "package" else None,
        price_no_vat=0,
        active=False,
    )
    await db.products.insert_one({**obj.dict(), "_id": obj.id})
    return obj


# ---------------- Product import from Excel ----------------
IMPORT_MAX_BYTES = 4 * 1024 * 1024
IMPORT_MAX_ROWS = 2000

# Normalized header (lowercase, no diacritics/spaces/punctuation) -> field.
_IMPORT_HEADERS = {
    "name": "name", "naziv": "name", "nazivproizvoda": "name", "proizvod": "name", "artikal": "name",
    "barcode": "barcode", "barkod": "barcode", "ean": "barcode",
    "packagebarcode": "package_barcode", "barkodkutije": "package_barcode", "barkodpakovanja": "package_barcode",
    "barkodkutija": "package_barcode", "eankutije": "package_barcode",
    "manufacturer": "manufacturer", "proizvodjac": "manufacturer",
    "pricenovat": "price_no_vat", "cena": "price_no_vat", "cenabezpdv": "price_no_vat", "cenabezpdva": "price_no_vat",
    "vatrate": "vat_rate", "pdv": "vat_rate", "stopapdv": "vat_rate", "stopapdva": "vat_rate",
    "piecesperpackage": "pieces_per_package", "komadaupakovanju": "pieces_per_package",
    "komadapopakovanju": "pieces_per_package", "pakovanje": "pieces_per_package",
    "boxespertransport": "boxes_per_transport", "transportnopakovanje": "boxes_per_transport",
    "transportno": "boxes_per_transport", "komadaunatransportnom": "boxes_per_transport",
    "stockqty": "stock_qty", "stanje": "stock_qty", "kolicina": "stock_qty", "nastanju": "stock_qty",
    "trackexpiry": "track_expiry", "pratirok": "track_expiry", "pratirokove": "track_expiry",
    "pratirokatrajanja": "track_expiry",
    "unit": "unit", "unitofmeasure": "unit", "jedinicamere": "unit", "jedinica": "unit", "mera": "unit",
}
_IMPORT_TEMPLATE_HEADERS = [
    ("name", "Naziv"), ("barcode", "Barkod"), ("manufacturer", "Proizvođač"), ("price_no_vat", "Cena bez PDV"),
    ("vat_rate", "PDV %"), ("pieces_per_package", "Komada u pakovanju"),
    ("boxes_per_transport", "Transportno pakovanje (komada)"), ("stock_qty", "Stanje (u jedinici mere)"),
    ("track_expiry", "Prati rok (da/ne)"), ("package_barcode", "Barkod kutije"), ("unit", "Jedinica mere"),
]


# Downloadable example: includes a product without a barcode (matched by name)
# and a text price with a comma, both of which the import accepts.
_IMPORT_EXAMPLE_ROWS = [
    ("Jaffa keks 150g", "8601000000018", "Jaffa", 46.5, 20, 12, 48, 240, "ne", "8601000100011", "kom"),
    ("Plazma keks 300g", "8601000000025", "Bambi", 189.9, 20, 10, 40, 120, "ne", "8601000100028", "kom"),
    ("Smoki 50g", None, "Bambi", 39, 20, 24, 96, 0, "ne", None, "kom"),
    ("Mleko 2.8% 1l", "8601000000049", "Imlek", "124,5", 10, 12, 72, 360, "da", "8601000100042", "kom"),
    ("Jogurt 2.8% 1l", "8601000000056", "Imlek", 118, 10, 12, 72, 300, "da", None, "kom"),
    ("Ulje suncokretovo 1l", "8601000000063", "Dijamant", 259, 20, 12, 60, 180, "ne", None, "kom"),
    ("Brašno T-500 1kg", "8601000000070", "Mlin", "79,9", 10, 10, 100, 500, "ne", None, "kom"),
    ("Kafa Grand 200g", "8601000000087", "Strauss", 299, 20, 12, 48, None, "ne", None, "kom"),
    ("Sir kačkavalj (po kg)", "8601000000094", "Mlekara", 1100, 10, 0, 0, "25,5", "ne", None, "kg"),
]


# Import columns that belong to a module: dropped for clients without it.
_IMPORT_FIELD_MODULE = {
    "barcode": Module.WAREHOUSE,
    "package_barcode": Module.WAREHOUSE,
    "stock_qty": Module.STOCK,
    "track_expiry": Module.EXPIRY,
}


async def _import_disabled_fields(client_id: Optional[str]) -> set:
    doc = await db.clients.find_one({"id": client_id}, {"_id": 0, "modules": 1}) if client_id else None
    enabled = set(effective_modules(doc))
    return {f for f, m in _IMPORT_FIELD_MODULE.items() if m.value not in enabled}


def _import_header_key(value: Any) -> str:
    # "Stanje (komada)" and "Prati rok (da/ne)" from the downloadable example
    # must match the plain names: drop the bracketed hint.
    text = re.sub(r"\(.*?\)", "", str(value or "").lower())
    for src, dst in (("đ", "dj"), ("č", "c"), ("ć", "c"), ("š", "s"), ("ž", "z")):
        text = text.replace(src, dst)
    return re.sub(r"[^a-z0-9]", "", text)


def _import_number(value: Any, label: str, integer: bool = False, maximum: Optional[float] = None):
    """Empty cell -> None (field untouched). Raises ValueError on bad input."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
    try:
        num = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label}: not a number")
    if num < 0 or (maximum is not None and num > maximum):
        raise ValueError(f"{label}: out of range")
    if integer:
        if not num.is_integer():
            raise ValueError(f"{label}: must be a whole number")
        return int(num)
    return num


@api_router.get("/products/import-template")
async def product_import_template(
    client_id: Optional[str] = Query(None), current_user: User = Depends(require_manager)
):
    from openpyxl import Workbook
    from fastapi.responses import Response
    import io

    # Only the columns the client's modules allow (superadmin: pass client_id, else all).
    disabled = await _import_disabled_fields(current_user.client_id or client_id)
    keep = [i for i, (field, _) in enumerate(_IMPORT_TEMPLATE_HEADERS) if field not in disabled]
    wb = Workbook()
    ws = wb.active
    ws.title = "Artikli"
    ws.append([_IMPORT_TEMPLATE_HEADERS[i][1] for i in keep])
    for row in _IMPORT_EXAMPLE_ROWS:
        ws.append([row[i] for i in keep])
    for pos, i in enumerate(keep, start=1):
        letter = ws.cell(row=1, column=pos).column_letter
        if _IMPORT_TEMPLATE_HEADERS[i][0] in ("barcode", "package_barcode"):  # text keeps leading zeros
            for cell in ws[letter][1:]:
                cell.number_format = "@"
        ws.column_dimensions[letter].width = (32, 18, 22, 14, 8, 20, 30, 22, 18, 18, 16)[i]
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="uvoz-artikala.xlsx"'},
    )


@api_router.post("/products/import")
async def import_products(
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    client_id: Optional[str] = Query(None),
    current_user: User = Depends(require_manager),
):
    """Excel import. Rows are matched to existing products by barcode, else by
    exact name; empty cells leave stored values alone. dry_run (default) only
    reports what would happen; a bad row never blocks the good ones."""
    target_client = await resolve_write_client_id(current_user, client_id)
    data = await file.read(IMPORT_MAX_BYTES + 1)
    if len(data) > IMPORT_MAX_BYTES:
        raise HTTPException(status_code=413, detail="File too large (max 4 MB)")
    try:
        from openpyxl import load_workbook
        import io

        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        raw_rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read the file - upload an .xlsx workbook")
    if not raw_rows:
        raise HTTPException(status_code=400, detail="The file is empty")
    columns = {}
    for idx, head in enumerate(raw_rows[0]):
        field = _IMPORT_HEADERS.get(_import_header_key(head))
        if field and field not in columns.values():
            columns[idx] = field
    # Columns of modules this client doesn't have are skipped (and reported).
    disabled = await _import_disabled_fields(target_client)
    ignored_columns = sorted({f for f in columns.values() if f in disabled})
    columns = {idx: f for idx, f in columns.items() if f not in disabled}
    if "name" not in columns.values() and "barcode" not in columns.values():
        raise HTTPException(status_code=400, detail="Missing a 'Naziv' (name) or 'Barkod' (barcode) column")
    body = [(i + 2, r) for i, r in enumerate(raw_rows[1:]) if any(c not in (None, "") for c in r)]
    if len(body) > IMPORT_MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"Too many rows (max {IMPORT_MAX_ROWS})")

    existing = await db.products.find({"client_id": target_client}, {"_id": 0}).to_list(None)
    by_barcode = {p["barcode"]: p for p in existing if p.get("barcode")}
    by_package = {p["package_barcode"]: p for p in existing if p.get("package_barcode")}
    by_name = {p["name"].strip().lower(): p for p in existing if p.get("name")}
    seen_barcodes: Dict[str, int] = {}
    seen_ids: Dict[str, int] = {}
    units_by_code: Dict[str, dict] = {}
    unit_lookup: Dict[str, dict] = {}
    async for u in db.units.find({}, {"_id": 0}):
        units_by_code[u["code"]] = u
        for label in (u["code"], u["symbol"], u["name_sr"], u["name_en"]):
            unit_lookup.setdefault(_import_header_key(label), u)
    rows_out = []
    for row_no, row in body:
        values = {field: (row[idx] if idx < len(row) else None) for idx, field in columns.items()}
        errors: List[str] = []
        parsed: Dict[str, Any] = {}
        name = str(values.get("name") or "").strip()
        try:
            code = normalize_barcode(values.get("barcode"))
        except ValueError as exc:
            code = None
            errors.append(str(exc))
        try:
            package_code = normalize_barcode(values.get("package_barcode"))
        except ValueError as exc:
            package_code = None
            errors.append(str(exc))
        for field, kwargs in (
            ("price_no_vat", {}), ("vat_rate", {"maximum": 100}), ("pieces_per_package", {"integer": True}),
            ("boxes_per_transport", {"integer": True}),
        ):
            if field in columns.values():
                try:
                    num = _import_number(values.get(field), field, **kwargs)
                except ValueError as exc:
                    errors.append(str(exc))
                    continue
                if num is not None:
                    parsed[field] = num
        manufacturer = str(values.get("manufacturer") or "").strip()
        if manufacturer:
            parsed["manufacturer"] = manufacturer
        raw_track = str(values.get("track_expiry") or "").strip().lower()
        if raw_track:
            if raw_track in ("da", "yes", "true", "1", "x"):
                parsed["track_expiry"] = True
            elif raw_track in ("ne", "no", "false", "0"):
                parsed["track_expiry"] = False
            else:
                errors.append("Prati rok must be da or ne")

        product = by_barcode.get(code) if code else None
        if product is None and package_code:
            product = by_package.get(package_code)
        if product is None and name:
            product = by_name.get(name.lower())
        # Unit of measure: the cell, else the product's own, else pieces. The stock
        # column is in that unit (25,5 kg) and is stored as whole base units (g).
        unit_cell = str(values.get("unit") or "").strip()
        unit_doc = None
        if unit_cell:
            unit_doc = unit_lookup.get(_import_header_key(unit_cell))
            if unit_doc is None or not unit_doc.get("active", True):
                errors.append(f"Unknown unit: {unit_cell}")
                unit_doc = None
        current_unit = units_by_code.get((product or {}).get("unit") or "kom") or units_by_code.get("kom")
        if unit_doc is not None and product is not None and unit_doc["code"] != current_unit["code"]:
            try:
                await _check_unit_change(product, current_unit, unit_doc)
            except HTTPException as exc:
                errors.append(str(exc.detail))
                unit_doc = None
        final_unit = unit_doc or current_unit
        if unit_doc is not None:
            parsed["unit"] = unit_doc["code"]
        if "stock_qty" in columns.values():
            try:
                shown = _import_number(values.get("stock_qty"), "stock_qty", maximum=1_000_000)
            except ValueError as exc:
                shown = None
                errors.append(str(exc))
            if shown is not None:
                base = shown * final_unit["factor"]
                if abs(base - round(base)) > 1e-6:
                    errors.append(f"stock_qty: too many decimals for {final_unit['symbol']}")
                else:
                    parsed["stock_qty"] = int(round(base))
        if final_unit["group"] != "count" and (
            (parsed.get("pieces_per_package", (product or {}).get("pieces_per_package")) or 0) > 0
            or (parsed.get("boxes_per_transport", (product or {}).get("boxes_per_transport")) or 0) > 0
            or package_code or (product or {}).get("package_barcode")
        ):
            errors.append("Packaging is only for products counted in pieces")
        if product is not None and code and product.get("barcode") not in (None, code):
            errors.append("Product already has a different barcode")
        if code and code in seen_barcodes:
            errors.append(f"Barcode repeated (row {seen_barcodes[code]})")
        if package_code:
            if package_code in seen_barcodes:
                errors.append(f"Box barcode repeated (row {seen_barcodes[package_code]})")
            if product is not None and product.get("package_barcode") not in (None, package_code):
                errors.append("Product already has a different box barcode")
            owner = by_package.get(package_code) or by_barcode.get(package_code)
            if owner and (product is None or owner["id"] != product["id"]):
                errors.append("Box barcode belongs to another product")
            pieces = parsed.get("pieces_per_package", (product or {}).get("pieces_per_package") or 0)
            if not pieces > 0:
                errors.append("Set 'Komada u pakovanju' before adding a box barcode")
        elif (
            product is not None
            and product.get("package_barcode")
            and "pieces_per_package" in parsed
            and not parsed["pieces_per_package"] > 0
        ):
            errors.append("Pieces per package can't be 0 while the product has a box barcode")
        if code:
            owner = by_package.get(code)
            if owner and (product is None or owner["id"] != product["id"]):
                errors.append("Barcode belongs to another product")
        final_piece = code or (product or {}).get("barcode")
        final_box = package_code or (product or {}).get("package_barcode")
        if final_piece and final_piece == final_box:
            errors.append("Piece and box barcode must differ")
        if product is not None and product["id"] in seen_ids:
            errors.append(f"Same product repeated (row {seen_ids[product['id']]})")
        if product is None and not name:
            errors.append("Name is required for a new product")
        if product is not None and code and not product.get("barcode"):
            clash = by_barcode.get(code)
            if clash and clash["id"] != product["id"]:
                errors.append("Barcode belongs to another product")
        if (
            product is not None
            and parsed.get("track_expiry") is False
            and product.get("track_expiry")
            and await db.stock_batches.find_one(
                {"product_id": product["id"], "expiry_date": {"$ne": None}, "qty": {"$ne": 0}}, {"_id": 1}
            )
        ):
            errors.append("Can't switch off expiry tracking while dated batches are in stock")
        if code:
            seen_barcodes.setdefault(code, row_no)
        if package_code:
            seen_barcodes.setdefault(package_code, row_no)
        if product is not None:
            seen_ids.setdefault(product["id"], row_no)
        entry = {
            "row": row_no,
            "name": name or (product or {}).get("name", ""),
            "barcode": code,
            "package_barcode": package_code,
            "action": "error" if errors else ("update" if product else "create"),
            "errors": errors,
        }
        if not errors:
            entry["_product"] = product
            entry["_values"] = {
                **parsed,
                **({"name": name} if name else {}),
                **({"barcode": code} if code else {}),
                **({"package_barcode": package_code} if package_code else {}),
            }
        rows_out.append(entry)

    if not dry_run:
        for entry in rows_out:
            if entry["action"] == "error":
                continue
            vals = dict(entry["_values"])
            stock = vals.pop("stock_qty", None)
            product = entry["_product"]
            if product is None:
                obj = Product(client_id=target_client, **{k: v for k, v in vals.items()})
                await db.products.insert_one({**obj.dict(), "_id": obj.id})
                product = obj.dict()
            elif vals:
                if "track_expiry" in vals and vals["track_expiry"] != bool(product.get("track_expiry")):
                    await _switch_expiry_tracking(product, vals["track_expiry"])
                await db.products.update_one({"id": product["id"]}, {"$set": vals})
            if stock is not None:
                fresh = await db.products.find_one({"id": product["id"]}, {"_id": 0})
                if fresh.get("stock_qty") != stock:
                    await _set_stock_count(fresh, stock, current_user, "Uvoz iz Excela")
    for entry in rows_out:
        entry.pop("_product", None)
        entry.pop("_values", None)
    summary = {k: sum(1 for r in rows_out if r["action"] == k) for k in ("create", "update", "error")}
    return {"dry_run": dry_run, "summary": summary, "rows": rows_out, "ignored_columns": ignored_columns}


# ---------------- Orders ----------------
def _allowed_supplier_discount(product: dict, requested: Optional[float]) -> float:
    """Supplier discount = product default, or another value the admin
    whitelisted in product.discounts. Anything else is rejected."""
    default = float(product.get("discount") or 0)
    if requested is None:
        return default
    allowed = {float(v) for v in (product.get("discounts") or [])} | {default}
    if float(requested) not in allowed:
        raise HTTPException(status_code=400, detail=f"Discount {requested}% is not allowed for {product['name']}")
    return float(requested)


def _allowed_additional_discount(product: dict, requested: Optional[float]) -> float:
    """Additional discount must be one of product.additional_discounts
    (0 = none is always allowed)."""
    if not requested:
        return 0.0
    allowed = {float(v) for v in (product.get("additional_discounts") or [])}
    if float(requested) not in allowed:
        raise HTTPException(
            status_code=400, detail=f"Additional discount {requested}% is not allowed for {product['name']}"
        )
    return float(requested)


def _rounded_totals(totals: dict) -> OrderTotals:
    return OrderTotals(
        subtotal=round(totals["subtotal"], 2), vat=round(totals["vat"], 2), grand=round(totals["grand"], 2)
    )


def _order_out(doc: dict) -> OrderOut:
    shipped = doc.get("status") == OrderStatus.SHIPPED.value
    items = [{**item, "line_net": round(line_net(item, shipped), 2)} for item in doc.get("items", [])]
    return OrderOut(
        **{**doc, "items": items},
        totals=_rounded_totals(compute_order_totals(doc)),
        ordered_totals=_rounded_totals(compute_order_totals(doc, by_ordered=True)),
    )


def _without_batches(doc: dict) -> dict:
    """Sales reps never see which batches (expiry dates) an order was packed from."""
    return {**doc, "items": [{k: v for k, v in i.items() if k != "picked_batches"} for i in doc.get("items", [])]}


def _order_scope_query(user: User) -> dict:
    """Tenant scope for orders, plus: an operator sees only their own."""
    query = _scope_query(user)
    if user.role == Role.OPERATOR:
        query = {**query, "created_by_user_id": user.id}
    return query


def _parse_date(value: str, name: str) -> datetime:
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{name} must be YYYY-MM-DD")


@api_router.get("/orders", response_model=List[OrderOut])
async def list_orders(
    customer_id: Optional[str] = None,
    portal: bool = False,
    status: Optional[List[str]] = Query(None),
    created_by_user_id: Optional[str] = None,
    from_date: Optional[str] = Query(None, description="YYYY-MM-DD, inclusive"),
    to_date: Optional[str] = Query(None, description="YYYY-MM-DD, inclusive"),
    limit: Optional[int] = Query(None, ge=1, le=500),
    skip: int = Query(0, ge=0),
    current_user: User = Depends(get_current_user),
):
    query = _order_scope_query(current_user)
    if customer_id:
        query = {**query, "customer_id": customer_id}
    if status:
        query = {**query, "status": {"$in": status}}
    # Ignored for operators: they are already limited to their own orders.
    if created_by_user_id and current_user.role != Role.OPERATOR:
        query = {**query, "created_by_user_id": created_by_user_id}
    date_range = {}
    if from_date:
        date_range["$gte"] = _parse_date(from_date, "from_date").isoformat()
    if to_date:
        date_range["$lt"] = (_parse_date(to_date, "to_date") + timedelta(days=1)).isoformat()
    if date_range:
        query = {**query, "created_at": date_range}
    cursor = db.orders.find(query, {"_id": 0}).sort("created_at", -1).skip(skip)
    if limit:
        cursor = cursor.limit(limit)
    docs = await cursor.to_list(None)
    if portal and current_user.role == Role.SUPERADMIN:
        # One query for all client names instead of one per order.
        client_ids = list({d.get("client_id") for d in docs})
        clients = await db.clients.find({"id": {"$in": client_ids}}, {"_id": 0, "id": 1, "name": 1}).to_list(None)
        names = {c["id"]: c.get("name") for c in clients}
        docs = [{**d, "client_name": names.get(d.get("client_id"))} for d in docs]
    if current_user.role == Role.OPERATOR:
        docs = [_without_batches(d) for d in docs]
    return [_order_out(d) for d in docs]


@api_router.get("/orders/{order_id}", response_model=OrderOut)
async def get_order(order_id: str, current_user: User = Depends(get_current_user)):
    order = await get_scoped_or_404("orders", order_id, current_user)
    if current_user.role == Role.OPERATOR and order.get("created_by_user_id") != current_user.id:
        raise HTTPException(status_code=404, detail="Not found")  # someone else's order
    return _order_out(_without_batches(order) if current_user.role == Role.OPERATOR else order)


async def _resolve_order_lines(client_id: str, inp: "OrderInput") -> tuple[dict, List[OrderItem]]:
    """Validates an order payload against this client's customer/products and
    returns (customer, items snapshotted from the stored products)."""
    if not inp.items:
        raise HTTPException(status_code=400, detail="Order must contain at least one item")
    # Customer and products are looked up inside this client only, so an
    # order can't reference another tenant's data.
    customer = await db.customers.find_one({"id": inp.customer_id, "client_id": client_id}, {"_id": 0})
    if not customer:
        raise HTTPException(status_code=400, detail="Customer not found")

    product_ids = list({line.product_id for line in inp.items})
    product_docs = await db.products.find(
        {"id": {"$in": product_ids}, "client_id": client_id}, {"_id": 0}
    ).to_list(None)
    products = {p["id"]: p for p in product_docs}
    unit_codes = list({p.get("unit") or "kom" for p in product_docs})
    units = {u["code"]: u async for u in db.units.find({"code": {"$in": unit_codes}}, {"_id": 0})}

    items: List[OrderItem] = []
    for line in inp.items:
        product = products.get(line.product_id)
        if not product:
            raise HTTPException(status_code=400, detail=f"Product not found: {line.product_id}")
        if line.ordered_qty <= 0:
            raise HTTPException(status_code=400, detail=f"Quantity must be positive for {product['name']}")
        if product.get("active", True) is False:
            raise HTTPException(status_code=400, detail=f"Product is not available: {product['name']}")
        # Snapshot from the stored product - never from the payload - so a
        # client can't set its own price/VAT/discount.
        items.append(OrderItem(
            product_id=product["id"],
            name=product["name"],
            image=product.get("image", ""),
            manufacturer=product.get("manufacturer", ""),
            price_no_vat=product.get("price_no_vat", 0),
            vat_rate=product.get("vat_rate", 20),
            pieces_per_package=product.get("pieces_per_package", 0),
            boxes_per_transport=product.get("boxes_per_transport", 0),
            unit=product.get("unit") or "kom",
            unit_symbol=(units.get(product.get("unit") or "kom") or {}).get("symbol") or (product.get("unit") or "kom"),
            unit_factor=(units.get(product.get("unit") or "kom") or {}).get("factor") or 1,
            discount=_allowed_supplier_discount(product, line.discount),
            additional_discount=_allowed_additional_discount(product, line.additional_discount),
            ordered_qty=line.ordered_qty,
        ))
    return customer, items


@api_router.post("/orders", response_model=OrderOut)
async def create_order(inp: OrderInput, current_user: User = Depends(require_order_creator)):
    client_id = await resolve_write_client_id(current_user, inp.client_id)
    if not inp.items:
        raise HTTPException(status_code=400, detail="Order must contain at least one item")

    customer, items = await _resolve_order_lines(client_id, inp)

    obj = Order(
        client_id=client_id,
        customer_id=customer["id"],
        customer_name=customer["name"],
        items=items,
        created_by_user_id=current_user.id,
        created_by_name=current_user.name,
        status_history=[StatusChange(
            from_status=None,
            to_status=OrderStatus.NEW.value,
            changed_by_user_id=current_user.id,
            changed_by_name=current_user.name,
            changed_by_role=current_user.role.value,
        )],
    )
    doc = obj.dict()
    await db.orders.insert_one({**doc, "_id": obj.id})
    return _order_out(doc)


@api_router.put("/orders/{order_id}", response_model=OrderOut)
async def update_order(order_id: str, inp: OrderInput, current_user: User = Depends(require_order_creator)):
    """Edit an order's customer and lines while it is still `new` - by its
    creator (operator) or an admin. Prices are re-snapshotted from the
    current catalog, like on creation."""
    order = await _get_order_for_processing(order_id, current_user)
    if _status_value(order) != OrderStatus.NEW.value:
        raise HTTPException(status_code=409, detail="Only new orders can be edited")
    customer, items = await _resolve_order_lines(order["client_id"], inp)
    # Conditional on status: if the warehouse took the order meanwhile, 409.
    result = await db.orders.update_one(
        {"id": order_id, "status": OrderStatus.NEW.value},
        {"$set": {
            "customer_id": customer["id"],
            "customer_name": customer["name"],
            "items": [i.dict() for i in items],
            "updated_at": now_iso(),
        }},
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=409, detail="Order was changed in the meantime")
    return _order_out(await db.orders.find_one({"id": order_id}, {"_id": 0}))


@api_router.delete("/orders/{order_id}")
async def delete_order(order_id: str, current_user: User = Depends(require_manager)):
    order = await get_scoped_or_404("orders", order_id, current_user)
    if _status_value(order) != OrderStatus.NEW.value:
        raise HTTPException(status_code=409, detail="Only new orders can be deleted - cancel the order instead")
    await db.orders.delete_one({"id": order_id})
    return {"ok": True}


# ---------------- Order status workflow (RBAC_PLAN.md section 4) ----------------
class StatusChangeInput(BaseModel):
    status: OrderStatus
    note: Optional[str] = None
    # Only for -> shipped on a client with manual invoice numbering.
    invoice_number: Optional[str] = Field(default=None, max_length=40)


class PickedBatchInput(BaseModel):
    batch_id: str
    qty: int = Field(gt=0, le=1_000_000)


class PickedItemInput(BaseModel):
    product_id: str
    picked_qty: Optional[int] = None  # None = un-check the line
    # Optional manual batch choice (must add up to picked_qty); without it the
    # server takes the earliest-expiring batches (FEFO) when the order ships.
    batches: Optional[List[PickedBatchInput]] = None


class PickedItemsInput(BaseModel):
    items: List[PickedItemInput]


def _status_value(order: dict) -> str:
    s = order.get("status") or OrderStatus.NEW.value
    return s.value if isinstance(s, Enum) else s


# (from, to) -> roles allowed; admin/superadmin may override anything (with a note).
_WAREHOUSE_FLOW = {
    ("new", "in_progress"): {Role.WAREHOUSE},
    ("in_progress", "new"): {Role.WAREHOUSE},
    ("in_progress", "shipped"): {Role.WAREHOUSE},
    ("new", "rejected"): {Role.WAREHOUSE},
    ("in_progress", "rejected"): {Role.WAREHOUSE},
    ("new", "canceled"): {Role.OPERATOR},
}
_MANAGER_ROLES = {Role.SUPERADMIN, Role.ADMIN}
_NOTE_REQUIRED_TARGETS = {"rejected"}


def _history_entry(user: User, from_status: Optional[str], to_status: str, note: Optional[str]) -> dict:
    return StatusChange(
        from_status=from_status,
        to_status=to_status,
        changed_by_user_id=user.id,
        changed_by_name=user.name,
        changed_by_role=user.role.value,
        note=note,
    ).dict()


async def _get_order_for_processing(order_id: str, user: User) -> dict:
    """Order lookup that applies the operator's own-orders rule like get_order."""
    order = await get_scoped_or_404("orders", order_id, user)
    if user.role == Role.OPERATOR and order.get("created_by_user_id") != user.id:
        raise HTTPException(status_code=404, detail="Not found")
    return order


async def _assign_invoice_number(order: dict, manual_number: Optional[str]) -> Dict[str, Any]:
    """Invoice number for the first shipment (RBAC_PLAN.md section 5): from
    the per-client counter ("auto") or typed by the warehouse ("manual")."""
    client = await db.clients.find_one({"id": order["client_id"]}, {"_id": 0}) or {}
    if client.get("invoice_numbering", "auto") == "manual":
        number = (manual_number or "").strip()
        if not number:
            raise HTTPException(status_code=400, detail="Invoice number is required")
        if await db.orders.find_one({"client_id": order["client_id"], "invoice_number": number}, {"_id": 1}):
            raise HTTPException(status_code=409, detail="Invoice number is already used")
        return {"invoice_number": number}
    prefix = client.get("invoice_prefix") or ""
    if not prefix:
        raise HTTPException(status_code=400, detail="Set the invoice prefix for this client first")
    counter = await db.counters.find_one_and_update(
        {"_id": _invoice_counter_id(order["client_id"])},
        {"$inc": {"seq": 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    return {"invoice_seq": counter["seq"], "invoice_number": f"{prefix}/{counter['seq']:04d}"}


@api_router.post("/orders/{order_id}/status", response_model=OrderOut)
async def change_order_status(
    order_id: str, inp: StatusChangeInput, current_user: User = Depends(get_current_user)
):
    order = await _get_order_for_processing(order_id, current_user)
    current = _status_value(order)
    target = inp.status.value
    note = (inp.note or "").strip() or None

    if current == target:
        raise HTTPException(status_code=400, detail="Order already has that status")
    base_flow = {OrderStatus.NEW.value, OrderStatus.CANCELED.value}
    if (current not in base_flow or target not in base_flow) and not has_module(current_user, Module.WAREHOUSE):
        # Without the warehouse module an order is only new or canceled.
        raise HTTPException(status_code=403, detail=f"Module not enabled: {Module.WAREHOUSE.value}")
    override = current_user.role in _MANAGER_ROLES
    if override:
        # Anything -> anything, but only with a reason, except the normal flow.
        if (current, target) not in _WAREHOUSE_FLOW and not note:
            raise HTTPException(status_code=400, detail="A note is required to override the status flow")
    else:
        allowed = _WAREHOUSE_FLOW.get((current, target))
        if allowed is None:
            raise HTTPException(status_code=400, detail=f"Transition {current} -> {target} is not allowed")
        if current_user.role not in allowed:
            raise HTTPException(status_code=403, detail="Forbidden")
    if target in _NOTE_REQUIRED_TARGETS and not note:
        raise HTTPException(status_code=400, detail="A note is required")

    if target == OrderStatus.SHIPPED.value:
        items = order.get("items", [])
        if any(i.get("picked_qty") is None for i in items):
            raise HTTPException(status_code=400, detail="Check every item (picked_qty) before shipping")
        if not any((i.get("picked_qty") or 0) > 0 for i in items):
            raise HTTPException(status_code=400, detail="Nothing was picked - reject the order instead")

    update: Dict[str, Any] = {"status": target}
    if target == OrderStatus.SHIPPED.value and not order.get("invoice_number"):
        # An order keeps the number from its first shipment, even if an admin
        # later moves it back and ships it again.
        update.update(await _assign_invoice_number(order, inp.invoice_number))
    if target == OrderStatus.IN_PROGRESS.value:
        update["assigned_to_user_id"] = current_user.id
        update["assigned_to_name"] = current_user.name
    elif target == OrderStatus.NEW.value:
        update["assigned_to_user_id"] = None
        update["assigned_to_name"] = None
    if target == OrderStatus.SHIPPED.value:
        update["shipped_at"] = now_iso()

    # Conditional update: if someone else changed the status meanwhile, 409.
    try:
        result = await db.orders.update_one(
            {"id": order_id, "status": current},
            {"$set": update, "$push": {"status_history": _history_entry(current_user, current, target, note)}},
        )
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail="Invoice number is already used")
    if result.matched_count == 0:
        raise HTTPException(status_code=409, detail="Order was changed in the meantime")
    updated_doc = await db.orders.find_one({"id": order_id}, {"_id": 0})
    if await client_has_module(updated_doc["client_id"], Module.STOCK):
        if target == OrderStatus.SHIPPED.value:
            await _move_stock_for_order(updated_doc, -1, "shipment", current_user)
        elif current == OrderStatus.SHIPPED.value:
            # An admin took a shipped order back: the goods return to stock.
            await _move_stock_for_order(updated_doc, +1, "reversal", current_user)
    if target in (OrderStatus.SHIPPED.value, OrderStatus.REJECTED.value):
        await _send_order_status_email(updated_doc, target, note, current_user)
    return _order_out(updated_doc)


async def _validated_picked_batches(item: dict, line: PickedItemInput) -> List[dict]:
    """The warehouse's own batch choice for one line: only for products with
    expiry tracking, from this product's batches, in-date, adding up to picked_qty."""
    product = await db.products.find_one({"id": item["product_id"]}, {"_id": 0, "track_expiry": 1})
    if not (product or {}).get("track_expiry"):
        raise HTTPException(status_code=400, detail=f"{item['name']} has no expiry tracking")
    if line.picked_qty is None or sum(b.qty for b in line.batches) != line.picked_qty:
        raise HTTPException(status_code=400, detail=f"Batches for {item['name']} must add up to picked_qty")
    today = _today().isoformat()
    picks: Dict[str, dict] = {}
    for b in line.batches:
        batch = await db.stock_batches.find_one({"id": b.batch_id, "product_id": item["product_id"]}, {"_id": 0})
        if not batch:
            raise HTTPException(status_code=400, detail=f"Unknown batch for {item['name']}")
        if batch.get("expiry_date") and batch["expiry_date"] < today:
            raise HTTPException(status_code=400, detail=f"Batch for {item['name']} has expired")
        pick = picks.setdefault(batch["id"], {"batch_id": batch["id"], "expiry_date": batch.get("expiry_date"), "qty": 0})
        pick["qty"] += b.qty
        if pick["qty"] > batch["qty"]:
            raise HTTPException(status_code=400, detail=f"Not enough pieces in the chosen batch of {item['name']}")
    return list(picks.values())


@api_router.patch("/orders/{order_id}/items", response_model=OrderOut)
async def update_picked_items(
    order_id: str,
    inp: PickedItemsInput,
    current_user: User = Depends(require_module(Module.WAREHOUSE, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE)),
):
    order = await get_scoped_or_404("orders", order_id, current_user)
    current = _status_value(order)
    if current not in (OrderStatus.NEW.value, OrderStatus.IN_PROGRESS.value):
        raise HTTPException(status_code=409, detail="Only new or in-progress orders can be packed")

    items = order.get("items", [])
    by_id = {i["product_id"]: i for i in items}
    for line in inp.items:
        item = by_id.get(line.product_id)
        if item is None:
            raise HTTPException(status_code=400, detail=f"Product not in order: {line.product_id}")
        if line.picked_qty is not None and (line.picked_qty < 0 or line.picked_qty > item.get("ordered_qty", 0)):
            raise HTTPException(
                status_code=400, detail=f"picked_qty for {item['name']} must be between 0 and {item.get('ordered_qty', 0)}"
            )
        if line.batches is not None:
            item["picked_batches"] = await _validated_picked_batches(item, line)
        elif line.picked_qty != item.get("picked_qty"):
            item.pop("picked_batches", None)  # a different quantity: let FEFO decide again
        item["picked_qty"] = line.picked_qty

    update: Dict[str, Any] = {"items": items}
    push = None
    if current == OrderStatus.NEW.value:
        update.update(
            status=OrderStatus.IN_PROGRESS.value,
            assigned_to_user_id=current_user.id,
            assigned_to_name=current_user.name,
        )
        push = {"status_history": _history_entry(current_user, current, OrderStatus.IN_PROGRESS.value, None)}
    result = await db.orders.update_one(
        {"id": order_id, "status": current}, {"$set": update, **({"$push": push} if push else {})}
    )
    if result.matched_count == 0:
        raise HTTPException(status_code=409, detail="Order was changed in the meantime")
    return _order_out(await db.orders.find_one({"id": order_id}, {"_id": 0}))


@api_router.get("/")
async def root():
    return {"message": "Easy Order API"}


@api_router.get("/app/update")
async def app_update():
    if not APP_UPDATE_FILE.exists():
        return {"enabled": False}
    try:
        return json.loads(APP_UPDATE_FILE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        raise HTTPException(status_code=503, detail="App update metadata is unavailable")


# ---------------- Statistics ----------------
_STATS_EXCLUDED_STATUSES = ("canceled", "rejected")


async def _compute_client_stats(client_id: str) -> ClientStats:
    client = await db.clients.find_one({"id": client_id}, {"_id": 0})
    client_name = client["name"] if client else "Unknown"
    # Canceled and rejected orders never became sales.
    orders = await db.orders.find(
        {"client_id": client_id, "status": {"$nin": list(_STATS_EXCLUDED_STATUSES)}}, {"_id": 0}
    ).to_list(None)
    customer_count = await db.customers.count_documents({"client_id": client_id})
    product_count = await db.products.count_documents({"client_id": client_id})

    total_net = 0.0
    total_vat = 0.0
    for o in orders:
        totals = compute_order_totals(o)
        total_net += totals["subtotal"]
        total_vat += totals["vat"]

    return ClientStats(
        client_id=client_id,
        client_name=client_name,
        order_count=len(orders),
        customer_count=customer_count,
        product_count=product_count,
        total_net=round(total_net, 2),
        total_vat=round(total_vat, 2),
        total_grand=round(total_net + total_vat, 2),
    )


# ---------------- Warehouse stock (PLAN_STANJE_MAGACINA.md) ----------------


class StockMovement(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    client_id: str
    product_id: str
    product_name: str = ""
    batch_id: Optional[str] = None
    expiry_date: Optional[str] = None
    delta: int
    balance_after: Optional[int] = None
    type: str  # receipt | adjustment | shipment | reversal
    order_id: Optional[str] = None
    note: Optional[str] = None
    created_by_user_id: Optional[str] = None
    created_by_name: Optional[str] = None
    created_at: str = Field(default_factory=now_iso)


class StockLineInput(BaseModel):
    product_id: str
    qty: int = Field(gt=0, le=1_000_000)
    # Required for products with expiry tracking (YYYY-MM-DD), ignored otherwise.
    expiry_date: Optional[str] = None


class StockReceiptInput(BaseModel):
    items: List[StockLineInput] = Field(min_length=1)
    note: Optional[str] = Field(default=None, max_length=300)


class BatchCount(BaseModel):
    expiry_date: Optional[str] = None  # None = undated stock
    counted_qty: int = Field(ge=0, le=1_000_000)


class StockAdjustmentInput(BaseModel):
    product_id: str
    # Plain products: counted_qty. Products with expiry tracking: `batches`,
    # the full count per expiry date (batches left out are set to 0).
    counted_qty: Optional[int] = Field(default=None, ge=0, le=1_000_000)
    batches: Optional[List[BatchCount]] = Field(default=None, max_length=100)
    note: str = Field(min_length=1, max_length=300)


class StockBatch(BaseModel):
    id: str
    client_id: str
    product_id: str
    expiry_date: Optional[str] = None
    qty: int = 0
    created_at: str = ""
    days_left: Optional[int] = None  # negative = already expired
    expired: bool = False


def _today() -> date:
    return datetime.now(timezone.utc).date()


def _normalize_expiry(value: Optional[str], required: bool = False) -> Optional[str]:
    """YYYY-MM-DD, not absurdly far away (typo guard); None = undated."""
    if value is None or not str(value).strip():
        if required:
            raise HTTPException(status_code=400, detail="expiry_date is required (YYYY-MM-DD)")
        return None
    try:
        parsed = datetime.strptime(str(value).strip(), "%Y-%m-%d").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="expiry_date must be YYYY-MM-DD")
    today = _today()
    if parsed.year < 2000 or parsed > today.replace(year=today.year + 20):
        raise HTTPException(status_code=400, detail="expiry_date is out of range")
    return parsed.isoformat()


def _batch_out(doc: dict) -> StockBatch:
    expiry = doc.get("expiry_date")
    days_left = (datetime.strptime(expiry, "%Y-%m-%d").date() - _today()).days if expiry else None
    return StockBatch(**{**doc, "days_left": days_left, "expired": days_left is not None and days_left < 0})


async def _expiry_info(scope: dict, product_id: Optional[str] = None) -> Dict[str, dict]:
    """Per product with batches in stock: pieces already expired, and the
    earliest expiry date among the pieces that are still fine."""
    query = {**scope, "qty": {"$gt": 0}}
    if product_id:
        query["product_id"] = product_id
    today = _today().isoformat()
    info: Dict[str, dict] = {}
    async for b in db.stock_batches.find(query, {"_id": 0, "product_id": 1, "expiry_date": 1, "qty": 1}):
        entry = info.setdefault(b["product_id"], {"expired_qty": 0, "next_expiry": None})
        expiry = b.get("expiry_date")
        if expiry is None:
            continue
        if expiry < today:
            entry["expired_qty"] += b["qty"]
        elif entry["next_expiry"] is None or expiry < entry["next_expiry"]:
            entry["next_expiry"] = expiry
    return info


async def _record_movement(
    product: dict,
    delta: int,
    kind: str,
    user: Optional[User],
    order_id: Optional[str] = None,
    note: Optional[str] = None,
    batch: Optional[dict] = None,
) -> None:
    fresh = await db.products.find_one({"id": product["id"]}, {"_id": 0, "stock_qty": 1})
    movement = StockMovement(
        client_id=product["client_id"],
        product_id=product["id"],
        product_name=product.get("name", ""),
        batch_id=(batch or {}).get("id"),
        expiry_date=(batch or {}).get("expiry_date"),
        delta=delta,
        balance_after=(fresh or {}).get("stock_qty"),
        type=kind,
        order_id=order_id,
        note=note,
        created_by_user_id=user.id if user else None,
        created_by_name=user.name if user else None,
    )
    await db.stock_movements.insert_one({**movement.dict(), "_id": movement.id})


async def _batch_inc(product: dict, expiry: Optional[str], delta: int) -> dict:
    """Adds `delta` pieces to the product's batch for that expiry date,
    creating the batch when it doesn't exist (expiry None = undated)."""
    new_id = str(uuid.uuid4())
    return await db.stock_batches.find_one_and_update(
        {"product_id": product["id"], "expiry_date": expiry},
        {
            "$inc": {"qty": delta},
            "$setOnInsert": {
                "_id": new_id,
                "id": new_id,
                "client_id": product["client_id"],
                "created_at": now_iso(),
            },
        },
        upsert=True,
        return_document=ReturnDocument.AFTER,
        projection={"_id": 0},
    )


async def _apply_delta(
    product: dict,
    delta: int,
    kind: str,
    user: Optional[User],
    order_id: Optional[str] = None,
    note: Optional[str] = None,
    expiry: Optional[str] = None,
    start_tracking: bool = False,
) -> None:
    """stock_qty += delta, the batch too for products with expiry tracking
    (so stock_qty stays the sum of the batches), plus one movement row.
    `start_tracking`: untracked stock (None) starts from 0 instead of failing."""
    if start_tracking:
        first = await db.products.update_one({"id": product["id"], "stock_qty": None}, {"$set": {"stock_qty": delta}})
        if first.matched_count == 0:
            await db.products.update_one({"id": product["id"]}, {"$inc": {"stock_qty": delta}})
    else:
        await db.products.update_one({"id": product["id"]}, {"$inc": {"stock_qty": delta}})
    batch = await _batch_inc(product, expiry, delta) if product.get("track_expiry") else None
    await _record_movement(product, delta, kind, user, order_id, note, batch)


async def _add_stock(
    product: dict, delta: int, kind: str, user: Optional[User], order_id=None, note=None, expiry: Optional[str] = None
) -> None:
    """Receipt-style change: starts tracking (from 0) when stock was untracked."""
    await _apply_delta(product, delta, kind, user, order_id, note, expiry, start_tracking=True)


async def _switch_expiry_tracking(product: dict, enable: bool) -> None:
    """Called when a manager flips products.track_expiry. Switching on turns
    the pieces already in stock into one undated batch (counted per expiry
    date later); switching off needs every dated batch to be empty."""
    if enable:
        stock = product.get("stock_qty")
        if stock:
            batch_id = str(uuid.uuid4())
            await db.stock_batches.update_one(
                {"product_id": product["id"], "expiry_date": None},
                {
                    "$set": {"qty": stock},
                    "$setOnInsert": {
                        "_id": batch_id,
                        "id": batch_id,
                        "client_id": product["client_id"],
                        "created_at": now_iso(),
                    },
                },
                upsert=True,
            )
        return
    if await db.stock_batches.find_one(
        {"product_id": product["id"], "expiry_date": {"$ne": None}, "qty": {"$ne": 0}}, {"_id": 1}
    ):
        raise HTTPException(status_code=409, detail="Write off or count all dated batches to zero before switching off")
    await db.stock_batches.delete_many({"product_id": product["id"]})


def _fefo_sort_key(batch: dict):
    # Undated stock first: its age is unknown, so use it up before dated stock.
    expiry = batch.get("expiry_date")
    return (expiry is not None, expiry or "")


async def _fefo_plan(product_id: str, qty: int) -> List[dict]:
    """Which batches to take `qty` pieces from: earliest expiry first, never
    an expired batch. A shortfall goes on undated stock (it may go negative,
    like stock itself)."""
    today = _today().isoformat()
    batches = await db.stock_batches.find({"product_id": product_id, "qty": {"$gt": 0}}, {"_id": 0}).to_list(None)
    usable = sorted((b for b in batches if b.get("expiry_date") is None or b["expiry_date"] >= today), key=_fefo_sort_key)
    plan: List[dict] = []
    left = qty
    for b in usable:
        if left <= 0:
            break
        take = min(left, b["qty"])
        plan.append({"batch_id": b["id"], "expiry_date": b.get("expiry_date"), "qty": take})
        left -= take
    if left > 0:
        shortfall = next((p for p in plan if p["expiry_date"] is None), None)
        if shortfall:
            shortfall["qty"] += left
        else:
            plan.append({"batch_id": None, "expiry_date": None, "qty": left})
    return plan


async def _move_stock_for_order(order: dict, direction: int, kind: str, user: User) -> None:
    """Shipment (-picked) / reversal (+picked) for products whose stock is
    tracked; untracked products are left alone. May go negative on purpose.
    Products with expiry tracking move batch by batch: the warehouse's own
    choice (item.picked_batches) or the earliest-expiring ones (FEFO); the
    result is kept on the order line."""
    changed_items = False
    for item in order.get("items", []):
        qty = item.get("picked_qty") or 0
        if qty <= 0:
            continue
        product = await db.products.find_one({"id": item["product_id"]}, {"_id": 0})
        if not product or product.get("stock_qty") is None:
            continue
        if not product.get("track_expiry"):
            await _apply_delta(product, direction * qty, kind, user, order_id=order["id"])
            continue
        picks = item.get("picked_batches") or []
        if sum(p["qty"] for p in picks) != qty:
            picks = await _fefo_plan(product["id"], qty) if direction < 0 else [
                {"batch_id": None, "expiry_date": None, "qty": qty}
            ]
        for pick in picks:
            await _apply_delta(
                product, direction * pick["qty"], kind, user, order_id=order["id"], expiry=pick.get("expiry_date")
            )
        if direction < 0:
            item["picked_batches"] = picks
            changed_items = True
    if changed_items:
        await db.orders.update_one({"id": order["id"]}, {"$set": {"items": order["items"]}})


@api_router.post("/stock/receipts")
async def stock_receipt(inp: StockReceiptInput, current_user: User = Depends(require_module(Module.STOCK, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))):
    products = []
    expiries = []
    for line in inp.items:
        product = await get_scoped_or_404("products", line.product_id, current_user)
        products.append(product)
        expiries.append(
            _normalize_expiry(line.expiry_date, required=True) if product.get("track_expiry") else None
        )
    for product, line, expiry in zip(products, inp.items, expiries):
        await _add_stock(product, line.qty, "receipt", current_user, note=inp.note, expiry=expiry)
    return {"ok": True, "count": len(products)}


async def _set_stock_count(product: dict, counted: int, user: Optional[User], note: Optional[str]) -> int:
    """Plain count. For a product with expiry tracking the difference lands on
    the undated batch (used by the Excel import, which has no dates)."""
    delta = counted - (product.get("stock_qty") or 0)
    if product.get("track_expiry"):
        await _apply_delta(product, delta, "adjustment", user, note=note, start_tracking=True)
        return delta
    await db.products.update_one({"id": product["id"]}, {"$set": {"stock_qty": counted}})
    await _record_movement(product, delta, "adjustment", user, note=note)
    return delta


def _validated_count(product: dict, counted: Optional[int], batches: Optional[List[BatchCount]]):
    """Checks that a count has the right shape for the product: counted_qty
    for plain ones, `batches` for products with expiry tracking."""
    if product.get("track_expiry"):
        if batches is None or counted is not None:
            raise HTTPException(status_code=400, detail=f"{product['name']}: count per expiry date (batches)")
        seen = set()
        out = []
        for b in batches:
            expiry = _normalize_expiry(b.expiry_date)
            if expiry in seen:
                raise HTTPException(status_code=400, detail=f"{product['name']}: duplicate expiry date")
            seen.add(expiry)
            out.append((expiry, b.counted_qty))
        return out
    if counted is None or batches is not None:
        raise HTTPException(status_code=400, detail=f"{product['name']}: counted_qty is required")
    return counted


async def _apply_count(product: dict, count, user: Optional[User], note: Optional[str]) -> int:
    if not product.get("track_expiry"):
        return await _set_stock_count(product, count, user, note)
    existing = {b.get("expiry_date"): b["qty"] async for b in db.stock_batches.find({"product_id": product["id"]})}
    wanted = dict(count)
    total = 0
    for expiry in sorted(set(existing) | set(wanted), key=lambda e: (e is not None, e or "")):
        delta = wanted.get(expiry, 0) - existing.get(expiry, 0)
        if delta:
            await _apply_delta(product, delta, "adjustment", user, note=note, expiry=expiry, start_tracking=True)
            total += delta
    return total


@api_router.post("/stock/adjustments")
async def stock_adjustment(inp: StockAdjustmentInput, current_user: User = Depends(require_module(Module.STOCK, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))):
    product = await get_scoped_or_404("products", inp.product_id, current_user)
    count = _validated_count(product, inp.counted_qty, inp.batches)
    delta = await _apply_count(product, count, current_user, inp.note)
    return {"ok": True, "delta": delta}


class StockCountLine(BaseModel):
    product_id: str
    counted_qty: Optional[int] = Field(default=None, ge=0, le=1_000_000)
    batches: Optional[List[BatchCount]] = Field(default=None, max_length=100)


class StockBatchAdjustmentInput(BaseModel):
    items: List[StockCountLine] = Field(min_length=1, max_length=500)
    note: str = Field(min_length=1, max_length=300)


@api_router.post("/stock/adjustments/batch")
async def stock_adjustment_batch(inp: StockBatchAdjustmentInput, current_user: User = Depends(require_module(Module.STOCK, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))):
    """A whole stocktake in one request; every product is checked first so a
    bad id doesn't leave the count half applied."""
    seen = set()
    products = []
    counts = []
    for line in inp.items:
        if line.product_id in seen:
            raise HTTPException(status_code=400, detail="Duplicate product in stocktake")
        seen.add(line.product_id)
        product = await get_scoped_or_404("products", line.product_id, current_user)
        products.append(product)
        counts.append(_validated_count(product, line.counted_qty, line.batches))
    deltas = []
    for product, count in zip(products, counts):
        deltas.append(await _apply_count(product, count, current_user, inp.note))
    return {"ok": True, "count": len(deltas), "deltas": deltas}


@api_router.get("/products/{product_id}/batches", response_model=List[StockBatch])
async def product_batches(product_id: str, current_user: User = Depends(require_module(Module.EXPIRY, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))):
    """Batches of one product with pieces in stock, earliest expiry first
    (undated stock first). Warehouse/admin/superadmin only."""
    await get_scoped_or_404("products", product_id, current_user)
    docs = await db.stock_batches.find({"product_id": product_id, "qty": {"$ne": 0}}, {"_id": 0}).to_list(None)
    return [_batch_out(d) for d in sorted(docs, key=_fefo_sort_key)]


class WriteOffInput(BaseModel):
    note: Optional[str] = Field(default=None, max_length=300)


@api_router.post("/stock/batches/{batch_id}/writeoff")
async def writeoff_batch(batch_id: str, inp: WriteOffInput, current_user: User = Depends(require_module(Module.EXPIRY, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE))):
    """Writes the pieces of a batch off the stock (e.g. expired goods)."""
    batch = await get_scoped_or_404("stock_batches", batch_id, current_user)
    if batch["qty"] <= 0:
        raise HTTPException(status_code=400, detail="Batch has no pieces to write off")
    product = await db.products.find_one({"id": batch["product_id"]}, {"_id": 0})
    if not product:
        raise HTTPException(status_code=404, detail="Not found")
    note = (inp.note or "").strip() or "Otpis - istekao rok"
    await _apply_delta(product, -batch["qty"], "adjustment", current_user, note=note, expiry=batch.get("expiry_date"))
    return {"ok": True, "written_off": batch["qty"]}


class ExpiringItem(BaseModel):
    batch_id: str
    product_id: str
    product_name: str
    barcode: Optional[str] = None
    expiry_date: str
    qty: int
    days_left: int  # negative = already expired
    level: str  # "expired" or the threshold (in days) this batch has reached


class ExpiringResponse(BaseModel):
    thresholds: List[int]
    total: int
    counts: Dict[str, int]  # per level, for the badge
    items: List[ExpiringItem]


@api_router.get("/stock/expiring", response_model=ExpiringResponse)
async def stock_expiring(current_user: User = Depends(require_module(Module.EXPIRY, Role.ADMIN, Role.WAREHOUSE))):
    """Batches that are expired or reach one of the client's alert thresholds
    (default 30/15/5 days). For the client's admin and warehouse only: no
    superadmin (not tied to a warehouse), no sales reps."""
    client = await db.clients.find_one({"id": current_user.client_id}, {"_id": 0, "expiry_alert_days": 1}) or {}
    thresholds = sorted(client.get("expiry_alert_days") or setting("expiry_alert_days"))
    limit = (_today() + timedelta(days=thresholds[-1])).isoformat()
    batches = await db.stock_batches.find(
        {"client_id": current_user.client_id, "qty": {"$gt": 0}, "expiry_date": {"$ne": None, "$lte": limit}},
        {"_id": 0},
    ).to_list(None)
    products = {
        p["id"]: p
        for p in await db.products.find(
            {"id": {"$in": list({b["product_id"] for b in batches})}}, {"_id": 0, "id": 1, "name": 1, "barcode": 1}
        ).to_list(None)
    }
    items = []
    counts: Dict[str, int] = {"expired": 0, **{str(t): 0 for t in thresholds}}
    for b in sorted(batches, key=lambda b: b["expiry_date"]):
        product = products.get(b["product_id"])
        if not product:
            continue
        days_left = _batch_out(b).days_left
        level = "expired" if days_left < 0 else str(next(t for t in thresholds if days_left <= t))
        counts[level] += 1
        items.append(
            ExpiringItem(
                batch_id=b["id"],
                product_id=b["product_id"],
                product_name=product["name"],
                barcode=product.get("barcode"),
                expiry_date=b["expiry_date"],
                qty=b["qty"],
                days_left=days_left,
                level=level,
            )
        )
    return ExpiringResponse(thresholds=thresholds, total=len(items), counts=counts, items=items)


@api_router.get("/stock/movements", response_model=List[StockMovement])
async def stock_movements(
    product_id: Optional[str] = None,
    limit: int = Query(100, ge=1, le=500),
    current_user: User = Depends(require_module(Module.STOCK, Role.SUPERADMIN, Role.ADMIN, Role.WAREHOUSE)),
):
    query = _scope_query(current_user)
    if product_id:
        query["product_id"] = product_id
    docs = await db.stock_movements.find(query, {"_id": 0}).sort("created_at", -1).limit(limit).to_list(None)
    return [StockMovement(**d) for d in docs]


# ---------------- Reports ----------------
class StatusCount(BaseModel):
    count: int = 0
    grand: float = 0  # incl. VAT; by packed quantity once shipped


class PersonReport(BaseModel):
    user_id: Optional[str] = None
    name: str
    total: int = 0
    shipped: int = 0
    rejected: int = 0
    canceled: int = 0


class OrdersReport(BaseModel):
    from_date: Optional[str] = None
    to_date: Optional[str] = None
    order_count: int
    by_status: Dict[str, StatusCount]
    by_creator: List[PersonReport]  # sales reps: orders they placed
    by_handler: List[PersonReport]  # who shipped/rejected orders (warehouse)
    avg_hours_to_ship: Optional[float] = None


def _last_transition(order: dict, to_status: str) -> Optional[dict]:
    for entry in reversed(order.get("status_history") or []):
        if entry.get("to_status") == to_status:
            return entry
    return None


@api_router.get("/reports/orders", response_model=OrdersReport)
async def orders_report(
    from_date: Optional[str] = Query(None, description="YYYY-MM-DD, inclusive (by created_at)"),
    to_date: Optional[str] = Query(None, description="YYYY-MM-DD, inclusive"),
    client_id: Optional[str] = None,
    current_user: User = Depends(require_module(Module.REPORTS, Role.SUPERADMIN, Role.ADMIN)),
):
    """Order counts per status, per sales rep and per warehouse handler for a
    period. Admin: own client. Superadmin: all clients, or one via client_id."""
    query = _scope_query(current_user)
    if current_user.role == Role.SUPERADMIN and client_id:
        query = {"client_id": client_id}
    date_range = {}
    if from_date:
        date_range["$gte"] = _parse_date(from_date, "from_date").isoformat()
    if to_date:
        date_range["$lt"] = (_parse_date(to_date, "to_date") + timedelta(days=1)).isoformat()
    if date_range:
        query["created_at"] = date_range
    orders = await db.orders.find(query, {"_id": 0}).to_list(None)

    by_status: Dict[str, StatusCount] = {s.value: StatusCount() for s in OrderStatus}
    creators: Dict[str, PersonReport] = {}
    handlers: Dict[str, PersonReport] = {}
    ship_hours: List[float] = []

    def person(bucket: Dict[str, PersonReport], user_id: Optional[str], name: Optional[str]) -> PersonReport:
        key = user_id or name or "?"
        if key not in bucket:
            bucket[key] = PersonReport(user_id=user_id, name=name or "-")
        return bucket[key]

    for order in orders:
        status = _status_value(order)
        bucket = by_status.setdefault(status, StatusCount())
        bucket.count += 1
        bucket.grand = round(bucket.grand + compute_order_totals(order)["grand"], 2)

        creator = person(creators, order.get("created_by_user_id"), order.get("created_by_name"))
        creator.total += 1
        if status in ("shipped", "rejected", "canceled"):
            setattr(creator, status, getattr(creator, status) + 1)

        for final in ("shipped", "rejected"):
            if status == final:
                entry = _last_transition(order, final)
                if entry:
                    h = person(handlers, entry.get("changed_by_user_id"), entry.get("changed_by_name"))
                    h.total += 1
                    setattr(h, final, getattr(h, final) + 1)
        if status == "shipped" and order.get("shipped_at"):
            try:
                delta = datetime.fromisoformat(order["shipped_at"]) - datetime.fromisoformat(order["created_at"])
                ship_hours.append(delta.total_seconds() / 3600)
            except ValueError:
                pass

    return OrdersReport(
        from_date=from_date,
        to_date=to_date,
        order_count=len(orders),
        by_status=by_status,
        by_creator=sorted(creators.values(), key=lambda p: -p.total),
        by_handler=sorted(handlers.values(), key=lambda p: -p.total),
        avg_hours_to_ship=round(sum(ship_hours) / len(ship_hours), 1) if ship_hours else None,
    )


@api_router.get("/stats/overview", response_model=StatsResponse)
async def stats_overview(current_user: User = Depends(require_roles(Role.SUPERADMIN, Role.ADMIN))):
    if current_user.role == Role.SUPERADMIN:
        client_docs = await db.clients.find({}, {"_id": 0, "id": 1}).to_list(None)
        by_client = [await _compute_client_stats(c["id"]) for c in client_docs]
        totals = ClientStats(
            client_id="",
            client_name="All Clients",
            order_count=sum(c.order_count for c in by_client),
            customer_count=sum(c.customer_count for c in by_client),
            product_count=sum(c.product_count for c in by_client),
            total_net=round(sum(c.total_net for c in by_client), 2),
            total_vat=round(sum(c.total_vat for c in by_client), 2),
            total_grand=round(sum(c.total_grand for c in by_client), 2),
        )
        return StatsResponse(scope="global", totals=totals, by_client=by_client)

    stats = await _compute_client_stats(current_user.client_id)
    return StatsResponse(scope="client", totals=stats, by_client=[stats])


def _parse_dashboard_date(value: Optional[str], field_name: str) -> Optional[datetime]:
    if not value:
        return None
    try:
        return datetime.fromisoformat(value).replace(tzinfo=timezone.utc)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"Invalid {field_name}; use YYYY-MM-DD")


@api_router.get("/stats/dashboard", response_model=DashboardStatsResponse)
async def stats_dashboard(
    from_date: Optional[str] = None,
    to_date: Optional[str] = None,
    current_user: User = Depends(require_roles(Role.SUPERADMIN, Role.ADMIN)),
):
    start = _parse_dashboard_date(from_date, "from_date")
    end = _parse_dashboard_date(to_date, "to_date")
    if start and end and start > end:
        raise HTTPException(status_code=400, detail="from_date must be before to_date")

    query = {} if current_user.role == Role.SUPERADMIN else {"client_id": current_user.client_id}
    # created_at is an ISO-8601 UTC string, so a string range works in Mongo
    # and avoids loading every order ever placed. The loop below re-checks
    # with real datetimes for any legacy format.
    created_range: Dict[str, str] = {}
    if start:
        created_range["$gte"] = start.isoformat()
    if end:
        created_range["$lt"] = (end + timedelta(days=1)).isoformat()
    order_query = {**query, "created_at": created_range} if created_range else query
    orders = await db.orders.find(order_query, {"_id": 0}).sort("created_at", -1).to_list(None)
    filtered_orders = []
    for order in orders:
        try:
            created_at = datetime.fromisoformat(order["created_at"].replace("Z", "+00:00"))
        except (KeyError, ValueError):
            continue
        if start and created_at < start:
            continue
        if end and created_at >= end + timedelta(days=1):
            continue
        if order.get("status") in _STATS_EXCLUDED_STATUSES:
            continue
        filtered_orders.append(order)

    product_rows: Dict[str, Dict[str, Any]] = {}
    customer_rows: Dict[str, Dict[str, Any]] = {}
    revenue_rows: Dict[str, Dict[str, Any]] = {}
    total_net = 0.0
    total_vat = 0.0
    total_items = 0

    for order in filtered_orders:
        totals = compute_order_totals(order)
        total_net += totals["subtotal"]
        total_vat += totals["vat"]
        total_items += len(order.get("items", []))
        period = order["created_at"][:7]
        revenue = revenue_rows.setdefault(period, {"period": period, "order_count": 0, "total_net": 0.0, "total_vat": 0.0})
        revenue["order_count"] += 1
        revenue["total_net"] += totals["subtotal"]
        revenue["total_vat"] += totals["vat"]

        customer_id = order.get("customer_id", "")
        customer = customer_rows.setdefault(customer_id, {
            "customer_id": customer_id,
            "customer_name": order.get("customer_name", "Unknown"),
            "order_count": 0,
            "total_grand": 0.0,
        })
        customer["order_count"] += 1
        customer["total_grand"] += totals["grand"]

        order_shipped = order.get("status") == OrderStatus.SHIPPED.value
        for item in order.get("items", []):
            product_id = item.get("product_id", item.get("name", "unknown"))
            product = product_rows.setdefault(product_id, {
                "product_id": product_id,
                "name": item.get("name", "Unknown"),
                "manufacturer": item.get("manufacturer", ""),
                # Orders come newest first, so this is the product's current unit.
                "unit_symbol": item.get("unit_symbol") or "kom",
                "unit_factor": item.get("unit_factor") or 1,
                "ordered_qty": 0,
                "order_count": 0,
                "total_net": 0.0,
                "total_vat": 0.0,
            })
            # Same basis as the order totals: what was packed once shipped, else what was ordered.
            product["ordered_qty"] += int(effective_qty(item, order_shipped))
            product["order_count"] += 1
            product["total_net"] += line_net(item, order_shipped)
            product["total_vat"] += line_vat(item, order_shipped)

    customer_count = await db.customers.count_documents(query)
    summary = DashboardSummary(
        order_count=len(filtered_orders),
        customer_count=customer_count,
        product_count=await db.products.count_documents(query),
        active_customer_count=len(customer_rows),
        total_net=round(total_net, 2),
        total_vat=round(total_vat, 2),
        total_grand=round(total_net + total_vat, 2),
        average_order_value=round((total_net + total_vat) / len(filtered_orders), 2) if filtered_orders else 0,
        average_items_per_order=round(total_items / len(filtered_orders), 2) if filtered_orders else 0,
    )

    revenue_by_period = [
        RevenuePoint(**{**row, "total_net": round(row["total_net"], 2), "total_vat": round(row["total_vat"], 2), "total_grand": round(row["total_net"] + row["total_vat"], 2)})
        for row in sorted(revenue_rows.values(), key=lambda row: row["period"])
    ]
    top_products = [
        ProductSalesStats(**{**row, "total_net": round(row["total_net"], 2), "total_vat": round(row["total_vat"], 2), "total_grand": round(row["total_net"] + row["total_vat"], 2)})
        for row in sorted(product_rows.values(), key=lambda row: row["total_net"] + row["total_vat"], reverse=True)[:10]
    ]
    top_customers = [CustomerSalesStats(**row) for row in sorted(customer_rows.values(), key=lambda row: row["total_grand"], reverse=True)[:10]]
    recent_orders = [
        RecentOrderStats(
            id=order["id"],
            customer_name=order.get("customer_name", "Unknown"),
            item_count=len(order.get("items", [])),
            total_grand=round(compute_order_totals(order)["grand"], 2),
            created_at=order["created_at"],
        )
        for order in filtered_orders[:10]
    ]

    return DashboardStatsResponse(
        scope="global" if current_user.role == Role.SUPERADMIN else "client",
        from_date=from_date,
        to_date=to_date,
        summary=summary,
        revenue_by_period=revenue_by_period,
        top_products=top_products,
        top_customers=top_customers,
        recent_orders=recent_orders,
    )


@api_router.get("/stats/clients/{client_id}", response_model=ClientStats)
async def stats_client(client_id: str, current_user: User = Depends(require_roles(Role.SUPERADMIN))):
    if not await db.clients.find_one({"id": client_id}, {"_id": 0}):
        raise HTTPException(status_code=404, detail="Client not found")
    return await _compute_client_stats(client_id)


# ---------------- Seed ----------------
DEMO_CUSTOMERS = [
    {"name": "Maxi Market d.o.o.", "address": "Bulevar oslobođenja 1, Novi Sad", "pib": "101234567",
     "phone": "+381 21 123456", "email": "nabavka@maxi-demo.rs"},
    {"name": "Delikates Prodavnica", "address": "Knez Mihailova 10, Beograd", "pib": "107654321",
     "phone": "+381 11 7654321", "email": "info@delikates-demo.rs"},
]
# Values the older integration tests (test_easy_order/iteration3/iteration4) assert on.
DEMO_PRODUCTS = [
    {"name": "Ulje Bundeve 250ml", "manufacturer": "Bački Dukat", "price_no_vat": 450, "vat_rate": 20,
     "discount": 5, "discounts": [0, 5, 10], "pieces_per_package": 12, "boxes_per_transport": 40},
    {"name": "Kokosovo ulje 150ml", "manufacturer": "Bački Dukat", "price_no_vat": 380, "vat_rate": 20,
     "discount": 0, "discounts": [0, 3, 7], "pieces_per_package": 12, "boxes_per_transport": 50},
    {"name": "Jabukovo sirce 1L", "manufacturer": "Zdrava Hrana", "price_no_vat": 220, "vat_rate": 10,
     "discount": 10, "discounts": [0, 10, 15], "pieces_per_package": 6, "boxes_per_transport": 60},
]


async def _seed_demo_data() -> None:
    """Local/CI only (never IS_PRODUCTION): a demo client with an admin,
    customers and products, so tests run against an empty database."""
    client = await db.clients.find_one({"name": "Demo Wholesaler"}, {"_id": 0})
    if not client:
        client_obj = Client(name="Demo Wholesaler")
        client = client_obj.dict()
        await db.clients.insert_one({**client, "_id": client_obj.id})
    client_id = client["id"]

    if not await find_user_by_email(DEMO_ADMIN_EMAIL):
        admin = User(email=normalize_email(DEMO_ADMIN_EMAIL), name="Demo Admin", role=Role.ADMIN, client_id=client_id)
        raw = {**admin.dict(), "token_version": 0, "password_hash": hash_password(DEMO_ADMIN_PASSWORD)}
        await db.users.insert_one({**raw, "_id": admin.id})

    for data in DEMO_CUSTOMERS:
        if not await db.customers.find_one({"client_id": client_id, "name": data["name"]}):
            obj = Customer(client_id=client_id, **data)
            await db.customers.insert_one({**obj.dict(), "_id": obj.id})
    for data in DEMO_PRODUCTS:
        if not await db.products.find_one({"client_id": client_id, "name": data["name"]}):
            obj = Product(client_id=client_id, **data)
            await db.products.insert_one({**obj.dict(), "_id": obj.id})


async def seed_data():
    # Opt-in, not just "not production": a local backend may point at the
    # real Atlas database (backend/.env), and a demo admin with a public
    # password must never land there.
    if SEED_DEMO_DATA:
        if IS_PRODUCTION:
            logger.error("SEED_DEMO_DATA is ignored in production")
        else:
            await _seed_demo_data()

    if not await find_user_by_email(SUPERADMIN_EMAIL):
        if SUPERADMIN_PASSWORD == DEFAULT_SUPERADMIN_PASSWORD:
            if IS_PRODUCTION:
                # The default password is public (it's in the repo), so never
                # create a production superadmin with it. Skip instead of
                # crashing so an already-seeded deployment keeps running.
                logger.error("SUPERADMIN_PASSWORD not set - refusing to create superadmin with the default password")
                return
            logger.warning(
                "SUPERADMIN_PASSWORD not set in backend/.env - using an insecure default password."
            )
        superadmin = User(email=normalize_email(SUPERADMIN_EMAIL), name="Super Admin", role=Role.SUPERADMIN, client_id=None)
        raw = superadmin.dict()
        raw["token_version"] = 0
        raw["password_hash"] = hash_password(SUPERADMIN_PASSWORD)
        await db.users.insert_one({**raw, "_id": superadmin.id})


app = FastAPI(lifespan=lifespan)
app.include_router(api_router)
# Not every OS registers .apk (Windows doesn't), and without it StaticFiles
# serves the APK as text/plain, which some Android browsers won't install.
mimetypes.add_type("application/vnd.android.package-archive", ".apk")
APP_DOWNLOADS_DIR = ROOT_DIR / "public" / "app"
if APP_DOWNLOADS_DIR.is_dir():
    app.mount("/app", StaticFiles(directory=APP_DOWNLOADS_DIR), name="app-downloads")

app.add_middleware(
    CORSMiddleware,
    allow_credentials=False,
    allow_origins=CORS_ORIGINS,
    allow_methods=["*"],
    allow_headers=["*"],
)
