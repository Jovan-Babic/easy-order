"""Modules (PLAN_MODULI.md): the superadmin turns features on per client."""
import io
import uuid

import pytest
import requests
from openpyxl import Workbook, load_workbook

from helpers import API, TEST_PASSWORD, activate_invited_user, login


class S:
    client_ids = []


def _new_client(superadmin_client, modules=None):
    suffix = uuid.uuid4().hex[:8]
    email = f"test-mod-{suffix}@easyorder.dev"
    body = {"name": f"TEST_Mod_{suffix}", "admin_name": "TEST Mod Admin", "admin_email": email}
    if modules is not None:
        body["modules"] = modules
    r = superadmin_client.post(f"{API}/clients", json=body)
    assert r.status_code == 200, r.text
    S.client_ids.append(r.json()["client"]["id"])
    admin = activate_invited_user(email, r.json()["temporary_password"])
    return r.json()["client"], admin


@pytest.fixture(scope="module", autouse=True)
def cleanup(superadmin_client):
    yield
    for cid in S.client_ids:
        superadmin_client.delete(f"{API}/clients/{cid}")


def _invite(admin, role):
    email = f"test-mod-{role}-{uuid.uuid4().hex[:8]}@easyorder.dev"
    r = admin.post(f"{API}/users", json={"email": email, "name": f"TEST {role}", "role": role})
    return r, email


def test_default_client_has_every_module(superadmin_client):
    client, admin = _new_client(superadmin_client)
    assert client["modules"] == ["warehouse", "stock", "expiry", "reports"]
    assert admin.get(f"{API}/auth/me").json()["modules"] == client["modules"]


def test_login_returns_modules(superadmin_client):
    client, _ = _new_client(superadmin_client, ["warehouse"])
    assert client["modules"] == ["warehouse"]
    # superadmin is not tied to a client: every module
    me = superadmin_client.get(f"{API}/auth/me").json()
    assert me["modules"] == ["warehouse", "stock", "expiry", "reports"]


def test_dependencies_are_validated(superadmin_client):
    suffix = uuid.uuid4().hex[:8]
    for bad in (["stock"], ["warehouse", "expiry"], ["expiry"]):
        r = superadmin_client.post(
            f"{API}/clients",
            json={"name": f"TEST_Mod_bad_{suffix}", "admin_name": "x", "admin_email": f"bad-{suffix}@easyorder.dev", "modules": bad},
        )
        assert r.status_code == 400, (bad, r.text)
    client, _ = _new_client(superadmin_client, [])
    assert superadmin_client.put(
        f"{API}/clients/{client['id']}", json={"name": client["name"], "modules": ["stock"]}
    ).status_code == 400
    assert superadmin_client.post(
        f"{API}/clients",
        json={"name": "x", "admin_name": "x", "admin_email": f"unk-{suffix}@easyorder.dev", "modules": ["bogus"]},
    ).status_code == 422


def test_update_keeps_modules_unless_sent(superadmin_client):
    client, _ = _new_client(superadmin_client, ["warehouse"])
    kept = superadmin_client.put(f"{API}/clients/{client['id']}", json={"name": "TEST_Mod_renamed"})
    assert kept.status_code == 200 and kept.json()["modules"] == ["warehouse"]
    changed = superadmin_client.put(
        f"{API}/clients/{client['id']}", json={"name": "TEST_Mod_renamed", "modules": ["warehouse", "stock"]}
    )
    assert changed.json()["modules"] == ["warehouse", "stock"]


def test_base_client_is_locked_to_the_base(superadmin_client):
    client, admin = _new_client(superadmin_client, [])
    assert admin.get(f"{API}/auth/me").json()["modules"] == []

    # no warehouse user without the module; operator is fine
    r, _ = _invite(admin, "warehouse")
    assert r.status_code == 403 and "warehouse" in r.json()["detail"]
    r, email = _invite(admin, "operator")
    assert r.status_code == 200
    operator = activate_invited_user(email, r.json()["temporary_password"])

    # fields of missing modules are ignored on products
    p = admin.post(
        f"{API}/products",
        json={"name": "TEST_Mod_P", "price_no_vat": 10, "barcode": "123456", "package_barcode": "654321",
              "pieces_per_package": 6, "track_expiry": True},
    ).json()
    assert p["barcode"] is None and p["package_barcode"] is None and p["track_expiry"] is False
    listed = admin.get(f"{API}/products").json()[0]
    assert listed["stock_qty"] is None and listed["available_qty"] is None and listed["barcode"] is None

    # module endpoints are 403
    assert admin.get(f"{API}/products/by-barcode/123456").status_code == 403
    assert admin.post(f"{API}/stock/receipts", json={"items": [{"product_id": p["id"], "qty": 1}]}).status_code == 403
    assert admin.get(f"{API}/stock/movements").status_code == 403
    assert admin.get(f"{API}/stock/expiring").status_code == 403
    assert admin.get(f"{API}/products/{p['id']}/batches").status_code == 403
    assert admin.get(f"{API}/reports/orders").status_code == 403
    assert admin.post(f"{API}/products/quick", json={"name": "TEST_q"}).status_code == 403

    # orders: create and cancel work, the warehouse flow does not
    customer = admin.post(f"{API}/customers", json={"name": "TEST_Mod_C"}).json()
    line = {"customer_id": customer["id"], "items": [{"product_id": p["id"], "ordered_qty": 2}]}
    order = operator.post(f"{API}/orders", json=line).json()
    assert admin.patch(
        f"{API}/orders/{order['id']}/items", json={"items": [{"product_id": p["id"], "picked_qty": 2}]}
    ).status_code == 403
    r = admin.post(f"{API}/orders/{order['id']}/status", json={"status": "in_progress", "note": "x"})
    assert r.status_code == 403
    r = admin.post(f"{API}/orders/{order['id']}/status", json={"status": "shipped", "note": "x"})
    assert r.status_code == 403
    r = operator.post(f"{API}/orders/{order['id']}/status", json={"status": "canceled"})
    assert r.status_code == 200 and r.json()["status"] == "canceled"


def test_import_drops_columns_of_missing_modules(superadmin_client):
    client, admin = _new_client(superadmin_client, ["warehouse"])  # no stock, no expiry
    wb = Workbook()
    ws = wb.active
    ws.append(["Naziv", "Barkod", "Cena bez PDV", "Stanje (komada)", "Prati rok (da/ne)"])
    ws.append(["TEST_Mod_Imp", "5551234", 10, 40, "da"])
    buf = io.BytesIO()
    wb.save(buf)
    r = admin.post(
        f"{API}/products/import",
        params={"dry_run": "false"},
        files={"file": ("uvoz.xlsx", buf.getvalue(), "application/octet-stream")},
        headers={"Content-Type": None},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ignored_columns"] == ["stock_qty", "track_expiry"] and body["summary"]["create"] == 1
    product = admin.get(f"{API}/products/by-barcode/5551234").json()
    assert product["stock_qty"] is None and product["track_expiry"] is False

    template = admin.get(f"{API}/products/import-template")
    headers = [c.value for c in load_workbook(io.BytesIO(template.content)).active[1]]
    assert "Barkod" in headers and "Stanje (komada)" not in headers and "Prati rok (da/ne)" not in headers


def test_enabling_a_module_takes_effect_immediately(superadmin_client):
    client, admin = _new_client(superadmin_client, [])
    assert admin.get(f"{API}/reports/orders").status_code == 403
    superadmin_client.put(f"{API}/clients/{client['id']}", json={"name": client["name"], "modules": ["reports"]})
    assert admin.get(f"{API}/reports/orders").status_code == 200
    superadmin_client.put(f"{API}/clients/{client['id']}", json={"name": client["name"], "modules": []})
    assert admin.get(f"{API}/reports/orders").status_code == 403
    # turning the warehouse module on lets the admin create warehouse users
    superadmin_client.put(f"{API}/clients/{client['id']}", json={"name": client["name"], "modules": ["warehouse"]})
    r, _ = _invite(admin, "warehouse")
    assert r.status_code == 200


def test_warehouse_module_without_stock_does_not_touch_stock(superadmin_client):
    client, admin = _new_client(superadmin_client, ["warehouse"])
    r, email = _invite(admin, "warehouse")
    wh = activate_invited_user(email, r.json()["temporary_password"])
    p = admin.post(f"{API}/products", json={"name": "TEST_Mod_W", "price_no_vat": 5}).json()
    customer = admin.post(f"{API}/customers", json={"name": "TEST_Mod_WC"}).json()
    order = admin.post(
        f"{API}/orders", json={"customer_id": customer["id"], "items": [{"product_id": p["id"], "ordered_qty": 3}]}
    ).json()
    set_prefix = superadmin_client.put(
        f"{API}/clients/{client['id']}", json={"name": client["name"], "invoice_prefix": "MOD"}
    )
    assert set_prefix.status_code == 200
    assert wh.patch(
        f"{API}/orders/{order['id']}/items", json={"items": [{"product_id": p["id"], "picked_qty": 3}]}
    ).status_code == 200
    r = wh.post(f"{API}/orders/{order['id']}/status", json={"status": "shipped"})
    assert r.status_code == 200 and r.json()["status"] == "shipped" and r.json()["invoice_number"].startswith("MOD/")
    assert wh.get(f"{API}/stock/movements").status_code == 403


def test_downloaded_template_headers_are_recognized(api_client):
    """Round trip: the example file's own headers (with bracketed hints) map to fields."""
    template = api_client.get(f"{API}/products/import-template")
    wb = load_workbook(io.BytesIO(template.content))
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        for cell in row:  # fresh codes/names so nothing collides with existing data
            if isinstance(cell.value, str) and cell.column_letter in ("A", "B", "J"):
                cell.value = f"{cell.value}-{uuid.uuid4().hex[:6]}"
    buf = io.BytesIO()
    wb.save(buf)
    r = api_client.post(
        f"{API}/products/import",
        params={"dry_run": "true"},
        files={"file": ("uvoz.xlsx", buf.getvalue(), "application/octet-stream")},
        headers={"Content-Type": None},
    )
    assert r.status_code == 200, r.text
    assert r.json()["ignored_columns"] == [] and r.json()["summary"]["error"] == 0, r.json()
